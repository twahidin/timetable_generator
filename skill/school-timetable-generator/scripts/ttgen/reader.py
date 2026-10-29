"""Read a timetable setup workbook back into a spec, and work out what the formulas show (codes, periods, lesson pattern)
in Python so a file never opened in Excel still reads correctly.

Reads the current layout (Teacher availability, Fixed lessons and rule rows on Settings, as in setup template v9 and later
template) and the older one (a 7 Rules sheet). Everything that constrains the timetable becomes a rule on spec['rules']."""
import datetime, re
import openpyxl
from .spec import defaults, lesson_split, SETTING_RULES

R0 = 6
COMPUTED = ("Periods given", "Spare", "(key", "Lesson rows", "Home room of", "Lessons listed", "Used by", "Group code", "Groups in", "Periods", "Lesson length",
            "Full", "Short", "Doubles", "Triples", "Singles", "Room kind", "Teacher periods", "Check", "Times", "Students in", "Difference", "Longest",
            "Total load", "Role guide", "% of")

def _v(x):
    if isinstance(x, str):
        x = x.strip()
        return x or None
    return x

def _table(ws, first):
    """rows of the table whose header row starts with `first` (searched in rows 1-8), as dicts by header. A header
    with a bracket, e.g. 'Class (2027)', also answers to its plain name when no column has that name."""
    hr = next((r for r in range(1, 9) if str(ws.cell(r, 1).value or "").strip() == first), 5)
    hdr = [(_v(c.value) or "") for c in ws[hr]]
    alias = {h.split(" (")[0]: h for h in hdr if " (" in str(h) and h.split(" (")[0] not in hdr}
    out = []
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        d = {h: _v(v) for h, v in zip(hdr, row) if h}
        if any(v not in (None, "") and not (isinstance(v, str) and v.startswith("=")) for k, v in d.items() if not str(k).startswith(COMPUTED)):
            for a, h in alias.items(): d[a] = d.get(h)
            out.append(d)
    return hdr, out

def _first_col(row, prefix, n):
    return [row.get(f"{prefix} {i}") for i in range(1, n + 1) if row.get(f"{prefix} {i}")]

def _minutes(t):
    """a clock time (time, datetime, '13:40', or an Excel day fraction) -> minutes after midnight"""
    if isinstance(t, datetime.datetime): t = t.time()
    if isinstance(t, datetime.time): return t.hour * 60 + t.minute
    if isinstance(t, (int, float)): return round(float(t) * 1440) if t < 1 else int(t) * 60
    h, m = (str(t).strip().replace(".", ":").split(":") + ["0"])[:2]
    return int(h) * 60 + int(m[:2] or 0)

def periods_between(sc, a, b):
    """clock times a-b -> (first, last) period that the block-out overlaps"""
    st, pl, NP = _minutes(sc["start_time"]), int(sc["period_minutes"]), int(sc["periods_per_day"])
    lo, hi = (_minutes(a) - st) // pl, -(-(_minutes(b) - st) // pl) - 1
    return max(0, lo), min(NP - 1, hi)

def _yes(v): return str(v or "").strip().lower() in ("yes", "y", "true", "1")

def _names(v):
    """a cell listing names ('Chemistry G3, SC, Sec 3') -> ['Chemistry G3', 'SC', 'Sec 3']"""
    return [s.strip() for s in re.split(r"[,;\n]", str(v or "")) if s.strip()]

def _week(v):
    """'Special room weeks' on a lesson row: Odd or Even (the special room on that week only, a classroom on the other)"""
    v = str(v or "").strip().lower()
    return "Odd" if v.startswith("odd") else "Even" if v.startswith("even") else None

ROOM_NOTE = re.compile(r"((?:[A-Z][\w&/-]*(?: [A-Z][\w&/-]*)*)(?:(?:, | and )[A-Z][\w&/-]*(?: [A-Z][\w&/-]*)*)*) takes? (\d+) classes")

def _room_rules_from_notes(rows, venues, rules):
    """a note on a Settings row such as 'Field and Hall take 2 classes each' (or 'Hall takes 3 classes'): those rooms hold
    that many classes at once, as an 'At most N at once' rule for the room (unless the workbook already has one)"""
    names = {v["room"] for v in venues}; have = {(r.get("scope"), r.get("target")) for r in rules if r.get("rule") == "At most N at once"}
    out = []
    for k, v, note in rows:
        for m in ROOM_NOTE.finditer(str(note or "")):
            for room in re.split(r", | and ", m.group(1)):
                if room in names and ("Room", room) not in have:
                    have.add(("Room", room))
                    out.append(dict(scope="Room", target=room, rule="At most N at once", number=int(m.group(2)), strength="Must", week="Both", day="All",
                                    note=f"From the Settings note on '{k}': {note}"))
    return out

def read_template(path):
    wb = openpyxl.load_workbook(path, data_only=False)
    st = wb["Settings"]
    rows = [(str(st.cell(r, 1).value).strip(), st.cell(r, 2).value, st.cell(r, 3).value) for r in range(5, 60) if st.cell(r, 1).value]
    kv = {k: v for k, v, _ in rows}
    kind = str(kv.get("School type") or "secondary").lower()
    sc: dict = defaults("primary" if kind.startswith("pri") else "secondary")
    t0 = kv.get("Period 0 starts at")
    if isinstance(t0, datetime.datetime): t0 = t0.time()
    weeks = kv.get("Weeks in cycle") or (re.match(r"\s*(\d)", str(kv.get("Cycle") or "")) or [None, None])[1]
    sc.update(name=kv.get("School") or sc["name"], timetable=kv.get("Timetable") or sc["timetable"], weeks=int(weeks or sc["weeks"]),
              period_minutes=int(kv.get("Period length (minutes)") or sc["period_minutes"]),
              start_time=t0.strftime("%H:%M") if isinstance(t0, datetime.time) else str(t0 or sc["start_time"]),
              periods_per_day=int(kv.get("Periods in a day") or sc["periods_per_day"]),
              lesson_length=int(kv.get("Lesson length (periods)") or kv.get("Standard lesson length (periods)") or sc["lesson_length"]))
    # rule rows on Settings (once a day, after school, PE at once, daily cap, recess); rows the app does not know are kept as they are
    rules, known = [], {"School", "Timetable", "School type", "Weeks in cycle", "Cycle", "Period length (minutes)", "Period 0 starts at", "Periods in a day",
                        "Standard lesson length (periods)", "Lesson length (periods)", "Lesson lengths"}
    sc["extra_settings"] = []
    for k, v, note in rows:
        what = next((s for s in SETTING_RULES if k.lower().startswith(s["label"].lower()[:30])), None)
        if what is None:
            if k not in known: sc["extra_settings"].append([k, v.strftime("%H:%M") if isinstance(v, datetime.time) else v, note])
            continue
        if what["key"] == "after_school": sc["after_school"] = _yes(v) if v not in (None, "") else True
        elif what["key"] == "pe_rain_home": sc["pe_rain_home"] = _yes(v)
        elif what["number"] and v not in (None, ""): rules.append(dict(what["rule"], number=int(float(str(v)))))
        elif not what["number"] and _yes(v): rules.append(dict(what["rule"]))
    # lists
    ls = wb["Lists"]; heads = {str(ls.cell(3, j).value): j for j in range(1, ls.max_column + 1) if ls.cell(3, j).value}
    def col(name):
        j = heads.get(name); out = []
        if not j: return out
        for r in range(4, ls.max_row + 1):
            v = ls.cell(r, j).value
            if v not in (None, ""): out.append(v)
        return out
    if col("Levels"): sc["levels"] = [str(x) for x in col("Levels")]
    if col("Departments"): sc["departments"] = [str(x) for x in col("Departments")]
    if col("Room kinds"): sc["room_kinds"] = [str(x) for x in col("Room kinds")]
    tags = col("Level tags") or col("G-levels")
    if tags:   # each tag's letter in group codes: the Tag letter column, else the usual one (G1 K, G2 A, G3 P; Std S, Fdn F, Higher H)
        known_letter = dict(sc["tags"])
        letters = [ls.cell(r, heads["Tag letter"]).value for r in range(4, 4 + len(tags))] if "Tag letter" in heads else [None] * len(tags)
        sc["tags"] = [[str(a), str(b if b is not None else known_letter.get(str(a), ""))] for a, b in zip(tags, letters)]
    if col("Level") and col("Band"): sc["bands"] = dict(zip([str(x) for x in col("Level")], [str(x) for x in col("Band")]))
    # day structure
    ds = wb["Day structure"]; NP = sc["periods_per_day"]; day_structure = {}
    for r in range(6, ds.max_row + 1):
        lv, d = ds.cell(r, 1).value, ds.cell(r, 2).value
        if not lv or not d: continue
        codes = "".join((str(ds.cell(r, 3 + p).value or ".").strip()[:1].upper() or ".") for p in range(NP))
        day_structure.setdefault(str(lv), {})[str(d)] = codes
    _, ven = _table(wb["1 Venues"], "Room")
    venues = [dict(room=str(x["Room"]), kind=x.get("Kind"), capacity=x.get("Capacity"), building=x.get("Building") or x.get("Block"), floor=x.get("Floor") or x.get("Floor / level"),
                   note=x.get("Note"), only_for=_names(next((v for h, v in x.items() if h.startswith("For")), None))) for x in ven if x.get("Room")]
    _, cl = _table(wb["2 Classes"], "Class")
    classes = [{"class": str(x["Class"]), "level": x.get("Level"), "size": x.get("Students") or 0, "home_room": x.get("Home room"), "note": x.get("Note")} for x in cl if x.get("Class")]
    sh, su = _table(wb["3 Subjects"], "Subject")
    lvcols = [h for h in sh if h in sc["levels"]]
    lenh = next((h for h in sh if h.startswith("Lesson length")), None); kindh = next((h for h in sh if h.startswith("Room kind")), None)
    subjects = [dict(subject=str(x["Subject"]), code=x.get("Code"), tag=x.get("Level tag") or x.get("G-level") or "-", dept=x.get("Department"), room_kind=x.get(kindh),
                     lesson_length=x.get(lenh), periods={lv: x.get(lv) for lv in lvcols if x.get(lv)}, note=x.get("Note")) for x in su if x.get("Subject")]
    _, te = _table(wb["4 Teachers"], "Teacher")
    teachers = []
    for x in te:
        if not x.get("Teacher"): continue
        ct = [dict(subject=x.get(f"Can teach {p}"), levels=x.get(f"at {p}") or "All") for p in range(1, 9) if x.get(f"Can teach {p}")]
        teachers.append(dict(name=str(x["Teacher"]), short=x.get("Short"), dept=x.get("Department"), allowance=x.get("Allowance (periods)"),
                             provisional=x.get("Provisional"), can_teach=ct, role=x.get("Role"), other_duties=x.get("Other duties"), note=x.get("Note")))
    _, bl = _table(wb["5 Blocks"], "Block")
    blocks = [dict(block=str(x["Block"]), level=x.get("Level"), classes=[str(c) for c in _first_col(x, "Class", 12)], note=x.get("Note")) for x in bl if x.get("Block")]
    lh, le = _table(wb["6 Lessons"], "Level")
    bh = next((h for h in lh if h.startswith("Block")), "Block")
    lessons = [dict(level=x.get("Level"), subject=x.get("Subject"), block=x.get(bh), classes=[str(c) for c in _first_col(x, "Class", 12)],
                    teachers=[str(t) for t in _first_col(x, "Teacher", 3)], students=x.get("Students"),
                    periods_override=x.get("Periods (only if different)"), code_override=x.get("Code (only to keep an old one)"), note=x.get("Note"),
                    room_weeks=_week(x.get("Special room weeks")))
               for x in le if x.get("Subject")]
    for name in ("7 Rules", "More rules"):
        if name not in wb.sheetnames: continue
        _, ru = _table(wb[name], "Applies to (kind)")
        for x in ru:
            if not x.get("Applies to (kind)") or not x.get("Rule"): continue
            rules.append({"scope": x.get("Applies to (kind)"), "target": x.get("Applies to"), "subject": x.get("Subject (optional)"), "rule": x.get("Rule"),
                          "week": x.get("Week") or "Both", "day": x.get("Day") or "All", "from": x.get("From period"), "to": x.get("To period"),
                          "number": x.get("Number"), "strength": x.get("Must / Prefer") or "Must", "note": x.get("Note")})
    if "Teacher availability" in wb.sheetnames:   # clock times -> the periods they overlap; the times are kept for writing back
        for x in _table(wb["Teacher availability"], "Teacher")[1]:
            if not x.get("Teacher") or x.get("From") in (None, "") or x.get("To") in (None, ""): continue
            a, b = periods_between(sc, x["From"], x["To"])
            rules.append(dict(scope="Teacher", target=str(x["Teacher"]), rule="Not available", week=x.get("Week") or "Both", day=x.get("Day") or "All",
                              **{"from": a, "to": b}, strength="Must", note=x.get("Reason"), times=[_clock(x["From"]), _clock(x["To"])]))
    class_names = {c["class"] for c in classes}
    if "Fixed lessons" in wb.sheetnames:
        for x in _table(wb["Fixed lessons"], "Level")[1]:
            who, start = x.get("Class(es) or group"), x.get("Start period")
            if not who or start in (None, ""): continue
            for w in [s.strip() for s in str(who).split(",") if s.strip()]:
                rules.append(dict(scope="Class" if w in class_names else "Group", target=w, subject=x.get("Subject"), rule="Fixed at", level=x.get("Level"),
                                  week=x.get("Week") or "Both", day=x.get("Day") or "All", **{"from": int(start), "to": None}, strength="Must", note=x.get("Note")))
    students = []
    if "Students" in wb.sheetnames:
        sh_, sx = _table(wb["Students"], "Student ID")
        for x in sx:
            if not (x.get("Student ID") or x.get("Name")): continue
            g = x.get("Groups")
            students.append(dict(id=x.get("Student ID"), name=x.get("Name"), level=x.get("Level"), **{"class": x.get("Class")}, combination=x.get("Combination"),
                                 choices=[x[f"Choice {i}"] for i in (1, 2, 3) if x.get(f"Choice {i}")], merit=x.get("Merit"),
                                 groups=[s.strip() for s in str(g).split(",") if s.strip()] if g else [], note=x.get("Note") or x.get("Source")))
    combinations = []
    if "Combinations" in wb.sheetnames:
        for x in _table(wb["Combinations"], "Level")[1]:
            subs = [x[k] for k in ("Subject 1", "Subject 2", "Subject 3", "Subject 4") if x.get(k)] or \
                   [s.strip() for s in str(x.get("Subjects") or "").split("+") if s.strip()]
            if x.get("Code") or subs: combinations.append(dict(level=x.get("Level"), code=x.get("Code"), name=x.get("Name"), subjects=subs, places=x.get("Places")))
    roles = []
    if "Roles" in wb.sheetnames:
        rs = wb["Roles"]
        for x in _table(rs, "Role")[1]:
            if x.get("Role"): roles.append(dict(role=str(x["Role"]), what=x.get("What it is"), allowance=x.get("Guide allowance"), group=x.get("Group")))
        spread = next((rs.cell(r, 3).value for r in range(1, 8) if str(rs.cell(r, 1).value or "").startswith("Even spread")), None)
        if isinstance(spread, (int, float)): sc["role_spread"] = int(spread)
    rules += _room_rules_from_notes(rows, venues, rules)
    # sheets the app does not use (e.g. a year's notes) are carried through as text
    extra = []
    for ws in wb.worksheets:
        if ws.title in KNOWN_SHEETS: continue
        rows_ = [[("" if v is None else v.strftime("%H:%M") if isinstance(v, datetime.time) else v) for v in row] for row in ws.iter_rows(values_only=True)]
        while rows_ and not any(v not in ("", None) for v in rows_[-1]): rows_.pop()
        extra.append(dict(name=ws.title, rows=[[v if isinstance(v, (int, float, str)) else str(v) for v in row] for row in rows_]))
    return dict(school=sc, day_structure=day_structure, venues=venues, classes=classes, subjects=subjects, teachers=teachers,
                blocks=blocks, lessons=lessons, rules=rules, students=students, combinations=combinations, roles=roles, extra_sheets=extra)

KNOWN_SHEETS = {"README", "Settings", "Period times", "Day structure", "1 Venues", "2 Classes", "3 Subjects", "4 Teachers", "Roles", "5 Blocks", "6 Lessons",
                "7 Rules", "More rules", "Rule types", "Teacher availability", "Fixed lessons", "Students", "Combinations", "Lists"}

def _clock(t):
    m = _minutes(t); return f"{m // 60:02d}:{m % 60:02d}"

def derive(spec):
    """add code, periods, lesson pattern, room kind to every lesson (same logic as the workbook formulas); returns problems found"""
    sc = spec["school"]; subj = {s["subject"]: s for s in spec["subjects"]}; letter = {a: b for a, b in sc["tags"]}
    seq = {}; problems = []
    for i, x in enumerate(spec["lessons"]):
        s = subj.get(x["subject"])
        if s is None: problems.append(f"Lesson row {i + 6}: subject '{x['subject']}' is not on Subjects."); continue
        p = x.get("periods_override") or (s.get("periods") or {}).get(x["level"])
        L = int(s.get("lesson_length") or sc["lesson_length"])
        full, short, L = lesson_split(p, L)
        if x.get("block"):
            k = (x["level"], x["subject"]); seq[k] = seq.get(k, 0) + 1
        else:
            seq.setdefault((x["level"], x["subject"]), seq.get((x["level"], x["subject"]), 0))
        if x.get("code_override"): code = str(x["code_override"])
        elif not x.get("block"): code = "Entire Class"
        else: code = f"{str(x['level'])[-1]}{s.get('code') or '??'}{letter.get(s.get('tag'), '')}{seq[(x['level'], x['subject'])]}"
        x.update(code=code, periods=int(p or 0), full=full, short=short, length=L, room_kind=s.get("room_kind") or ("classroom" if x.get("block") or x.get("code_override") else None),
                 tag=s.get("tag") or "-", dept=s.get("dept"))
        if x.get("pattern") and sum(int(v) for v in x["pattern"]) != x["periods"]: x.pop("pattern")   # lengths set by hand no longer add up: back to the usual split
        if not p: problems.append(f"Lesson row {i + 6}: {x['subject']} has no periods for {x['level']}.")
        if not x["classes"]: problems.append(f"Lesson row {i + 6}: no class.")
        if not x["teachers"]: problems.append(f"Lesson row {i + 6}: no teacher ({x['subject']} {', '.join(x['classes'])}).")
    return problems
