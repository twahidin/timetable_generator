"""Write a guided timetable setup workbook from a spec (see spec.py), in the setup template layout (v9 and later):
Settings (with the rule rows), Period times, Day structure, steps 1-6, Roles, Teacher availability, Fixed lessons,
More rules (anything else), Students, Combinations and Lists. Drop-downs flow from step to step; group codes, lesson
patterns, loads and checks are Excel formulas."""
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.workbook.defined_name import DefinedName
from .spec import RULE_TYPES, SCOPES, CODES, RECESS_ROWS, SETTING_RULES, ROLES, day_names, order_days, scopes_for, split_rules

NAVY, SLATE = "1F4E78", "5B6B7F"
HF = Font(name="Arial", bold=True, color="FFFFFF", size=10)
INF = Font(name="Arial", size=10, color="0000FF"); INFILL = PatternFill("solid", fgColor="FFF9DB")
BF = Font(name="Arial", size=10); CF = PatternFill("solid", fgColor="EDEDED")
TITLE = Font(name="Arial", size=14, bold=True, color=NAVY); NOTE = Font(name="Arial", size=9, italic=True, color="595959")
EXF = Font(name="Arial", size=10, italic=True, color="7F7F7F"); RED = PatternFill("solid", fgColor="F4B6B6")
HR, R0 = 5, 6          # header row and first data row of steps 1-6
NTCH, NPAIR = 3, 8
BOX = Border(*(Side(style="thin", color="C9D3E0"),) * 4)
LENGTH_NAME = {1: "Singles", 2: "Doubles", 3: "Triples", 4: "Quadruples"}

class Book:
    def __init__(self):
        self.wb = openpyxl.Workbook(); self.wb.remove(self.wb.active)
    def sheet(self, name, title, notes, cols, inputs, hr=HR):
        ws = self.wb.create_sheet(name)
        ws["A1"] = title; ws["A1"].font = TITLE
        for i, n in enumerate(notes[:hr - 2], start=2): ws.cell(i, 1, n).font = NOTE
        for j, (h, w) in enumerate(cols, start=1):
            c = ws.cell(hr, j, h); c.font = HF; c.alignment = Alignment(wrap_text=True, vertical="center")
            c.fill = PatternFill("solid", fgColor=NAVY if j in inputs else SLATE)
            ws.column_dimensions[CL(j)].width = w
        ws.row_dimensions[hr].height = 30; ws.freeze_panes = ws.cell(hr + 1, 1)
        return ws
    def paint(self, ws, last, inputs, ncols, first=R0):
        for r in range(first, last + 1):
            for j in range(1, ncols + 1):
                c = ws.cell(r, j)
                if j in inputs: c.font = INF; c.fill = INFILL
                else: c.font = BF; c.fill = CF
    def dv(self, ws, rng, src, msg):
        d = DataValidation(type="list", formula1=src, allow_blank=True, showErrorMessage=True, errorTitle="Pick from the list", error=msg)
        ws.add_data_validation(d)
        for part in rng.split(): d.add(part)
    def name(self, nm, ref): self.wb.defined_names[nm] = DefinedName(nm, attr_text=ref)
    def dyn(self, sheet, col, first=R0, n=1000):
        return f"OFFSET('{sheet}'!${col}${first},0,0,MAX(1,COUNTA('{sheet}'!${col}${first}:${col}${first + n - 1})),1)"
    def head(self, ws, row, names, fill=NAVY):
        for j, h in enumerate(names, start=1):
            c = ws.cell(row, j, h); c.font = HF; c.fill = PatternFill("solid", fgColor=fill)

def _t(s):
    if isinstance(s, datetime.time): return s
    h, m = str(s or "07:30").split(":")[:2]; return datetime.time(int(h), int(m))

def _clock(sc, p):
    """start time of period p (p = periods_per_day: end of the day) as a time"""
    h, m = map(int, str(sc["start_time"]).split(":")[:2]); t = h * 60 + m + p * int(sc["period_minutes"])
    return datetime.time(min(23, t // 60), t % 60)

def write_template(spec, path, extra_rows=120):
    sc = spec["school"]; levels = sc["levels"]; NP = int(sc["periods_per_day"]); L_std = int(sc["lesson_length"])
    days = order_days(list(next((d for d in spec.get("day_structure", {}).values() if d), {}))) or day_names(sc["weeks"])
    secondary = sc.get("type", "secondary") == "secondary"
    TAGH, TAGS = ("G-level", "G-levels") if secondary else ("Level tag", "Level tags")
    B = Book(); wb = B.wb
    V, C, S, T, BL, L = "1 Venues", "2 Classes", "3 Subjects", "4 Teachers", "5 Blocks", "6 Lessons"
    TA, FX, MR = "Teacher availability", "Fixed lessons", "More rules"
    rules = split_rules(spec)
    NCLS = max([9] + [len(x.get("classes") or []) for x in spec["blocks"] + spec["lessons"]])
    nV, nC, nS = len(spec["venues"]), len(spec["classes"]), len(spec["subjects"])
    nT, nB, nL = len(spec["teachers"]), len(spec["blocks"]), len(spec["lessons"])
    LAST = lambda n: R0 + n + extra_rows
    year = str(sc["timetable"])[:4] if str(sc["timetable"])[:4].isdigit() else ""
    fortnight = "fortnight" if int(sc["weeks"]) == 2 else "week"

    # ---------------- README
    rd = wb.create_sheet("README")
    steps = [("0", "Settings, Day structure", "Start time, period length, rules; what each period is for each level and day."),
             ("1", V, "List every room and pick its kind."),
             ("2", C, "List the classes; pick each one's level and home room (from Step 1)."),
             ("3", S, f"List the subjects with code, {TAGH}, department, room kind (Step 1 kinds) and periods a {fortnight} per level. "
                      f"{LENGTH_NAME.get(L_std - 1, 'Shorter lessons')} and {LENGTH_NAME.get(L_std, 'full lessons').lower()} follow from the periods."),
             ("4", T, "List the teachers; pick the subjects (Step 3) and levels each can teach. Periods given and Spare update as you fill Step 6. "
                      "Pick each teacher's role (HOD, SH, EO...); the Roles sheet shows the load by role."),
             ("5", BL, "Name each block of groups that run together and pick its classes (Step 2). Its groups, student totals and longest periods fill in from Step 6."),
             ("6", L, "One row per group (or per class for whole-class subjects): pick level, subject (Step 3), block (Step 5), classes (Step 2) and teachers (Step 4). "
                      "The group code, periods, lesson pattern, room kind and checks fill in."),
             ("7", f"{TA}, {FX}, {MR}, Students", "Optional: block-out times, lessons at set times, any other rule, and each student's class, combination and groups.")]
    rd["A1"] = f"{sc['name']} {sc['timetable']} timetable setup"; rd["A1"].font = TITLE
    rd["A2"] = "Fill the sheets in order. Each step's list becomes the drop-down for the next, so nothing is typed twice and nothing can be misspelt."; rd["A2"].font = BF
    rd["A3"] = "Blue text on yellow = type or pick.  Grey = worked out by formulas.  Red in a Check column = something to fix."; rd["A3"].font = BF
    B.head(rd, 5, ["Step", "Sheet", "What you do"])
    for i, (a, b, c_) in enumerate(steps, start=6):
        rd.cell(i, 1, a).font = Font(name="Arial", bold=True, size=10); rd.cell(i, 2, b).font = Font(name="Arial", bold=True, size=10)
        x = rd.cell(i, 3, c_); x.font = BF; x.alignment = Alignment(wrap_text=True, vertical="top")
    r = 6 + len(steps) + 1
    letters = ", ".join(f"{a} {b}" for a, b in sc["tags"] if b)
    for t_ in ["How group codes are made", f"Level digit + subject code (Step 3) + level letter ({letters}) + a running number, e.g. 1ELP1, 1ELP2, 1ELA1. "
               "To keep an existing school code (such as 1ELR), type it in 'Code (only to keep an old one)'.", "",
               "Lists", f"The Lists sheet holds the fixed choices (levels, departments, room kinds, {TAGS}, 'teaches at'). Add an item at the bottom of a column and it appears in the drop-downs."]:
        c = rd.cell(r, 1, t_ or None); c.font = Font(name="Arial", size=11, bold=True, color=NAVY) if len(t_) < 40 else BF
        rd.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3); c.alignment = Alignment(wrap_text=True, vertical="top"); r += 1
    rd.column_dimensions["A"].width = 7; rd.column_dimensions["B"].width = 26; rd.column_dimensions["C"].width = 110

    # ---------------- sheets carried through from the workbook this setup came from (e.g. a year's notes)
    extra = [x["name"] for x in spec.get("extra_sheets", []) if x.get("name") and x["name"] not in ("README", "Lists")]
    for x in spec.get("extra_sheets", []):
        if x.get("name") not in extra: continue
        es = wb.create_sheet(x["name"])
        for i, row in enumerate(x.get("rows") or [], start=1):
            for j, v in enumerate(row, start=1):
                if v not in ("", None): c = es.cell(i, j, v); c.font = TITLE if (i == 1 and j == 1) else BF; c.alignment = Alignment(wrap_text=True, vertical="top")
        es.column_dimensions["A"].width = 110

    # ---------------- Settings (with the rule rows) / Period times
    st = wb.create_sheet("Settings")
    st["A1"] = "Timetable settings"; st["A1"].font = TITLE
    st["A2"] = "Blue text on yellow = cells to edit. Everything else in the workbook reads from here."; st["A2"].font = NOTE
    B.head(st, 4, ["Setting", "Value", "Notes"])
    cycle = "2 weeks (Odd, Even)" if int(sc["weeks"]) == 2 else "1 week"
    rows = [("School", sc["name"], ""), ("Timetable", sc["timetable"], "")] + ([] if secondary else [("School type", sc["type"], "primary or secondary")]) + \
           [("Cycle", cycle, "Each week is Monday to Friday; the Day structure sheet has a row for each day." if int(sc["weeks"]) == 2 else "Every week is the same."),
            ("Period length (minutes)", int(sc["period_minutes"]), ""), ("Period 0 starts at", _t(sc["start_time"]), ""),
            ("Periods in a day", NP, f"Numbered 0-{NP - 1}."),
            ("Lesson length (periods)", L_std, f"A subject's leftover periods become one or two lessons of {L_std - 1} period{'s' if L_std != 2 else ''}.")]
    rs = rules["settings"]
    vals = dict(once_a_day="Yes" if "once_a_day" in rs else "No", after_school="Yes" if sc.get("after_school", True) else "No",
                pe_at_once=rs.get("pe_at_once"), max_day=rs.get("max_day"), recess="Yes" if "recess" in rs else "No",
                pe_rain_home="Yes" if sc.get("pe_rain_home") else "No")
    notes = dict(once_a_day="Yes / No", after_school="Yes: the 'A' periods on Day structure are used only when a class cannot fit. No: never.",
                 pe_at_once="Leave empty for no limit (e.g. the field and hall take 2 classes each).", max_day="Leave empty for no limit. 15 (5 hours) is a common cap.",
                 recess="Yes: every teacher is free for a recess each day (a whole recess, or the recess length inside a recess window). The solver keeps it.",
                 pe_rain_home="Yes: while a class has PE (a 'pe venue' lesson) nothing else uses its home room, so the class can go there when it rains. The timetables show it as the rain room.")
    rows += [(s["label"], vals[s["key"]], notes[s["key"]]) for s in SETTING_RULES]
    clock_txt = lambda v: _t(v) if isinstance(v, str) and len(v) <= 5 and v.count(":") == 1 and v.replace(":", "").isdigit() else v   # '15:20' back to a time
    more_rows = list(sc.get("extra_settings", []))
    if any("W" in str(row) for d in spec.get("day_structure", {}).values() for row in d.values()):   # recess windows: their Settings rows
        have = [str(x[0]).strip().lower() for x in more_rows]
        more_rows += [[k, v, n] for k, v, n in RECESS_ROWS if not any(h.startswith(k.lower()) for h in have)]
    rows += [(x[0], clock_txt(x[1] if len(x) > 1 else None), x[2] if len(x) > 2 else "") for x in more_rows]
    at = {}
    for i, (k, v, n) in enumerate(rows, start=5):
        at[k] = i
        st.cell(i, 1, k).font = Font(name="Arial", bold=True, size=10); c = st.cell(i, 2, v); c.font = INF; c.fill = INFILL
        st.cell(i, 3, n).font = NOTE
        if isinstance(v, datetime.time): c.number_format = "h:mm AM/PM"
    yn = " ".join(f"B{at[s['label']]}" for s in SETTING_RULES if not s["number"])
    B.dv(st, yn, '"Yes,No"', "Yes or No.")
    for k, r in at.items():
        if str(k).lower().startswith("students have no free periods"): B.dv(st, f"B{r}", '"Yes,Prefer,No"', "Yes, Prefer or No.")
        if str(k).lower().startswith("floating recess goes by"): B.dv(st, f"B{r}", '"Level,Class"', "Level or Class.")
    st.column_dimensions["A"].width = 50; st.column_dimensions["B"].width = 28; st.column_dimensions["C"].width = 70
    SLEN, SSTART, SLL = f"Settings!$B${at['Period length (minutes)']}", f"Settings!$B${at['Period 0 starts at']}", f"Settings!$B${at['Lesson length (periods)']}"
    pt = wb.create_sheet("Period times")
    pt["A1"] = "Period times"; pt["A1"].font = TITLE; pt["A2"] = "Worked out from Settings (period 0 start, period length). Do not type here."; pt["A2"].font = NOTE
    B.head(pt, 4, ["Period", "Starts", "Ends"], SLATE)
    for p in range(NP):
        rr = 5 + p; pt.cell(rr, 1, p).font = BF
        pt.cell(rr, 2, f"={SSTART}+A{rr}*{SLEN}/1440").number_format = "h:mm AM/PM"
        pt.cell(rr, 3, f"=B{rr}+{SLEN}/1440").number_format = "h:mm AM/PM"
    PTB, PTC = f"'Period times'!$B$5:$B${4 + NP}", f"'Period times'!$C$5:$C${4 + NP}"

    # ---------------- Day structure
    ds = wb.create_sheet("Day structure")
    ds["A1"] = "Day structure: what each period is, per level and day"; ds["A1"].font = TITLE
    ds["A2"] = "Codes: " + "   ".join(f"{k} = {v.lower()}" for k, v in CODES.items() if k != ".") + "   blank = no school"; ds["A2"].font = NOTE
    for j, h in enumerate(["Level", "Day"] + list(range(NP)), start=1):
        c = ds.cell(4, j, h); c.font = HF; c.fill = PatternFill("solid", fgColor=NAVY); c.alignment = Alignment(horizontal="center")
    ds.cell(5, 2, "starts").font = NOTE
    for p in range(NP):
        c = ds.cell(5, 3 + p, f"=TEXT('Period times'!B{5 + p},\"h:mm\")"); c.font = Font(name="Arial", size=7, color="595959")
        c.alignment = Alignment(horizontal="center", textRotation=90)
    ds.row_dimensions[5].height = 30
    rr = 6
    for lv in [lv for lv in levels if lv in spec["day_structure"]] or levels:   # a level with no day structure (e.g. Sec 5) gets no rows
        for d in days:
            ds.cell(rr, 1, lv).font = Font(name="Arial", bold=True, size=10); ds.cell(rr, 2, d).font = Font(name="Arial", bold=True, size=10)
            code = (spec["day_structure"].get(lv, {}).get(d, "") or "").ljust(NP, ".")
            for p in range(NP):
                v = code[p] if code[p] != "." else None
                c = ds.cell(rr, 3 + p, v); c.font = INF; c.alignment = Alignment(horizontal="center"); c.border = BOX
            rr += 1
    rng = f"C6:{CL(2 + NP)}{rr - 1}"
    B.dv(ds, rng, '"F,L,R,W,U,A"', "Use F, L, R, W, U, A or leave blank.")
    for code, col in (("F", "C9C9C9"), ("L", "CFE0F3"), ("R", "F6C28B"), ("W", "FBE2C8"), ("U", "A9D8B0"), ("A", "FFE699")):
        ds.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{code}"'], fill=PatternFill("solid", fgColor=col)))
    ds.column_dimensions["A"].width = 8; ds.column_dimensions["B"].width = 11
    for p in range(NP): ds.column_dimensions[CL(3 + p)].width = 4.6
    ds.freeze_panes = "C6"

    # ---------------- Lists (the first six columns as on v9; the rest feed Fixed lessons and More rules)
    ls = wb.create_sheet("Lists")
    ls["A1"] = "Lists behind the drop-downs (add an item at the bottom of a column to make it pickable)"; ls["A1"].font = TITLE
    lists = [("Levels", levels), ("Departments", sc["departments"]), ("Room kinds", sc["room_kinds"]), (TAGS, [a for a, b in sc["tags"]]),
             ("Teaches at", scopes_for(sc)), ("Yes/No", ["Yes", "No"]), ("Tag letter", [b for a, b in sc["tags"]]), ("Level", list(sc["bands"])),
             ("Band", list(sc["bands"].values())), ("Rule applies to", SCOPES), ("School", ["Whole school"]),
             ("Week", ["Both", "Odd", "Even"] if int(sc["weeks"]) == 2 else ["Both"]), ("Day", ["All"] + ["Mon", "Tue", "Wed", "Thu", "Fri"]),
             ("Must/Prefer", ["Must", "Prefer"]), ("Periods", list(range(NP)))]
    for j, (h, vals_) in enumerate(lists, start=1):
        c = ls.cell(3, j, h); c.font = HF; c.fill = PatternFill("solid", fgColor=NAVY); ls.column_dimensions[CL(j)].width = 15
        for i, v in enumerate(vals_, start=4): ls.cell(i, j, v).font = INF; ls.cell(i, j).fill = INFILL
    for nm, col in (("LevelList", "A"), ("DeptList", "B"), ("KindList", "C"), ("GList", "D"), ("ScopeList", "E"), ("YesNoList", "F"),
                    ("RuleScopeList", "J"), ("SchoolList", "K"), ("WeekList", "L"), ("DayList", "M"), ("StrengthList", "N"), ("PeriodList", "O")):
        B.name(nm, B.dyn("Lists", col, 4, 1000))
    letter_of = lambda tag: f'IFERROR(INDEX(Lists!$G$4:$G$60,MATCH({tag},Lists!$D$4:$D$60,0)),"")'
    band_of = lambda lv: f'IFERROR(INDEX(Lists!$I$4:$I$60,MATCH({lv},Lists!$H$4:$H$60,0)),"")'

    # ---------------- 1 Venues
    VIN = {1, 2, 3, 4, 5, 7, 8}
    vw = B.sheet(V, "Step 1 · Venues: every room", ["One row per room. Kind decides which subjects may use it (see Subjects). Capacity is seats: a group only goes in a room that seats it.",
                  "For (optional): the subjects, departments or levels the room is kept for, separated by commas; blank = anyone. Blue on yellow = type or pick. Grey = worked out for you."],
                 [("Room", 16), ("Kind", 16), ("Capacity", 10), ("Block", 12), ("Floor / level", 9), ("Used by lessons", 12), ("Note", 34), ("For (subjects, departments or levels)", 30)], VIN)
    for i, x in enumerate(spec["venues"]):
        for j, k in enumerate(["room", "kind", "capacity", "building", "floor", None, "note"], start=1):
            if k: vw.cell(R0 + i, j, x.get(k))
        if x.get("only_for"): vw.cell(R0 + i, 8, ", ".join(x["only_for"]))
    B.paint(vw, LAST(nV), VIN, 8)
    for r in range(R0, LAST(nV) + 1): vw.cell(r, 6, f"=IF(A{r}=\"\",\"\",COUNTIF('{C}'!$D${R0}:$D${R0 + 999},A{r}))")
    B.dv(vw, f"B{R0}:B{LAST(nV)}", "=KindList", "Pick a room kind (add kinds on the Lists sheet).")
    B.name("RoomList", B.dyn(V, "A"))

    # ---------------- 2 Classes
    cw = B.sheet(C, "Step 2 · Classes", ["One row per form class. Home room is picked from the Venues list.", "Blue on yellow = type or pick. Grey = worked out for you."],
                 [("Class", 10), ("Level", 9), ("Students", 10), ("Home room", 12), ("Lessons listed", 13), ("Note", 34)], {1, 2, 3, 4, 6})
    for i, x in enumerate(spec["classes"]):
        for j, k in enumerate(["class", "level", "size", "home_room", None, "note"], start=1):
            if k: cw.cell(R0 + i, j, x.get(k))
    B.paint(cw, LAST(nC), {1, 2, 3, 4, 6}, 6)
    lcls = f"'{L}'!$D${R0}:${CL(3 + NCLS)}${R0 + 1999}"
    for r in range(R0, LAST(nC) + 1): cw.cell(r, 5, f'=IF(A{r}="","",SUMPRODUCT(({lcls}=A{r})*1))')
    B.dv(cw, f"B{R0}:B{LAST(nC)}", "=LevelList", "Pick a level."); B.dv(cw, f"D{R0}:D{LAST(nC)}", "=RoomList", "Pick a room from Step 1.")
    B.name("ClassList", B.dyn(C, "A"))

    # ---------------- 3 Subjects (a Lesson length column only when some subject differs from the standard)
    own_len = any(x.get("lesson_length") not in (None, "", L_std) for x in spec["subjects"])
    scol = [("Subject", 36), ("Code", 7), (TAGH, 8), ("Department", 11), ("Room kind", 14)] + ([("Lesson length (periods)", 10)] if own_len else []) + \
           [(lv, 7) for lv in levels] + [("Lesson rows", 10), ("Note", 30)]
    nlv = len(levels); FL = 6 + own_len; LL = FL + nlv - 1
    SIN = set(range(1, LL + 1)) | {LL + 2}
    sw = B.sheet(S, f"Step 3 · Subjects: periods a {fortnight} at each level", [f"One row per subject. Give the periods a {fortnight} for each level that takes it (blank = not offered).",
                 f"Lessons are {L_std}-period lessons; a total that is not a multiple of {L_std} gets one or two shorter lessons." +
                 (" A subject's own Lesson length overrides this." if own_len else ""), "Room kind: leave blank for the class's home room."], scol, SIN)
    for i, x in enumerate(spec["subjects"]):
        rr = R0 + i
        for j, k in enumerate(["subject", "code", "tag", "dept", "room_kind"] + (["lesson_length"] if own_len else []), start=1): sw.cell(rr, j, x.get(k))
        for k, lv in enumerate(levels): sw.cell(rr, FL + k, (x.get("periods") or {}).get(lv))
        sw.cell(rr, LL + 2, x.get("note"))
    B.paint(sw, LAST(nS), SIN, LL + 2)
    for r in range(R0, LAST(nS) + 1): sw.cell(r, LL + 1, f"=IF(A{r}=\"\",\"\",COUNTIF('{L}'!$B${R0}:$B${R0 + 1999},A{r}))")
    B.dv(sw, f"C{R0}:C{LAST(nS)}", "=GList", f"Pick a {TAGH}."); B.dv(sw, f"D{R0}:D{LAST(nS)}", "=DeptList", "Pick a department.")
    B.dv(sw, f"E{R0}:E{LAST(nS)}", "=KindList", "Pick a room kind.")
    B.name("SubjectList", B.dyn(S, "A"))
    SA = f"'{S}'!$A${R0}:$A${R0 + 999}"
    SLV = f"'{S}'!${CL(FL)}${R0}:${CL(LL)}${R0 + 999}"; SLH = f"'{S}'!${CL(FL)}${HR}:${CL(LL)}${HR}"

    # ---------------- 6 Lessons columns (needed by Teachers and Blocks)
    short_h, full_h = (LENGTH_NAME.get(L_std - 1, "Short lessons"), LENGTH_NAME.get(L_std, "Full lessons")) if not own_len else ("Short lessons", "Full lessons")
    names = ["Level", "Subject", "Block"] + [f"Class {i}" for i in range(1, NCLS + 1)] + [f"Teacher {i}" for i in range(1, NTCH + 1)] + \
            ["Students", "Periods (only if different)", "Code (only to keep an old one)", "Group code", "Periods"] + (["Lesson length"] if own_len else []) + \
            [short_h, full_h, "Room kind", "Teacher periods", "Check", "Note", "Special room weeks"]
    COL = {n: CL(j) for j, n in enumerate(names, start=1)}
    LQ = f"'{L}'"; TC = [COL[f"Teacher {k}"] for k in range(1, NTCH + 1)]; TPC = COL["Teacher periods"]

    # ---------------- 4 Teachers (load, role and load check as on v9)
    tcols = [("Teacher", 28), ("Short", 12), ("Department", 11), ("Allowance (periods)", 11), ("Provisional", 10)]
    for p in range(1, NPAIR + 1): tcols += [(f"Can teach {p}", 28), (f"at {p}", 9)]
    LO = 6 + 2 * NPAIR   # Periods given
    tcols += [("Periods given", 10), ("Spare", 8), ("(key, used by checks)", 12), ("Note", 30), ("Role", 9), ("Other duties (periods)", 10),
              ("Total load", 9), ("Role guide (periods)", 9), ("% of allowance", 9), ("Load check", 44)]
    PG, SP, KY, NO, RO, OD, TL, RG, PC, LC = (CL(LO + k) for k in range(10))
    TIN = set(range(1, LO)) | {LO + 3, LO + 4, LO + 5}
    tw = B.sheet(T, "Step 4 · Teachers: who they are and what they can teach", ["Pick each subject a teacher can teach and the levels ('All', a level, or a band such as Sec 1-2). Lessons are checked against this.",
                 "Periods given adds up every lesson the teacher has on Step 6; Spare = Allowance - Periods given (red when over).",
                 f"Blue on yellow = type or pick. Grey = worked out for you.  Role (column {RO}) comes from the Roles sheet; Load check flags management teaching as much as teachers, and teachers well below or above an even share."], tcols, TIN)
    for i, x in enumerate(spec["teachers"]):
        rr = R0 + i
        for j, k in enumerate(["name", "short", "dept", "allowance", "provisional"], start=1): tw.cell(rr, j, x.get(k))
        for p, ct in enumerate((x.get("can_teach") or [])[:NPAIR]):
            tw.cell(rr, 6 + 2 * p, ct.get("subject")); tw.cell(rr, 7 + 2 * p, ct.get("levels") or "All")
        tw[f"{NO}{rr}"] = x.get("note"); tw[f"{RO}{rr}"] = x.get("role"); tw[f"{OD}{rr}"] = x.get("other_duties")
    B.paint(tw, LAST(nT), TIN, LO + 9)
    RN = "Roles!$A$9:$A$28"
    for r in range(R0, LAST(nT) + 1):
        load = "+".join(f"SUMIF({LQ}!${c}${R0}:${c}${R0 + 1999},A{r},{LQ}!${TPC}${R0}:${TPC}${R0 + 1999})" for c in TC)
        tw[f"{PG}{r}"] = f'=IF(A{r}="","",{load})'
        tw[f"{SP}{r}"] = f'=IF(OR(A{r}="",D{r}=""),"",D{r}-{PG}{r})'
        key = "&".join(f'"|"&{CL(6 + 2 * p)}{r}&"@"&IF({CL(7 + 2 * p)}{r}="","All",{CL(7 + 2 * p)}{r})' for p in range(NPAIR)) + '&"|"'
        tw[f"{KY}{r}"] = f'=IF(A{r}="","",{key})'; tw[f"{KY}{r}"].font = Font(name="Arial", size=7, color="8C8C8C")
        tw[f"{TL}{r}"] = f'=IF(A{r}="","",N({PG}{r})+N({OD}{r}))'
        tw[f"{RG}{r}"] = f'=IF(OR(A{r}="",{RO}{r}=""),"",IFERROR(INDEX(Roles!$C$9:$C$28,MATCH({RO}{r},{RN},0)),""))'
        tw[f"{PC}{r}"] = f'=IF(OR(A{r}="",N(D{r})=0),"",{TL}{r}/D{r})'; tw[f"{PC}{r}"].number_format = "0%"
        tw[f"{LC}{r}"] = (f'=IF(A{r}="","",TRIM(IF({RO}{r}="","No role. ","")'
                          f'&IF(AND(N({RG}{r})>0,N(D{r})>0,D{r}<>{RG}{r}),"Allowance differs from the "&{RO}{r}&" guide ("&{RG}{r}&"). ","")'
                          f'&IF(AND(N(D{r})>0,{TL}{r}>D{r}),"Over allowance. ","")'
                          f'&IF(AND(IFERROR(INDEX(Roles!$D$9:$D$28,MATCH({RO}{r},{RN},0)),"")="Management",N(Roles!$C$3)>0,{TL}{r}>=Roles!$C$3),"Management teaching as much as the average teacher ("&ROUND(Roles!$C$3,0)&"). ","")'
                          f'&IF(AND({RO}{r}="EO",{TL}{r}>0,N(Roles!$C$3)>0,Roles!$C$3-{TL}{r}>=Roles!$C$5),"Below the teacher average by "&ROUND(Roles!$C$3-{TL}{r},0)&" periods. ","")'
                          f'&IF(AND({RO}{r}="EO",N(Roles!$C$3)>0,{TL}{r}-Roles!$C$3>=Roles!$C$5),"Above the teacher average by "&ROUND({TL}{r}-Roles!$C$3,0)&" periods. ","")'
                          f'&IF(AND({RO}{r}<>"",N(D{r})>0,{TL}{r}=0),"No lessons yet. ","")))')
        tw[f"{LC}{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    tw.conditional_formatting.add(f"{SP}{R0}:{SP}{LAST(nT)}", CellIsRule(operator="lessThan", formula=["0"], fill=RED))
    B.dv(tw, f"C{R0}:C{LAST(nT)}", "=DeptList", "Pick a department."); B.dv(tw, f"E{R0}:E{LAST(nT)}", "=YesNoList", "Yes or No.")
    B.dv(tw, f"{RO}{R0}:{RO}{LAST(nT)}", "=RoleList", "Pick a role from the Roles sheet.")
    for p in range(NPAIR):
        B.dv(tw, f"{CL(6 + 2 * p)}{R0}:{CL(6 + 2 * p)}{LAST(nT)}", "=SubjectList", "Pick a subject from Step 3.")
        B.dv(tw, f"{CL(7 + 2 * p)}{R0}:{CL(7 + 2 * p)}{LAST(nT)}", "=ScopeList", "Pick All, a level or a band.")
    B.name("TeacherList", B.dyn(T, "A"))
    KEYC = KY

    # ---------------- Roles
    ro = wb.create_sheet("Roles")
    ro["A1"] = "Roles: teaching load by role"; ro["A1"].font = TITLE
    ro["A2"] = (f"Pick each teacher's role on {T} (column {RO}). Management (HOD, SH) should teach fewer periods than teachers; teachers (EO) should be spread evenly. "
                f"The guide allowance is what the role usually carries; Allowance on {T} is the figure the checks use."); ro["A2"].font = NOTE
    ro.merge_cells("A2:I2"); ro["A2"].alignment = Alignment(wrap_text=True, vertical="top"); ro.row_dimensions[2].height = 30
    TZ, TAB, TAD = f"'{T}'!${RO}$6:${RO}$1000", f"'{T}'!${TL}$6:${TL}$1000", f"'{T}'!${PC}$6:${PC}$1000"
    summary = [("Teachers (EO): average load (lessons + other duties)", f'=IFERROR(AVERAGEIFS({TAB},{TZ},"EO",{TAB},">0"),"")'),
               ("Management (HOD, SH): average load", f'=IFERROR((SUMIF({TZ},"HOD",{TAB})+SUMIF({TZ},"SH",{TAB}))/(COUNTIFS({TZ},"HOD",{TAB},">0")+COUNTIFS({TZ},"SH",{TAB},">0")),"")'),
               ("Even spread for teachers (EO): flag a teacher this many periods or more from the teacher average", int(sc.get("role_spread") or 15))]
    for i, (k, v) in enumerate(summary, start=3):
        ro.cell(i, 1, k).font = Font(name="Arial", bold=True, size=10); ro.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)
        c = ro.cell(i, 3, v); c.number_format = "0"
        c.font, c.fill = (INF, INFILL) if i == 5 else (BF, CF)
    B.head(ro, 8, ["Role", "What it is", "Guide allowance (periods a cycle)", "Group", "Teachers", "Average load (periods)", "Lowest", "Highest", "Average % of allowance"])
    roles = spec.get("roles") or [dict(role=a, what=b, allowance=c, group=d) for a, b, c, d in ROLES]
    for i, x in enumerate(roles[:20]):
        for j, k in enumerate(["role", "what", "allowance", "group"], start=1): ro.cell(9 + i, j, x.get(k))
    B.paint(ro, 28, {1, 2, 3, 4}, 9, first=9)
    for r in range(9, 29):
        ro[f"E{r}"] = f'=IF(A{r}="","",COUNTIF({TZ},A{r}))'
        ro[f"F{r}"] = f'=IF(OR(A{r}="",E{r}=0),"",IFERROR(AVERAGEIF({TZ},A{r},{TAB}),""))'
        ro[f"G{r}"] = f'=IF(OR(A{r}="",E{r}=0),"",_xlfn.MINIFS({TAB},{TZ},A{r}))'
        ro[f"H{r}"] = f'=IF(OR(A{r}="",E{r}=0),"",_xlfn.MAXIFS({TAB},{TZ},A{r}))'
        ro[f"I{r}"] = f'=IF(OR(A{r}="",E{r}=0),"",IFERROR(AVERAGEIF({TZ},A{r},{TAD}),""))'; ro[f"I{r}"].number_format = "0%"
        for c in "FGH": ro[f"{c}{r}"].number_format = "0"
    for col, w in zip("ABCDEFGHI", (12, 40, 14, 18, 10, 12, 9, 9, 12)): ro.column_dimensions[col].width = w
    B.name("RoleList", RN)

    # ---------------- 5 Blocks
    bcols = [("Block", 40), ("Level", 8)] + [(f"Class {i}", 7) for i in range(1, NCLS + 1)] + \
            [("Groups in this block (from Step 6)", 46), ("Groups", 8), ("Students in groups", 10), ("Students in classes", 10), ("Difference", 10), ("Longest (periods)", 10), ("Note", 30)]
    G0 = 3 + NCLS
    BIN = set(range(1, G0)) | {G0 + 6}
    bw = B.sheet(BL, "Step 5 · Blocks: groups that run at the same time", ["A block takes the students of its classes and splits them into groups that all have lessons at the same time (e.g. Sec 1 English for 1A-1D).",
                 "Name the block and pick its classes here; then on Step 6 give each group the block's name. The groups list and counts fill in by themselves.",
                 "Difference = students in groups - students in classes: 0 means everyone placed once; negative is normal for option blocks."], bcols, BIN)
    for i, x in enumerate(spec["blocks"]):
        rr = R0 + i; bw.cell(rr, 1, x["block"]); bw.cell(rr, 2, x["level"])
        for k, c in enumerate(x["classes"][:NCLS]): bw.cell(rr, 3 + k, c)
        bw.cell(rr, G0 + 6, x.get("note"))
    B.paint(bw, LAST(nB), BIN, G0 + 6)
    lc, lv_, ls_, lp = (f"{LQ}!${COL[n]}${R0}:${COL[n]}${R0 + 1999}" for n in ("Block", "Group code", "Students", "Periods"))
    for r in range(R0, LAST(nB) + 1):
        g = CL(G0)
        bw[f"{g}{r}"] = ArrayFormula(f"{g}{r}", f'=IF(A{r}="","",_xlfn.TEXTJOIN(", ",TRUE,IF({lc}=A{r},{lv_},"")))')
        bw.cell(r, G0 + 1, f'=IF(A{r}="","",COUNTIF({lc},A{r}))')
        bw.cell(r, G0 + 2, f'=IF(A{r}="","",SUMIF({lc},A{r},{ls_}))')
        bw.cell(r, G0 + 3, f"=IF(A{r}=\"\",\"\",SUMPRODUCT(SUMIF('{C}'!$A${R0}:$A${R0 + 999},C{r}:{CL(2 + NCLS)}{r},'{C}'!$C${R0}:$C${R0 + 999})))")
        bw.cell(r, G0 + 4, f'=IF(A{r}="","",{CL(G0 + 2)}{r}-{CL(G0 + 3)}{r})')
        bw.cell(r, G0 + 5, f'=IF(A{r}="","",_xlfn.MAXIFS({lp},{lc},A{r}))')
        for j in range(G0, G0 + 6): bw.cell(r, j).font = BF; bw.cell(r, j).fill = CF
        bw.cell(r, G0).alignment = Alignment(wrap_text=True, vertical="top")
    dcol = CL(G0 + 4)
    bw.conditional_formatting.add(f"{dcol}{R0}:{dcol}{LAST(nB)}", CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor="FCE4CC")))
    bw.conditional_formatting.add(f"{dcol}{R0}:{dcol}{LAST(nB)}", CellIsRule(operator="greaterThan", formula=["0"], fill=RED))
    B.dv(bw, f"B{R0}:B{LAST(nB)}", "=LevelList", "Pick a level."); B.dv(bw, f"C{R0}:{CL(2 + NCLS)}{LAST(nB)}", "=ClassList", "Pick a class from Step 2.")
    B.name("BlockList", B.dyn(BL, "A"))

    # ---------------- 6 Lessons
    widths = {"Level": 8, "Subject": 32, "Block": 32, "Students": 9, "Periods (only if different)": 11, "Code (only to keep an old one)": 11, "Group code": 11,
              "Periods": 8, "Lesson length": 8, short_h: 8, full_h: 8, "Room kind": 13, "Teacher periods": 9, "Check": 40, "Note": 28, "Special room weeks": 11}
    lcols = [(n, widths.get(n, 7 if n.startswith("Class") else 20)) for n in names]
    lcols[2] = ("Block (empty = whole class)", 32)
    LIN = set(range(1, 4 + NCLS + NTCH + 3)) | {len(names) - 1, len(names)}
    lw = B.sheet(L, "Step 6 · Lessons: one row per group (or per class for whole-class subjects)",
                 ["Pick level, subject and block; pick the class(es) and teacher(s). Leave Block empty for a whole-class lesson.",
                  f"Group code, periods, {short_h.lower()}/{full_h.lower()} and room kind fill in from Steps 3-5. Check flags a teacher not set to teach the subject, a teacher twice in a block, missing periods or a repeated code."], lcols, LIN)
    for i, x in enumerate(spec["lessons"]):
        rr = R0 + i
        lw.cell(rr, 1, x["level"]); lw.cell(rr, 2, x["subject"]); lw.cell(rr, 3, x.get("block") or None)
        for k, c in enumerate(x["classes"][:NCLS]): lw.cell(rr, 4 + k, c)
        for k, t in enumerate((x.get("teachers") or [])[:NTCH]): lw.cell(rr, 4 + NCLS + k, t)
        c0 = 4 + NCLS + NTCH
        lw.cell(rr, c0, x.get("students")); lw.cell(rr, c0 + 1, x.get("periods_override")); lw.cell(rr, c0 + 2, x.get("code_override")); lw.cell(rr, len(names) - 1, x.get("note"))
        if x.get("room_weeks"): lw.cell(rr, len(names), x["room_weeks"])
    nLast = LAST(nL) + 100
    B.paint(lw, nLast, LIN, len(names))
    c_ = COL; CF1, CFN = c_["Class 1"], c_[f"Class {NCLS}"]
    def idx(col, r): return f"INDEX('{S}'!${col}${R0}:${col}${R0 + 999},MATCH(B{r},{SA},0))"
    for r in range(R0, nLast + 1):
        A_, B_, C_ = f"A{r}", f"B{r}", f"C{r}"
        PO, CO, GC, P_, SH, FU, RK, TP, CK = (f"{c_[n]}{r}" for n in ("Periods (only if different)", "Code (only to keep an old one)", "Group code", "Periods",
                                                                      short_h, full_h, "Room kind", "Teacher periods", "Check"))
        seq = f'COUNTIFS($A${R0}:A{r},A{r},$B${R0}:B{r},B{r},$C${R0}:C{r},"<>")'
        code = f'RIGHT({A_},1)&IFERROR({idx("B", r)},"??")&{letter_of(idx("C", r))}&{seq}'
        lw[GC] = f'=IF({B_}="","",IF({CO}<>"",{CO},IF({C_}="","Entire Class",{code})))'
        lw[P_] = f'=IF({B_}="","",IF({PO}<>"",{PO},IFERROR(INDEX({SLV},MATCH({B_},{SA},0),MATCH({A_},{SLH},0)),"")))'
        if own_len:
            LN = f"{c_['Lesson length']}{r}"
            lw[LN] = f'=IF({B_}="","",IFERROR(IF({idx("F", r)}="",{SLL},{idx("F", r)}),{SLL}))'
        else: LN = SLL
        lw[SH] = f'=IF(N({P_})=0,"",IF({LN}<=1,0,IF(MOD(-{P_},{LN})*({LN}-1)>{P_},0,MOD(-{P_},{LN}))))'
        lw[FU] = f'=IF(N({P_})=0,"",({P_}-{SH}*({LN}-1))/{LN})'
        lw[RK] = f'=IF({B_}="","",IFERROR(IF({idx("E", r)}="","home room",{idx("E", r)}),""))'
        lw[TP] = f'=IF(N({P_})=0,0,{P_}*IF({C_}="",MAX(1,COUNTA({CF1}{r}:{CFN}{r})),1))'
        def can(tc):
            k = f"IFERROR(INDEX('{T}'!${KEYC}${R0}:${KEYC}${R0 + 999},MATCH({tc}{r},'{T}'!$A${R0}:$A${R0 + 999},0)),\"\")"
            ok = f'ISNUMBER(FIND("|"&{B_}&"@All|",{k}))+ISNUMBER(FIND("|"&{B_}&"@"&{A_}&"|",{k}))+ISNUMBER(FIND("|"&{B_}&"@"&{band_of(A_)}&"|",{k}))'
            twice = "+".join(f"COUNTIFS($C${R0}:$C${R0 + 1999},{C_},${c}${R0}:${c}${R0 + 1999},{tc}{r})" for c in TC)
            return f'IF({tc}{r}="","",IF({ok}=0,{tc}{r}&" not set to teach this at "&{A_}&". ","")&IF(AND({C_}<>"",{twice}>1),{tc}{r}&" twice in this block. ",""))'
        lw[CK] = (f'=IF({B_}="","",TRIM(IF(N({P_})=0,"No periods for "&{A_}&" on Subjects. ","")&IF(COUNTA({CF1}{r}:{CFN}{r})=0,"No class. ","")'
                  f'&IF(COUNTA({TC[0]}{r}:{TC[-1]}{r})=0,"No teacher. ","")&' + "&".join(can(tc) for tc in TC) +
                  f'&IF(AND({GC}<>"Entire Class",COUNTIF(${c_["Group code"]}${R0}:${c_["Group code"]}${R0 + 1999},{GC})>1),"Code used twice. ","")'
                  f"&IF(AND({C_}<>\"\",COUNTIF('{BL}'!$A${R0}:$A${R0 + 999},{C_})=0),\"Block not on Step 5. \",\"\")))")
        lw[CK].alignment = Alignment(wrap_text=True, vertical="top")
    ck = c_["Check"]
    lw.conditional_formatting.add(f"{ck}{R0}:{ck}{nLast}", FormulaRule(formula=[f"LEN({ck}{R0})>0"], fill=RED))
    B.dv(lw, f"A{R0}:A{nLast}", "=LevelList", "Pick a level."); B.dv(lw, f"B{R0}:B{nLast}", "=SubjectList", "Pick a subject from Step 3.")
    B.dv(lw, f"C{R0}:C{nLast}", "=BlockList", "Pick a block from Step 5, or leave empty for a whole-class lesson.")
    if int(sc["weeks"]) == 2: B.dv(lw, f"{CL(len(names))}{R0}:{CL(len(names))}{nLast}", '"Odd,Even"', "Odd or Even: the group uses its special room (lab, studio...) on that week only, and a classroom on the other. Blank = every week.")
    B.dv(lw, f"{CF1}{R0}:{CFN}{nLast}", "=ClassList", "Pick a class from Step 2."); B.dv(lw, f"{TC[0]}{R0}:{TC[-1]}{nLast}", "=TeacherList", "Pick a teacher from Step 4.")
    lw.auto_filter.ref = f"A{HR}:{CL(len(names))}{R0 + max(nL, 1) - 1}"
    week_src = '"Odd,Even,Both"' if int(sc["weeks"]) == 2 else '"Both"'; day_src = '"All,Mon,Tue,Wed,Thu,Fri"'

    # ---------------- Teacher availability (clock times; a rule whose periods were changed gets the times of its periods)
    from .reader import periods_between
    aw = B.sheet(TA, "Times a teacher cannot teach", ["One row per block-out. Week: Odd, Even or Both. Day: Mon-Fri or All. Times as on Period times.",
                 ""], [("Teacher", 26), ("Week", 8), ("Day", 7), ("From", 10), ("To", 10), ("Reason", 60)], set(range(1, 7)), hr=4)
    av = rules["availability"]
    for i, x in enumerate(av):
        a, b = int(x["from"]), int(x["to"] if x.get("to") not in (None, "") else x["from"])
        tm = x.get("times") if x.get("times") and periods_between(sc, *x["times"]) == (a, b) else [_clock(sc, a).strftime("%H:%M"), _clock(sc, b + 1).strftime("%H:%M")]
        for j, v in enumerate([x.get("target"), x.get("week") or "Both", x.get("day") or "All", _t(tm[0]), _t(tm[1]), x.get("note")], start=1):
            c = aw.cell(5 + i, j, v)
            if j in (4, 5): c.number_format = "h:mm AM/PM"
    aL = 5 + len(av) + 60
    B.paint(aw, aL, set(range(1, 7)), 6, first=5)
    for r in range(5, aL + 1): aw[f"D{r}"].number_format = aw[f"E{r}"].number_format = "h:mm AM/PM"
    B.dv(aw, f"A5:A{aL}", "=TeacherList", "Pick a teacher from Step 4."); B.dv(aw, f"B5:B{aL}", week_src, "Odd, Even or Both.")
    B.dv(aw, f"C5:C{aL}", day_src, "All or a day.")

    # ---------------- Fixed lessons
    fw = B.sheet(FX, "Lessons that must sit at a set time (optional)", ["Leave empty to let the timetable place everything. Start period as numbered on Period times.", ""],
                 [("Level", 8), ("Class(es) or group", 18), ("Subject", 32), ("Week", 8), ("Day", 7), ("Start period", 9), ("Note", 40)], set(range(1, 8)), hr=4)
    lv_of = {c["class"]: c.get("level") for c in spec["classes"]}
    fx = rules["fixed"]
    for i, x in enumerate(fx):
        for j, v in enumerate([x.get("level") or lv_of.get(x.get("target")), x.get("target"), x.get("subject"), x.get("week") or "Both", x.get("day") or "All",
                               x.get("from"), x.get("note")], start=1): fw.cell(5 + i, j, v)
    fL = 5 + len(fx) + 40
    B.paint(fw, fL, set(range(1, 8)), 7, first=5)
    for i, x in enumerate(fx):
        if str(x.get("note") or "").startswith("example"):
            for j in range(1, 8): fw.cell(5 + i, j).font = EXF
    B.dv(fw, f"A5:A{fL}", "=LevelList", "Pick a level."); B.dv(fw, f"C5:C{fL}", "=SubjectList", "Pick a subject from Step 3.")
    B.dv(fw, f"D5:D{fL}", week_src, "Odd, Even or Both."); B.dv(fw, f"E5:E{fL}", day_src, "All or a day."); B.dv(fw, f"F5:F{fL}", "=PeriodList", "A period number (see Period times).")

    # ---------------- More rules + Rule types (every rule that has no place above)
    rt = wb.create_sheet("Rule types")
    rt["A1"] = f"Rule types you can use on {MR}"; rt["A1"].font = TITLE
    B.head(rt, 3, ["Rule", "Can apply to", "Needs From/To", "Needs Number", "What it means", "Example"])
    for i, row in enumerate(RULE_TYPES, start=4):
        for j, v in enumerate(row, start=1): c = rt.cell(i, j, v); c.font = BF; c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEF", (26, 30, 10, 10, 60, 44)): rt.column_dimensions[col].width = w
    nRT = len(RULE_TYPES); B.name("RuleList", f"'Rule types'!$A$4:$A${3 + nRT}")
    rcols = [("Applies to (kind)", 12), ("Applies to", 30), ("Subject (optional)", 28), ("Rule", 26), ("Week", 8), ("Day", 7), ("From period", 8), ("To period", 8),
             ("Number", 8), ("Must / Prefer", 9), ("Note", 40), ("Times", 20), ("Check", 36)]
    RIN = set(range(1, 12))
    more = rules["more"]
    rw = B.sheet(MR, "More rules (optional): any other constraint", [f"For rules that have no place on Settings, {TA} or {FX}: rooms, levels, subjects, Prefer rules and so on. First pick what a rule applies to; the next drop-down then lists only those.",
                 "Must = never broken; Prefer = kept where possible. See 'Rule types' for every rule and an example.", "Periods are numbered as on Period times (Times shows the clock)."], rcols, RIN)
    for i, x in enumerate(more):
        for j, k in enumerate(["scope", "target", "subject", "rule", "week", "day", "from", "to", "number", "strength", "note"], start=1): rw.cell(R0 + i, j, x.get(k))
    rL = R0 + len(more) + 80
    B.paint(rw, rL, RIN, 13)
    SC = "{" + ",".join(f'"{s}"' for s in SCOPES) + "}"
    LISTS = "SchoolList,LevelList,ClassList,TeacherList,DeptList,SubjectList,BlockList,BlockList,RoomList"   # Group: a code, so any text is accepted
    rtA = f"'Rule types'!$A$4:$A${3 + nRT}"
    for r in range(R0, rL + 1):
        rw[f"L{r}"] = f'=IF(G{r}="","",TEXT(INDEX({PTB},G{r}+1),"h:mm AM/PM")&" - "&TEXT(INDEX({PTC},IF(H{r}="",G{r},H{r})+1),"h:mm AM/PM"))'
        nt, nn, scp = (f"INDEX('Rule types'!${c}$4:${c}${3 + nRT},MATCH(D{r},{rtA},0))" for c in ("C", "D", "B"))
        rw[f"M{r}"] = (f'=IF(A{r}="","",TRIM(IF(AND(A{r}<>"Group",COUNTIF(CHOOSE(MATCH(A{r},{SC},0),{LISTS}),B{r})=0),"Pick the "&LOWER(A{r})&" from the list. ","")'
                       f'&IF(D{r}="","Pick a rule. ",IF(ISERROR(FIND(A{r},{scp})),D{r}&" does not apply to a "&LOWER(A{r})&". ","")'
                       f'&IF(AND(IFERROR({nt},"N")="Y",G{r}="",H{r}=""),"Give the period(s). ","")&IF(AND(IFERROR({nn},"N")="Y",I{r}=""),"Give the Number. ",""))'
                       f'&IF(AND(OR(D{r}="Same time across classes",D{r}="Not on the same day as"),C{r}=""),"Pick the Subject. ","")))')
        rw[f"M{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    rw.conditional_formatting.add(f"M{R0}:M{rL}", FormulaRule(formula=[f"LEN(M{R0})>0"], fill=RED))
    B.dv(rw, f"A{R0}:A{rL}", "=RuleScopeList", "Pick what the rule applies to.")
    B.dv(rw, f"C{R0}:C{rL}", "=SubjectList", "Pick a subject."); B.dv(rw, f"D{R0}:D{rL}", "=RuleList", "Pick a rule (see Rule types).")
    B.dv(rw, f"E{R0}:E{rL}", "=WeekList", "Pick a week."); B.dv(rw, f"F{R0}:F{rL}", "=DayList", "All or a day.")
    B.dv(rw, f"G{R0}:H{rL}", "=PeriodList", "A period number (see Period times)."); B.dv(rw, f"J{R0}:J{rL}", "=StrengthList", "Must or Prefer.")

    # ---------------- Students
    nst = len(spec.get("students", []))
    SIN_ = set(range(1, 12))
    stw = B.sheet("Students", "Students (recommended)", ["Every student listed with their class and groups (group codes from Step 6, separated by commas) is checked for clashes. "
                  "Without it, students in a class are assumed to follow one level tag across their banded subjects.", ""],
                  [("Student ID", 12), ("Name", 30), ("Level", 8), (f"Class ({year})" if year else "Class", 10), ("Combination", 12), ("Choice 1", 30), ("Choice 2", 20),
                   ("Choice 3", 20), ("Merit", 8), ("Groups (codes)", 60), ("Source", 40)], SIN_, hr=4)
    for i, x in enumerate(spec.get("students", [])):
        ch = (x.get("choices") or []) + [None] * 3
        for j, v in enumerate([x.get("id"), x.get("name"), x.get("level"), x.get("class"), x.get("combination"), ch[0], ch[1], ch[2], x.get("merit"),
                               ", ".join(x.get("groups") or []), x.get("note")], start=1): stw.cell(5 + i, j, v)
    B.paint(stw, 5 + nst + 60, SIN_, 11, first=5)
    stw.freeze_panes = "C5"
    B.dv(stw, f"C5:C{5 + nst + 400}", "=LevelList", "Pick a level."); B.dv(stw, f"D5:D{5 + nst + 400}", "=ClassList", "Pick a class.")
    B.dv(stw, f"E5:E{5 + nst + 400}", "=ComboList", "Pick a combination from the Combinations sheet.")

    # ---------------- Combinations
    combos = spec.get("combinations") or []
    nsub = max([3] + [len(x.get("subjects") or []) for x in combos])
    cb = B.sheet("Combinations", "Subject combinations (optional)", ["Used to plan option blocks: subjects taken together must be in different blocks."],
                 [("Level", 8), ("Code", 8), ("Name", 60)] + [(f"Subject {k}", 30) for k in range(1, nsub + 1)] + [("Places", 8)], set(range(1, nsub + 5)), hr=3)
    for i, x in enumerate(combos):
        subs = (x.get("subjects") or []) + [None] * nsub
        for j, v in enumerate([x.get("level"), x.get("code"), x.get("name") or " + ".join(x.get("subjects") or [])] + subs[:nsub] + [x.get("places")], start=1):
            cb.cell(4 + i, j, v)
    B.paint(cb, 4 + len(combos) + 40, set(range(1, nsub + 5)), nsub + 4, first=4)
    B.name("ComboList", "OFFSET(Combinations!$B$4,0,0,MAX(1,COUNTA(Combinations!$B$4:$B$500)),1)")

    order = ["README"] + extra + ["Settings", "Period times", "Day structure", V, C, S, T, "Roles", BL, L, TA, FX, MR, "Rule types", "Students", "Combinations", "Lists"]
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)     # Excel works out every formula when the file is opened
    wb.save(path)
    return path
