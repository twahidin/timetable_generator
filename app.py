"""School timetable app (one school per deployment).

Committee (COMMITTEE_CODE): fill the setup step by step, generate the timetable (with the AI timetabler if one is set
up in Railway, or without), talk through conflicts in the chat, publish, and download everything.
Teachers (TEACHER_CODE): every published class, teacher, room and student timetable, Who's free, PDFs.
Students (STUDENT_CODE): class timetables and 'Build my timetable'; personal codes show one student's own timetable."""
import copy, io, json, os, shutil, tempfile, zipfile
from functools import wraps
from flask import Flask, jsonify, request, send_file, abort, session, redirect, Response
from werkzeug.middleware.proxy_fix import ProxyFix
from ttgen.samples import secondary, primary
from ttgen.template import write_template
from ttgen.reader import read_template
from ttgen.spec import RULE_TYPES, SCOPES, CODES, SETTING_RULES, ROLES, defaults
import access, store, runs, ai

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL_DIR = os.path.join(HERE, "skill", "school-timetable-generator")
app = Flask(__name__, static_folder="static", static_url_path="/static")
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)
app.secret_key = os.environ.get("SECRET_KEY") or os.urandom(32)
app.config.update(MAX_CONTENT_LENGTH=40 * 1024 * 1024, SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax",
                  SESSION_COOKIE_SECURE=bool(os.environ.get("RAILWAY_ENVIRONMENT")), PERMANENT_SESSION_LIFETIME=60 * 60 * 24 * 120)
limiter = access.Limiter()
_cache = {}
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# ------------------------------------------------------------------------------------------------ access
def role(): return session.get("role")

def need(*roles):
    def deco(f):
        @wraps(f)
        def inner(*a, **k):
            if role() not in roles:
                if request.path.startswith(("/api/", "/pdf/", "/files/")): return {"error": "Please sign in again."}, 401
                return redirect("/login")
            return f(*a, **k)
        return inner
    return deco

def esc(s):
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")

def page(title, body):
    return f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{esc(title)}</title>
<link rel="stylesheet" href="/static/login.css"></head><body>{body}</body></html>"""

# the three choices on the front page; each one accepts only its own kind of code
CHOICES = {
    "committee": dict(title="Generate a new timetable", blurb="Timetable committee: fill the setup, generate and publish.",
                      label="Committee code", hint="Use the code given to the timetable committee.", roles=("committee",)),
    "teacher": dict(title="Teacher's timetable", blurb="Every class, teacher and room timetable, Who's free and PDFs.",
                    label="Teacher code", hint="Use the teacher code the school gave you.", roles=("teacher",)),
    "student": dict(title="Student timetable", blurb="Your class timetable, or your own with a personal code.",
                    label="Student code or personal code", hint="Use the student code the school gave you. A personal code shows only your own timetable.",
                    roles=("student", "personal")),
}
CHOICE_OF = {r: k for k, c in CHOICES.items() for r in c["roles"]}

@app.get("/login")
def login_page(msg="", choice=None, fix=None):
    pub = store.published(); school = (pub["meta"]["school"] if pub else None) or os.environ.get("SCHOOL_NAME") or "School timetable"
    _, probs = access.configured()
    warn = "".join(f"<p class='warn'>{esc(p)}</p>" for p in probs)
    link = f' <a href="/login?as={fix}">Open {esc(CHOICES[fix]["title"])}</a>' if fix else ""
    err = f"<p class='err'>{esc(msg)}{link}</p>" if msg else ""
    choice = choice or request.args.get("as")
    c = CHOICES.get(choice or "")
    if not c:
        cards = "".join(f"""<a class="choice" href="/login?as={k}"><img src="/static/choice-{k}.png" alt="" width="112" height="112">
<span class="txt"><b>{esc(v["title"])}</b><span class="muted small">{esc(v["blurb"])}</span></span><span class="go">Enter code &rarr;</span></a>""" for k, v in CHOICES.items())
        return page("Sign in", f"""<main class="front"><header class="hero"><h1>{esc(school)}</h1><p>Timetable</p></header>
<nav class="choices">{cards}</nav>{f'<div class="notes">{err}{warn}</div>' if err or warn else ""}</main>""")
    return page(c["title"], f"""<main class="login"><p class="small"><a class="back" href="/login">&larr; Other choices</a></p>
<div class="picked"><img src="/static/choice-{choice}.png" alt="" width="56" height="56"><span><h1>{esc(c["title"])}</h1><p class="muted">{esc(school)}</p></span></div>
<form method="post" action="/login"><input type="hidden" name="as" value="{choice}"><label for="code">{esc(c["label"])}</label>
<input id="code" name="code" autocomplete="off" autocapitalize="characters" autofocus required>
<button type="submit">Open</button></form>{err}
<p class="muted small">{esc(c["hint"])}</p>{warn}</main>""")

@app.post("/login")
def login():
    who = request.remote_addr or "?"
    choice = request.form.get("as")
    if choice not in CHOICES: return redirect("/login")
    if limiter.blocked(who): return login_page("Too many wrong codes. Wait ten minutes and try again.", choice), 429
    pub = store.published()
    r, sid = access.match(request.form.get("code"), pub["codes"] if pub else {})
    if not r:
        limiter.fail(who); return login_page("That code does not work. Check it and try again.", choice), 401
    if r not in CHOICES[choice]["roles"]:   # a real code, but for another choice: not a guess, so no limiter strike
        return login_page(f"That code is not for {CHOICES[choice]['title']}.", choice, fix=CHOICE_OF[r]), 403
    session.clear(); session.permanent = True; session["role"] = r
    if sid is not None: session["sid"] = sid
    return redirect({"committee": "/", "personal": "/me"}.get(r, "/timetable"))

@app.get("/logout")
def logout():
    session.clear(); return redirect("/login")

@app.get("/healthz")
def health(): return {"ok": True}

# ------------------------------------------------------------------------------------------------ committee: setup editor
@app.get("/")
def index():
    if role() == "committee": return app.send_static_file("index.html")
    if role() == "personal": return redirect("/me")
    if role() in ("teacher", "student"): return redirect("/timetable")
    return redirect("/login")

def kind_ok(kind):
    if kind not in ("primary", "secondary"): abort(404)
    return kind

def sample(kind):
    if ("s", kind) not in _cache: _cache[("s", kind)] = secondary() if kind == "secondary" else primary()
    return copy.deepcopy(_cache[("s", kind)])

def blank(kind):
    s = sample(kind)
    keep = [r for r in s["rules"] if r["rule"] in ("Once a day per group", "Avoid after-school slots")]
    s.update(venues=[], classes=[], teachers=[], blocks=[], lessons=[], students=[], combinations=[], rules=keep)
    s["school"]["name"] = os.environ.get("SCHOOL_NAME") or "My School"
    return s

def xlsx_bytes(spec):
    buf = io.BytesIO(); write_template(spec, buf); buf.seek(0); return buf

def normalise(spec):
    """the editor's JSON -> the exact spec the solver would read from the workbook"""
    for k in ("venues", "classes", "subjects", "teachers", "blocks", "lessons", "rules", "students"): spec.setdefault(k, [])
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp: path = tmp.name
    try:
        write_template(spec, path); return read_template(path)
    finally: os.unlink(path)

@app.get("/api/meta/<kind>")
@need("committee")
def meta(kind):
    sc = defaults(kind_ok(kind))
    return jsonify(rule_types=[dict(rule=a, scopes=b, needs_time=c == "Y", needs_number=d == "Y", meaning=e, example=f) for a, b, c, d, e, f in RULE_TYPES],
                   scopes=SCOPES, codes=CODES, defaults=sc, setting_rules=SETTING_RULES,
                   roles=[dict(role=a, what=b, allowance=c, group=d) for a, b, c, d in ROLES])

@app.get("/api/sample/<kind>")
@need("committee")
def api_sample(kind): return jsonify(sample(kind_ok(kind)))

@app.get("/api/blank/<kind>")
@need("committee")
def api_blank(kind): return jsonify(blank(kind_ok(kind)))

@app.post("/api/build")
@need("committee")
def build():
    spec = request.get_json(force=True, silent=True)
    if not spec or "school" not in spec: return {"error": "No setup data sent."}, 400
    for k in ("venues", "classes", "subjects", "teachers", "blocks", "lessons", "rules", "students"): spec.setdefault(k, [])
    try: buf = xlsx_bytes(spec)
    except Exception as e: return {"error": f"Could not build the workbook: {e}"}, 400
    name = "".join(ch for ch in (spec["school"].get("name") or "School") if ch.isalnum() or ch in " -_").strip().replace(" ", "_") or "School"
    return send_file(buf, as_attachment=True, download_name=f"{name}_Timetable_Setup.xlsx", mimetype=XLSX)

@app.post("/api/import")
@need("committee")
def import_xlsx():
    f = request.files.get("file")
    if not f: return {"error": "No file."}, 400
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        f.save(tmp.name); path = tmp.name
    try: spec = read_template(path)
    except Exception as e: return {"error": f"That does not look like a setup workbook from this app ({e})."}, 400
    finally: os.unlink(path)
    return jsonify(spec)

@app.post("/api/read-table")
@need("committee")
def read_table():
    """the rows of a spreadsheet uploaded on the Students step: its Students sheet if it has one, else its first sheet"""
    import openpyxl
    f = request.files.get("file")
    if not f: return {"error": "No file."}, 400
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        f.save(tmp.name); path = tmp.name
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb["Students"] if "Students" in wb.sheetnames else wb.worksheets[0]
        cell = lambda v: "" if v is None else str(int(v)) if isinstance(v, float) and v.is_integer() else str(v).strip()
        rows = [[cell(v) for v in r] for r in ws.iter_rows(max_row=5000, values_only=True)]
        wb.close()
    except Exception as e: return {"error": f"Could not read that file ({e})."}, 400
    finally: os.unlink(path)
    return {"rows": [r for r in rows if any(r)]}

@app.get("/api/setup")
@need("committee")
def get_setup():
    spec, at = store.load_setup()
    return jsonify(spec=spec, saved_at=at)

@app.put("/api/setup")
@need("committee")
def put_setup():
    spec = request.get_json(force=True, silent=True)
    if not spec or "school" not in spec: return {"error": "No setup data sent."}, 400
    store.save_setup(spec); return {"ok": True, "saved_at": store.now()}

@app.post("/api/export")
@need("committee")
def export_printouts():
    """setup + timetable workbooks -> zip with the combined Excel, the HTML finder and the PDFs"""
    from ttgen.export import model_from_files, export_all, safe_name
    a, b = request.files.get("setup"), request.files.get("timetable")
    if not a or not b: return {"error": "Send the setup workbook and the timetable workbook."}, 400
    work = tempfile.mkdtemp(prefix="tt-")
    try:
        sp, tp = os.path.join(work, "setup.xlsx"), os.path.join(work, "timetable.xlsx"); a.save(sp); b.save(tp)
        try: m = model_from_files(sp, tp)
        except Exception as e: return {"error": f"Could not read the workbooks: {e}"}, 400
        if not m["lessons"]: return {"error": "The timetable workbook's Lessons sheet has no lessons."}, 400
        base = safe_name(f"{m['school']['name']} {m['school']['timetable']}")
        out = os.path.join(work, "out", base)
        export_all(m, out, individual=request.form.get("single", "1") == "1", zip_individual=False, log=lambda *x: None)
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for root, _, files in os.walk(out):
                for fn in sorted(files): full = os.path.join(root, fn); z.write(full, os.path.join(base, os.path.relpath(full, out)))
        buf.seek(0); ch = m["checks"]
        resp = send_file(buf, as_attachment=True, download_name=f"{base} - timetables.zip", mimetype="application/zip")
        resp.headers["X-Checks"] = "All checks passed." if ch["ok"] else f"{len(ch['problems'])} problems found: see the Read me sheet or the Checks tab."
        return resp
    finally: shutil.rmtree(work, ignore_errors=True)

@app.get("/download/template/<kind>")
@need("committee")
def dl_sample(kind):
    kind_ok(kind)
    return send_file(xlsx_bytes(sample(kind)), as_attachment=True, download_name=f"Sample_{kind.title()}_Timetable_Setup.xlsx", mimetype=XLSX)

@app.get("/download/blank/<kind>")
@need("committee")
def dl_blank(kind):
    kind_ok(kind)
    return send_file(xlsx_bytes(blank(kind)), as_attachment=True, download_name=f"Blank_{kind.title()}_Timetable_Setup.xlsx", mimetype=XLSX)

@app.get("/download/SKILL.md")
@need("committee")
def dl_skill_md():
    return send_file(os.path.join(SKILL_DIR, "SKILL.md"), as_attachment=True, download_name="SKILL.md", mimetype="text/markdown")

@app.get("/download/skill.zip")
@need("committee")
def dl_skill_zip():
    if "zip" not in _cache:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for root, _, files in os.walk(SKILL_DIR):
                for fn in files:
                    if fn.endswith(".pyc"): continue
                    full = os.path.join(root, fn); z.write(full, os.path.join("school-timetable-generator", os.path.relpath(full, SKILL_DIR)))
            for kind in ("secondary", "primary"):
                z.writestr(f"school-timetable-generator/templates/Sample_{kind.title()}_Timetable_Setup.xlsx", xlsx_bytes(sample(kind)).getvalue())
        _cache["zip"] = buf.getvalue()
    return send_file(io.BytesIO(_cache["zip"]), as_attachment=True, download_name="school-timetable-generator.zip", mimetype="application/zip")

# ------------------------------------------------------------------------------------------------ committee: generate, chat, publish
def pub_info():
    pub = store.published()
    if not pub: return None
    m = pub["meta"]
    return dict(at=m["published_at"], ok=m["ok"], after=len(m.get("after_school", [])), students=m.get("students", 0), classes=m.get("classes", 0),
                files=m.get("files", {}), problems=m.get("problems", []))

@app.get("/api/status")
@need("committee")
def status():
    have, probs = access.configured()
    _, at = store.load_setup(); r = runs.current()
    return jsonify(ai=ai.public_config(), codes=have, code_problems=probs, volume=store.on_volume(), setup_saved_at=at,
                   run=r.snapshot(int(request.args.get("after") or 0)) if r else None, published=pub_info(),
                   solve_seconds=int(os.environ.get("SOLVE_SECONDS") or 300))

@app.post("/api/generate")
@need("committee")
def generate():
    body = request.get_json(force=True, silent=True) or {}
    spec = body.get("spec")
    if spec: store.save_setup(spec)
    else: spec, _ = store.load_setup()
    if not spec: return {"error": "Fill and save a setup first."}, 400
    use_ai = bool(body.get("ai")) and ai.config()["ok"]
    try: spec = normalise(copy.deepcopy(spec))
    except Exception as e: return {"error": f"The setup could not be read: {e}"}, 400
    try: r = runs.start(spec, use_ai)
    except RuntimeError as e: return {"error": str(e)}, 409
    return jsonify(id=r.id, ai=use_ai)

@app.post("/api/message")
@need("committee")
def message():
    text = ((request.get_json(force=True, silent=True) or {}).get("text") or "").strip()
    if not text: return {"error": "Empty message."}, 400
    try: runs.message(text[:4000])
    except RuntimeError as e: return {"error": str(e)}, 409
    return {"ok": True}

@app.post("/api/stop")
@need("committee")
def stop():
    runs.stop(); return {"ok": True}

@app.post("/api/keep-changes")
@need("committee")
def keep_changes():
    r = runs.current()
    if not r or not r.changes: return {"error": "There are no changes to keep."}, 400
    store.save_setup(r.spec); r.emit("status", text="The committee kept these changes: the saved setup now includes them.")
    return jsonify(spec=r.spec)

def _publish_job(r):
    s = r.solved
    r.emit("status", text="Publishing: making the PDFs, the Excel workbook and the timetable pages (about a minute)...")
    m = store.publish(s["setup"], s["timetable"], run_id=r.id, log=lambda t: None)
    r.emit("published", text=f"Published. Teachers and students now see this timetable. {len(m['students'])} students have personal codes.")

@app.post("/api/publish")
@need("committee")
def publish():
    r = runs.current()
    if not r or not r.solved: return {"error": "Generate a timetable first."}, 400
    if r.busy: return {"error": "Wait for the current work to finish (or stop it) before publishing."}, 409
    runs._thread(r, _publish_job)
    return {"ok": True}

@app.get("/preview")
@need("committee")
def preview():
    """the latest generated (not yet published) timetable, committee view"""
    from ttgen.export import model_from_files, render_html
    r = runs.current()
    if not r or not r.solved: return page("Preview", "<main class='login'><p>Nothing generated yet.</p><p><a href='/'>Back</a></p></main>")
    key = ("preview", r.id, r.solved["n"])
    if _cache.get("pv_key") != key:
        _cache["pv"] = render_html(model_from_files(r.solved["setup"], r.solved["timetable"]), "full",
                                   app=dict(mode="committee", pdf=False, home="/", homeLabel="Back to the app", logout="/logout"),
                                   title="Preview (not published)")
        _cache["pv_key"] = key
    return Response(_cache["pv"], mimetype="text/html")

@app.get("/api/codes.csv")
@need("committee")
def codes_csv():
    pub = store.published()
    if not pub: abort(404)
    return Response(store.codes_csv(pub["model"]), mimetype="text/csv", headers={"Content-Disposition": 'attachment; filename="personal timetable codes.csv"'})

@app.get("/api/code-slips.pdf")
@need("committee")
def code_slips():
    from ttgen.export import code_slips_pdf
    pub = store.published()
    if not pub: abort(404)
    buf = io.BytesIO()
    code_slips_pdf(pub["model"], {s["id"]: access.personal_code(s["id"]) for s in pub["model"]["students"]}, request.host_url.rstrip("/"), buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="timetable code slips.pdf", mimetype="application/pdf")

# ------------------------------------------------------------------------------------------------ everyone: the published timetable
TABS = dict(committee=None, teacher=["classes", "teachers", "rooms", "students", "lessons", "free"], student=["classes", "students"])

def files_for(r):
    if r == "committee": return [["Excel (all)", "/files/excel"], ["PDFs (zip)", "/files/zip"]]
    if r == "teacher": return [["Class PDFs", "/files/book/classes"], ["Teacher PDFs", "/files/book/teachers"]]
    return []

@app.get("/timetable")
@need("committee", "teacher", "student")
def timetable():
    from ttgen.export import render_html
    pub = store.published()
    if not pub:
        return page("Timetable", "<main class='login'><h1>Not published yet</h1><p class='muted'>The timetable committee has not published the timetable yet. Please check again later.</p><p><a href='/logout'>Sign out</a></p></main>")
    r = role(); key = ("view", r)
    if key not in pub["views"]:
        view = "student" if r == "student" else "full"
        app_cfg = dict(mode=r, tabs=TABS[r], pdf=True, logout="/logout", files=files_for(r),
                       placeholder="Search a class" if r == "student" else None, home="/" if r == "committee" else None, homeLabel="Committee page")
        pub["views"][key] = render_html(pub["model"], view, app=app_cfg)
    return Response(pub["views"][key], mimetype="text/html")

@app.get("/me")
@need("personal")
def me():
    from ttgen.export import render_html
    pub = store.published(); sid = session.get("sid")
    if not pub or sid not in {s["id"] for s in pub["model"]["students"]}:
        session.clear()
        return page("Timetable", "<main class='login'><h1>Code no longer valid</h1><p class='muted'>Your timetable is not in the published timetable. Ask your form teacher for a new code.</p><p><a href='/login'>Back</a></p></main>")
    return Response(render_html(pub["model"], "personal", student_id=sid, app=dict(mode="personal", tabs=[], pdf=True, logout="/logout")), mimetype="text/html")

def _pdf(v, name):
    from ttgen.export import pdf_bytes, safe_name
    return send_file(io.BytesIO(pdf_bytes(store.published()["model"], v)), as_attachment=True, download_name=safe_name(name) + ".pdf", mimetype="application/pdf")

@app.get("/pdf/me")
@need("personal")
def pdf_me():
    from ttgen.export import student_view
    pub = store.published(); sid = session.get("sid")
    s = next((x for x in pub["model"]["students"] if x["id"] == sid), None) if pub else None
    if not s: abort(404)
    return _pdf(student_view(pub["model"], s), f"{s['name']} timetable")

@app.get("/pdf/custom")
@need("committee", "teacher", "student")
def pdf_custom():
    from ttgen.export import custom_view, group_options
    pub = store.published()
    if not pub: abort(404)
    m = pub["model"]; cls = request.args.get("cls", "")
    if cls not in m["class_ix"]: abort(404)
    allowed = {g for gs in group_options(m, cls).values() for g in gs}
    groups = [g for g in (request.args.get("groups") or "").split(",") if g in allowed]
    return _pdf(custom_view(m, cls, groups), f"Class {cls} my timetable")

@app.get("/pdf/<kind>/<path:key>")
@need("committee", "teacher", "student")
def pdf_one(kind, key):
    from ttgen.export import class_view, teacher_view, room_view, student_view
    pub = store.published()
    if not pub: abort(404)
    m = pub["model"]; r = role()
    if r == "student" and kind != "classes": abort(403)
    if kind == "classes" and key in m["class_ix"]: return _pdf(class_view(m, key), f"Class {key}")
    if kind == "teachers" and key in m["teacher_ix"]: return _pdf(teacher_view(m, key), key)
    if kind == "rooms" and key in m["room_ix"]: return _pdf(room_view(m, key), f"Room {key}")
    if kind == "students":
        s = next((x for x in m["students"] if x["id"] == key), None)
        if s: return _pdf(student_view(m, s), f"{s['id']} {s['name']}")
    abort(404)

@app.get("/files/<what>")
@app.get("/files/<what>/<kind>")
@need("committee", "teacher")
def files(what, kind=None):
    pub = store.published()
    if not pub: abort(404)
    f = pub["meta"].get("files", {}); r = role()
    if what == "book":
        allowed = {"committee": ("classes", "teachers", "rooms", "students"), "teacher": ("classes", "teachers", "rooms")}[r]
        if kind not in allowed or kind not in f.get("books", {}): abort(404)
        return send_file(os.path.join(store.PUB, f["books"][kind]), as_attachment=True, download_name=os.path.basename(f["books"][kind]))
    if r != "committee": abort(403)
    if what in ("excel", "zip", "html") and f.get(what):
        return send_file(os.path.join(store.PUB, f[what]), as_attachment=True, download_name=f[what])
    abort(404)

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 8000)), debug=False)
