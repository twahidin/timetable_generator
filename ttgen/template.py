"""Write a guided timetable setup workbook from a spec (see spec.py). Drop-downs flow from step to step;
group codes, lesson patterns, loads and checks are Excel formulas."""
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.workbook.defined_name import DefinedName
from .spec import RULE_TYPES, SCOPES, CODES, day_names, scopes_for

NAVY, SLATE = "1F4E78", "5B6B7F"
HF = Font(name="Arial", bold=True, color="FFFFFF", size=10)
INF = Font(name="Arial", size=10, color="0000FF"); INFILL = PatternFill("solid", fgColor="FFF9DB")
BF = Font(name="Arial", size=10); CF = PatternFill("solid", fgColor="EDEDED")
TITLE = Font(name="Arial", size=14, bold=True, color=NAVY); NOTE = Font(name="Arial", size=9, italic=True, color="595959")
EXF = Font(name="Arial", size=10, italic=True, color="7F7F7F")
HR, R0 = 5, 6          # header row, first data row
NCLS, NTCH, NPAIR = 12, 3, 8
BOX = Border(*(Side(style="thin", color="C9D3E0"),) * 4)

class Book:
    def __init__(self):
        self.wb = openpyxl.Workbook(); self.wb.remove(self.wb.active)
    def sheet(self, name, title, notes, cols, inputs):
        ws = self.wb.create_sheet(name)
        ws["A1"] = title; ws["A1"].font = TITLE
        for i, n in enumerate(notes[:3], start=2): ws.cell(i, 1, n).font = NOTE
        for j, (h, w) in enumerate(cols, start=1):
            c = ws.cell(HR, j, h); c.font = HF; c.alignment = Alignment(wrap_text=True, vertical="center")
            c.fill = PatternFill("solid", fgColor=NAVY if j in inputs else SLATE)
            ws.column_dimensions[CL(j)].width = w
        ws.row_dimensions[HR].height = 30; ws.freeze_panes = ws.cell(R0, 1)
        return ws
    def paint(self, ws, last, inputs, ncols):
        for r in range(R0, last + 1):
            for j in range(1, ncols + 1):
                c = ws.cell(r, j)
                if j in inputs: c.font = INF; c.fill = INFILL
                else: c.font = BF; c.fill = CF
    def dv(self, ws, rng, src, msg):
        d = DataValidation(type="list", formula1=src, allow_blank=True, showErrorMessage=True, errorTitle="Pick from the list", error=msg)
        ws.add_data_validation(d); d.add(rng)
    def name(self, nm, ref): self.wb.defined_names[nm] = DefinedName(nm, attr_text=ref)
    def dyn(self, sheet, col, first=R0, n=1000):
        return f"OFFSET('{sheet}'!${col}${first},0,0,MAX(1,COUNTA('{sheet}'!${col}${first}:${col}${first + n - 1})),1)"

def _t(s):
    if isinstance(s, datetime.time): return s
    h, m = str(s or "07:30").split(":")[:2]; return datetime.time(int(h), int(m))

def write_template(spec, path, extra_rows=120):
    sc = spec["school"]; levels = sc["levels"]; days = day_names(sc["weeks"]); NP = int(sc["periods_per_day"])
    B = Book(); wb = B.wb
    V, C, S, T, BL, L, RU = "1 Venues", "2 Classes", "3 Subjects", "4 Teachers", "5 Blocks", "6 Lessons", "7 Rules"
    nV, nC, nS = len(spec["venues"]), len(spec["classes"]), len(spec["subjects"])
    nT, nB, nL, nR = len(spec["teachers"]), len(spec["blocks"]), len(spec["lessons"]), len(spec.get("rules", []))
    LAST = lambda n: R0 + n + extra_rows

    # ---------------- README
    rd = wb.create_sheet("README")
    steps = [("0", "Settings, Day structure", "School name, cycle, period length and start time; what each period is for each level and day (F flag/assembly, L lesson, R recess, U lunch, A after-school slot)."),
             ("1", V, "List every room and pick its kind."),
             ("2", C, "List the classes; pick each one's level and home room (from step 1)."),
             ("3", S, "List subjects with code, level tag, department, room kind and periods per cycle for each level. Lesson length is in periods; leftovers become shorter lessons."),
             ("4", T, "List teachers; pick the subjects (step 3) and levels each can teach. Periods given and Spare update from step 6."),
             ("5", BL, "Name each block of groups that run at the same time and pick its classes (step 2). Its groups and student totals fill in from step 6."),
             ("6", L, "One row per group (or per class for whole-class subjects): level, subject, block, classes, teachers. Group code, periods, lesson pattern, room kind and checks fill in."),
             ("7", RU, "Constraints: pick what a rule applies to (school, level, class, teacher, department, subject, block, room), the rule and when. 'Rule types' lists every rule."),
             ("8", "Students", "Optional: each student's class and groups, so every student is checked for clashes.")]
    rd["A1"] = f"{sc['name']} · timetable setup ({sc['timetable']})"; rd["A1"].font = TITLE
    rd["A2"] = "Fill the sheets in order. Each step's list becomes the drop-down for the next, so nothing is typed twice and nothing is misspelt."; rd["A2"].font = BF
    rd["A3"] = "Blue text on yellow = type or pick.  Grey = worked out by formulas.  Red in a Check column = something to fix.  Grey italic rows = examples to replace."; rd["A3"].font = BF
    for j, h in enumerate(["Step", "Sheet", "What you do"], start=1):
        c = rd.cell(5, j, h); c.font = HF; c.fill = PatternFill("solid", fgColor=NAVY)
    for i, (a, b, c_) in enumerate(steps, start=6):
        rd.cell(i, 1, a).font = Font(name="Arial", bold=True, size=10); rd.cell(i, 2, b).font = Font(name="Arial", bold=True, size=10)
        x = rd.cell(i, 3, c_); x.font = BF; x.alignment = Alignment(wrap_text=True, vertical="top")
    r = 6 + len(steps) + 1
    for t_ in ["How group codes are made", f"Level digit + subject code + level letter ({', '.join(f'{a} {b}' for a, b in sc['tags'] if b)}) + a running number, e.g. 1ELP1. To keep an existing school code, type it in 'Code (only to keep an old one)'.",
               "Generating the timetable", "Give this workbook to an AI agent with the timetable skill (SKILL.md), or run: python -m ttgen solve <this file> --out timetable.xlsx"]:
        c = rd.cell(r, 1, t_); c.font = Font(name="Arial", size=11, bold=True, color=NAVY) if len(t_) < 40 else BF
        rd.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3); c.alignment = Alignment(wrap_text=True, vertical="top"); r += 1
    rd.column_dimensions["A"].width = 7; rd.column_dimensions["B"].width = 26; rd.column_dimensions["C"].width = 110

    # ---------------- Settings / Period times
    st = wb.create_sheet("Settings")
    st["A1"] = "Settings"; st["A1"].font = TITLE
    st["A2"] = "Blue on yellow = edit. Period times and the Day structure headers follow from these."; st["A2"].font = NOTE
    for j, h in enumerate(["Setting", "Value", "Notes"], start=1):
        c = st.cell(4, j, h); c.font = HF; c.fill = PatternFill("solid", fgColor=NAVY)
    rows = [("School", sc["name"], ""), ("Timetable", sc["timetable"], ""), ("School type", sc["type"], "primary or secondary"),
            ("Weeks in cycle", int(sc["weeks"]), "1 = same every week; 2 = Odd and Even weeks"),
            ("Period length (minutes)", int(sc["period_minutes"]), ""), ("Period 0 starts at", _t(sc["start_time"]), ""),
            ("Periods in a day", NP, f"Numbered 0 to {NP - 1}."), ("Standard lesson length (periods)", int(sc["lesson_length"]), "A subject's leftover periods become one or two shorter lessons.")]
    for i, (k, v, n) in enumerate(rows, start=5):
        st.cell(i, 1, k).font = Font(name="Arial", bold=True, size=10); c = st.cell(i, 2, v); c.font = INF; c.fill = INFILL
        st.cell(i, 3, n).font = NOTE
        if isinstance(v, datetime.time): c.number_format = "h:mm AM/PM"
    st.column_dimensions["A"].width = 34; st.column_dimensions["B"].width = 28; st.column_dimensions["C"].width = 60
    pt = wb.create_sheet("Period times")
    pt["A1"] = "Period times"; pt["A1"].font = TITLE; pt["A2"] = "Worked out from Settings. Do not type here."; pt["A2"].font = NOTE
    for j, h in enumerate(["Period", "Starts", "Ends"], start=1):
        c = pt.cell(4, j, h); c.font = HF; c.fill = PatternFill("solid", fgColor=SLATE)
    for p in range(NP):
        rr = 5 + p; pt.cell(rr, 1, p).font = BF
        pt.cell(rr, 2, f"=Settings!$B$10+A{rr}*Settings!$B$9/1440").number_format = "h:mm AM/PM"
        pt.cell(rr, 3, f"=B{rr}+Settings!$B$9/1440").number_format = "h:mm AM/PM"
    PTB, PTC = f"'Period times'!$B$5:$B${4 + NP}", f"'Period times'!$C$5:$C${4 + NP}"

    # ---------------- Day structure
    ds = wb.create_sheet("Day structure")
    ds["A1"] = "Day structure: what each period is, per level and day"; ds["A1"].font = TITLE
    ds["A2"] = "Codes: " + "   ".join(f"{k} = {v.lower()}" for k, v in CODES.items() if k != ".") + "   blank = no school"; ds["A2"].font = NOTE
    for j, h in enumerate(["Level", "Day"] + list(range(NP)), start=1):
        c = ds.cell(4, j, h); c.font = HF; c.fill = PatternFill("solid", fgColor=NAVY); c.alignment = Alignment(horizontal="center")
    for p in range(NP):
        c = ds.cell(5, 3 + p, f"=TEXT('Period times'!B{5 + p},\"h:mm\")"); c.font = Font(name="Arial", size=7, color="595959")
        c.alignment = Alignment(horizontal="center", textRotation=90)
    ds.row_dimensions[5].height = 30
    rr = 6
    for lv in levels:
        for d in days:
            ds.cell(rr, 1, lv).font = Font(name="Arial", bold=True, size=10); ds.cell(rr, 2, d).font = Font(name="Arial", bold=True, size=10)
            code = (spec["day_structure"].get(lv, {}).get(d, "") or "").ljust(NP, ".")
            for p in range(NP):
                v = code[p] if code[p] != "." else None
                c = ds.cell(rr, 3 + p, v); c.font = INF; c.alignment = Alignment(horizontal="center"); c.border = BOX
            rr += 1
    rng = f"C6:{CL(2 + NP)}{rr - 1}"
    B.dv(ds, rng, '"F,L,R,U,A"', "Use F, L, R, U, A or leave blank.")
    for code, col in (("F", "C9C9C9"), ("L", "CFE0F3"), ("R", "F6C28B"), ("U", "A9D8B0"), ("A", "FFE699")):
        ds.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{code}"'], fill=PatternFill("solid", fgColor=col)))
    ds.column_dimensions["A"].width = 8; ds.column_dimensions["B"].width = 11
    for p in range(NP): ds.column_dimensions[CL(3 + p)].width = 4.6
    ds.freeze_panes = "C6"

    # ---------------- Lists
    ls = wb.create_sheet("Lists")
    ls["A1"] = "Lists behind the drop-downs (add an item at the bottom of a column to make it pickable)"; ls["A1"].font = TITLE
    lists = [("Levels", levels), ("Departments", sc["departments"]), ("Room kinds", sc["room_kinds"]), ("Level tags", [a for a, b in sc["tags"]]),
             ("Tag letter", [b for a, b in sc["tags"]]), ("Teaches at", scopes_for(sc)), ("Level", list(sc["bands"])), ("Band", list(sc["bands"].values())),
             ("Yes/No", ["Yes", "No"]), ("Rule applies to", SCOPES), ("School", ["Whole school"]),
             ("Week", ["Both", "Odd", "Even"] if int(sc["weeks"]) == 2 else ["Both"]), ("Day", ["All"] + ["Mon", "Tue", "Wed", "Thu", "Fri"]),
             ("Must/Prefer", ["Must", "Prefer"]), ("Periods", list(range(NP)))]
    for j, (h, vals) in enumerate(lists, start=1):
        c = ls.cell(3, j, h); c.font = HF; c.fill = PatternFill("solid", fgColor=NAVY); ls.column_dimensions[CL(j)].width = 15
        for i, v in enumerate(vals, start=4): ls.cell(i, j, v).font = INF; ls.cell(i, j).fill = INFILL
    for nm, col in (("LevelList", "A"), ("DeptList", "B"), ("KindList", "C"), ("TagList", "D"), ("ScopeList", "F"), ("YesNoList", "I"),
                    ("RuleScopeList", "J"), ("SchoolList", "K"), ("WeekList", "L"), ("DayList", "M"), ("StrengthList", "N"), ("PeriodList", "O")):
        B.name(nm, B.dyn("Lists", col, 4, 100))
    TAGT, BANDT = "Lists!$D$4:$E$60", "Lists!$G$4:$H$60"

    # ---------------- 1 Venues
    vw = B.sheet(V, "Step 1 · Venues: every room", ["One row per room. Kind decides which subjects may use it (see Subjects).", "Blue on yellow = type or pick. Grey = worked out for you."],
                 [("Room", 16), ("Kind", 16), ("Capacity", 10), ("Building", 12), ("Floor", 8), ("Home room of", 12), ("Note", 34)], {1, 2, 3, 4, 5, 7})
    for i, x in enumerate(spec["venues"]):
        for j, k in enumerate(["room", "kind", "capacity", "building", "floor", None, "note"], start=1):
            if k: vw.cell(R0 + i, j, x.get(k))
    B.paint(vw, LAST(nV), {1, 2, 3, 4, 5, 7}, 7)
    for r in range(R0, LAST(nV) + 1): vw.cell(r, 6, f"=IF(A{r}=\"\",\"\",COUNTIF('{C}'!$D${R0}:$D${R0 + 999},A{r}))")
    B.dv(vw, f"B{R0}:B{LAST(nV)}", "=KindList", "Pick a room kind (add kinds on the Lists sheet).")
    B.name("RoomList", B.dyn(V, "A"))

    # ---------------- 2 Classes
    cw = B.sheet(C, "Step 2 · Classes", ["One row per form class. Home room is picked from the Venues list.", "Blue on yellow = type or pick. Grey = worked out for you."],
                 [("Class", 10), ("Level", 9), ("Students", 10), ("Home room", 12), ("Lessons listed", 13), ("Note", 34)], {1, 2, 3, 4, 6})
    for i, x in enumerate(spec["classes"]):
        for j, k in enumerate(["class", "level", "size", "home_room", None, "note"], start=1):
            if k: cw.cell(R0 + i, j, x.get(k))
    B.paint(cw, LAST(nC), {1, 2, 3, 4, 6}, 6)
    lcls = f"'{L}'!$D${R0}:${CL(3 + NCLS)}${R0 + 1999}"
    for r in range(R0, LAST(nC) + 1): cw.cell(r, 5, f'=IF(A{r}="","",SUMPRODUCT(({lcls}=A{r})*1))')
    B.dv(cw, f"B{R0}:B{LAST(nC)}", "=LevelList", "Pick a level."); B.dv(cw, f"D{R0}:D{LAST(nC)}", "=RoomList", "Pick a room from step 1.")
    B.name("ClassList", B.dyn(C, "A"))

    # ---------------- 3 Subjects
    scol = [("Subject", 36), ("Code", 7), ("Level tag", 9), ("Department", 11), ("Room kind (blank = home room)", 16), ("Lesson length (periods, blank = standard)", 12)] + \
           [(lv, 7) for lv in levels] + [("Lesson rows", 10), ("Note", 30)]
    nlv = len(levels); FL = 7; LL = 6 + nlv
    SIN = set(range(1, 7 + nlv)) | {8 + nlv}
    sw = B.sheet(S, "Step 3 · Subjects: periods a cycle at each level", [f"Periods per {'fortnight' if int(sc['weeks']) == 2 else 'week'} for each level that takes the subject (blank = not offered).",
                 f"Lessons are {sc['lesson_length']} periods long unless the subject says otherwise; leftover periods become one or two shorter lessons. Code names the groups.",
                 "Level tag: the level of study (e.g. G3 or Standard) used for banding; '-' if none."], scol, SIN)
    for i, x in enumerate(spec["subjects"]):
        rr = R0 + i
        for j, k in enumerate(["subject", "code", "tag", "dept", "room_kind", "lesson_length"], start=1): sw.cell(rr, j, x.get(k))
        for k, lv in enumerate(levels): sw.cell(rr, FL + k, (x.get("periods") or {}).get(lv))
        sw.cell(rr, LL + 2, x.get("note"))
    B.paint(sw, LAST(nS), SIN, LL + 2)
    for r in range(R0, LAST(nS) + 1): sw.cell(r, LL + 1, f"=IF(A{r}=\"\",\"\",COUNTIF('{L}'!$B${R0}:$B${R0 + 1999},A{r}))")
    B.dv(sw, f"C{R0}:C{LAST(nS)}", "=TagList", "Pick a level tag."); B.dv(sw, f"D{R0}:D{LAST(nS)}", "=DeptList", "Pick a department.")
    B.dv(sw, f"E{R0}:E{LAST(nS)}", "=KindList", "Pick a room kind.")
    B.name("SubjectList", B.dyn(S, "A"))
    SA = f"'{S}'!$A${R0}:$A${R0 + 999}"
    SLV = f"'{S}'!${CL(FL)}${R0}:${CL(LL)}${R0 + 999}"; SLH = f"'{S}'!${CL(FL)}${HR}:${CL(LL)}${HR}"

    # ---------------- 4 Teachers
    tcols = [("Teacher", 28), ("Short", 12), ("Department", 11), ("Allowance (periods)", 11), ("Provisional", 10)]
    for p in range(1, NPAIR + 1): tcols += [(f"Can teach {p}", 28), (f"at {p}", 9)]
    LO = 6 + 2 * NPAIR
    tcols += [("Periods given", 10), ("Spare", 8), ("(key, used by checks)", 12), ("Note", 30)]
    TIN = {1, 2, 3, 4, 5} | set(range(6, LO)) | {LO + 3}
    tw = B.sheet(T, "Step 4 · Teachers: who they are and what they can teach", ["Pick each subject a teacher can teach and the levels ('All', a level or a band). Lessons are checked against this.",
                 "Periods given adds every lesson the teacher has on step 6; Spare = Allowance - Periods given (red when over).", "Blue on yellow = type or pick. Grey = worked out for you."], tcols, TIN)
    for i, x in enumerate(spec["teachers"]):
        rr = R0 + i
        for j, k in enumerate(["name", "short", "dept", "allowance", "provisional"], start=1): tw.cell(rr, j, x.get(k))
        for p, ct in enumerate((x.get("can_teach") or [])[:NPAIR]):
            tw.cell(rr, 6 + 2 * p, ct.get("subject")); tw.cell(rr, 7 + 2 * p, ct.get("levels") or "All")
        tw.cell(rr, LO + 3, x.get("note"))
    B.paint(tw, LAST(nT), TIN, LO + 3)
    LQ = f"'{L}'"
    TC = [CL(4 + NCLS + k) for k in range(NTCH)]           # teacher columns on Lessons
    COL = {}
    names = ["Level", "Subject", "Block"] + [f"Class {i}" for i in range(1, NCLS + 1)] + [f"Teacher {i}" for i in range(1, NTCH + 1)] + \
            ["Students", "Periods (only if different)", "Code (only to keep an old one)", "Group code", "Periods", "Lesson length", "Full lessons", "Short lessons",
             "Room kind", "Teacher periods", "Check", "Note"]
    for j, n in enumerate(names, start=1): COL[n] = CL(j)
    TPC = COL["Teacher periods"]
    for r in range(R0, LAST(nT) + 1):
        load = "+".join(f"SUMIF({LQ}!${c}${R0}:${c}${R0 + 1999},A{r},{LQ}!${TPC}${R0}:${TPC}${R0 + 1999})" for c in TC)
        tw.cell(r, LO, f'=IF(A{r}="","",{load})')
        tw.cell(r, LO + 1, f'=IF(OR(A{r}="",D{r}=""),"",D{r}-{CL(LO)}{r})')
        key = "&".join(f'"|"&{CL(6 + 2 * p)}{r}&"@"&IF({CL(7 + 2 * p)}{r}="","All",{CL(7 + 2 * p)}{r})' for p in range(NPAIR)) + '&"|"'
        tw.cell(r, LO + 2, f'=IF(A{r}="","",{key})'); tw.cell(r, LO + 2).font = Font(name="Arial", size=7, color="8C8C8C")
    tw.conditional_formatting.add(f"{CL(LO + 1)}{R0}:{CL(LO + 1)}{LAST(nT)}", CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor="F4B6B6")))
    B.dv(tw, f"C{R0}:C{LAST(nT)}", "=DeptList", "Pick a department."); B.dv(tw, f"E{R0}:E{LAST(nT)}", "=YesNoList", "Yes or No.")
    for p in range(NPAIR):
        B.dv(tw, f"{CL(6 + 2 * p)}{R0}:{CL(6 + 2 * p)}{LAST(nT)}", "=SubjectList", "Pick a subject from step 3.")
        B.dv(tw, f"{CL(7 + 2 * p)}{R0}:{CL(7 + 2 * p)}{LAST(nT)}", "=ScopeList", "Pick All, a level or a band.")
    B.name("TeacherList", B.dyn(T, "A"))
    KEYC = CL(LO + 2)

    # ---------------- 5 Blocks
    bcols = [("Block", 40), ("Level", 8)] + [(f"Class {i}", 7) for i in range(1, NCLS + 1)] + \
            [("Groups in this block (from step 6)", 46), ("Groups", 8), ("Students in groups", 10), ("Students in classes", 10), ("Difference", 10), ("Longest (periods)", 10), ("Note", 30)]
    G0 = 3 + NCLS
    BIN = set(range(1, G0)) | {G0 + 6}
    bw = B.sheet(BL, "Step 5 · Blocks: groups that run at the same time", ["A block takes the students of its classes and splits them into groups that all have lessons at the same time (e.g. Sec 1 English for 1A-1D).",
                 "Name the block and pick its classes; then on step 6 give each group the block's name. The groups list and counts fill in by themselves.",
                 "Difference = students in groups - students in classes: 0 means everyone placed once; negative is normal for option blocks."], bcols, BIN)
    for i, x in enumerate(spec["blocks"]):
        rr = R0 + i; bw.cell(rr, 1, x["block"]); bw.cell(rr, 2, x["level"])
        for k, c in enumerate(x["classes"][:NCLS]): bw.cell(rr, 3 + k, c)
        bw.cell(rr, G0 + 6, x.get("note"))
    B.paint(bw, LAST(nB), BIN, G0 + 6)
    lc, lv_, ls_, lp = (f"{LQ}!${COL[n]}${R0}:${COL[n]}${R0 + 1999}" for n in ("Block", "Group code", "Students", "Periods"))
    for r in range(R0, LAST(nB) + 1):
        g = CL(G0)
        bw[f"{g}{r}"] = ArrayFormula(f"{g}{r}", f'=IF(A{r}="","",_xlfn.TEXTJOIN(", ",TRUE,IF({lc}=A{r},{lv_},"")))')
        bw.cell(r, G0 + 1, f'=IF(A{r}="","",COUNTIF({lc},A{r}))')
        bw.cell(r, G0 + 2, f'=IF(A{r}="","",SUMIF({lc},A{r},{ls_}))')
        bw.cell(r, G0 + 3, f"=IF(A{r}=\"\",\"\",SUMPRODUCT(SUMIF('{C}'!$A${R0}:$A${R0 + 999},C{r}:{CL(2 + NCLS)}{r},'{C}'!$C${R0}:$C${R0 + 999})))")
        bw.cell(r, G0 + 4, f'=IF(A{r}="","",{CL(G0 + 2)}{r}-{CL(G0 + 3)}{r})')
        bw.cell(r, G0 + 5, f'=IF(A{r}="","",_xlfn.MAXIFS({lp},{lc},A{r}))')
        for j in range(G0, G0 + 6): bw.cell(r, j).font = BF; bw.cell(r, j).fill = CF
        bw.cell(r, G0).alignment = Alignment(wrap_text=True, vertical="top")
    dcol = CL(G0 + 4)
    bw.conditional_formatting.add(f"{dcol}{R0}:{dcol}{LAST(nB)}", CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor="FCE4CC")))
    bw.conditional_formatting.add(f"{dcol}{R0}:{dcol}{LAST(nB)}", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="F4B6B6")))
    B.dv(bw, f"B{R0}:B{LAST(nB)}", "=LevelList", "Pick a level."); B.dv(bw, f"C{R0}:{CL(2 + NCLS)}{LAST(nB)}", "=ClassList", "Pick a class from step 2.")
    B.name("BlockList", B.dyn(BL, "A"))

    # ---------------- 6 Lessons
    widths = {"Level": 8, "Subject": 32, "Block": 32, "Students": 9, "Periods (only if different)": 11, "Code (only to keep an old one)": 11, "Group code": 11,
              "Periods": 8, "Lesson length": 8, "Full lessons": 8, "Short lessons": 8, "Room kind": 13, "Teacher periods": 9, "Check": 40, "Note": 28}
    lcols = [(n, widths.get(n, 7 if n.startswith("Class") else 20)) for n in names]
    lcols[2] = ("Block (empty = whole class)", 32)
    LIN = set(range(1, 4 + NCLS + NTCH + 3)) | {len(names)}
    lw = B.sheet(L, "Step 6 · Lessons: one row per group (or per class for whole-class subjects)",
                 ["Pick level, subject and block, then the class(es) and teacher(s). Leave Block empty for a whole-class lesson.",
                  "Group code, periods, lesson pattern and room kind fill in from steps 3-5. Check flags a teacher not set to teach the subject, a teacher twice in a block, missing periods or a repeated code."], lcols, LIN)
    for i, x in enumerate(spec["lessons"]):
        rr = R0 + i
        lw.cell(rr, 1, x["level"]); lw.cell(rr, 2, x["subject"]); lw.cell(rr, 3, x.get("block") or None)
        for k, c in enumerate(x["classes"][:NCLS]): lw.cell(rr, 4 + k, c)
        for k, t in enumerate((x.get("teachers") or [])[:NTCH]): lw.cell(rr, 4 + NCLS + k, t)
        c0 = 4 + NCLS + NTCH
        lw.cell(rr, c0, x.get("students")); lw.cell(rr, c0 + 1, x.get("periods_override")); lw.cell(rr, c0 + 2, x.get("code_override")); lw.cell(rr, len(names), x.get("note"))
    nLast = LAST(nL) + 100
    B.paint(lw, nLast, LIN, len(names))
    c_ = COL; CF1, CFN = c_["Class 1"], c_[f"Class {NCLS}"]
    def idx(col, r): return f"INDEX('{S}'!${col}${R0}:${col}${R0 + 999},MATCH(B{r},{SA},0))"
    for r in range(R0, nLast + 1):
        A_, B_, C_ = f"A{r}", f"B{r}", f"C{r}"
        PO, CO, GC, P_, LL_, FU, SH, RK, TP, CK = (f"{c_[n]}{r}" for n in ("Periods (only if different)", "Code (only to keep an old one)", "Group code", "Periods",
                                                                          "Lesson length", "Full lessons", "Short lessons", "Room kind", "Teacher periods", "Check"))
        seq = f'COUNTIFS($A${R0}:A{r},A{r},$B${R0}:B{r},B{r},$C${R0}:C{r},"<>")'
        code = f'RIGHT({A_},1)&IFERROR({idx("B", r)},"??")&IFERROR(VLOOKUP({idx("C", r)},{TAGT},2,FALSE),"")&{seq}'
        lw[GC] = f'=IF({B_}="","",IF({CO}<>"",{CO},IF({C_}="","Entire Class",{code})))'
        lw[P_] = f'=IF({B_}="","",IF({PO}<>"",{PO},IFERROR(INDEX({SLV},MATCH({B_},{SA},0),MATCH({A_},{SLH},0)),"")))'
        lw[LL_] = f'=IF({B_}="","",IFERROR(IF({idx("F", r)}="",Settings!$B$12,{idx("F", r)}),Settings!$B$12))'
        lw[SH] = f'=IF(N({P_})=0,"",IF({LL_}<=1,0,IF(MOD(-{P_},{LL_})*({LL_}-1)>{P_},0,MOD(-{P_},{LL_}))))'
        lw[FU] = f'=IF(N({P_})=0,"",({P_}-{SH}*({LL_}-1))/{LL_})'
        lw[RK] = f'=IF({B_}="","",IFERROR(IF({idx("E", r)}="","home room",{idx("E", r)}),""))'
        lw[TP] = f'=IF(N({P_})=0,0,{P_}*IF({C_}="",MAX(1,COUNTA({CF1}{r}:{CFN}{r})),1))'
        def can(tc):
            k = f"IFERROR(INDEX('{T}'!${KEYC}${R0}:${KEYC}${R0 + 999},MATCH({tc}{r},'{T}'!$A${R0}:$A${R0 + 999},0)),\"\")"
            band = f"IFERROR(VLOOKUP({A_},{BANDT},2,FALSE),\"\")"
            ok = f'ISNUMBER(FIND("|"&{B_}&"@All|",{k}))+ISNUMBER(FIND("|"&{B_}&"@"&{A_}&"|",{k}))+ISNUMBER(FIND("|"&{B_}&"@"&{band}&"|",{k}))'
            twice = "+".join(f"COUNTIFS($C${R0}:$C${R0 + 1999},{C_},${c}${R0}:${c}${R0 + 1999},{tc}{r})" for c in TC)
            return f'IF({tc}{r}="","",IF({ok}=0,{tc}{r}&" not set to teach this at "&{A_}&". ","")&IF(AND({C_}<>"",{twice}>1),{tc}{r}&" twice in this block. ",""))'
        lw[CK] = (f'=IF({B_}="","",TRIM(IF(N({P_})=0,"No periods for "&{A_}&" on Subjects. ","")&IF(COUNTA({CF1}{r}:{CFN}{r})=0,"No class. ","")'
                  f'&IF(COUNTA({TC[0]}{r}:{TC[-1]}{r})=0,"No teacher. ","")&' + "&".join(can(tc) for tc in TC) +
                  f'&IF(AND({GC}<>"Entire Class",COUNTIF(${c_["Group code"]}${R0}:${c_["Group code"]}${R0 + 1999},{GC})>1),"Code used twice. ","")'
                  f"&IF(AND({C_}<>\"\",COUNTIF('{BL}'!$A${R0}:$A${R0 + 999},{C_})=0),\"Block not on step 5. \",\"\")))")
        lw[CK].alignment = Alignment(wrap_text=True, vertical="top")
    ck = c_["Check"]
    lw.conditional_formatting.add(f"{ck}{R0}:{ck}{nLast}", FormulaRule(formula=[f"LEN({ck}{R0})>0"], fill=PatternFill("solid", fgColor="F4B6B6")))
    B.dv(lw, f"A{R0}:A{nLast}", "=LevelList", "Pick a level."); B.dv(lw, f"B{R0}:B{nLast}", "=SubjectList", "Pick a subject from step 3.")
    B.dv(lw, f"C{R0}:C{nLast}", "=BlockList", "Pick a block from step 5, or leave empty for a whole-class lesson.")
    B.dv(lw, f"{CF1}{R0}:{CFN}{nLast}", "=ClassList", "Pick a class from step 2."); B.dv(lw, f"{TC[0]}{R0}:{TC[-1]}{nLast}", "=TeacherList", "Pick a teacher from step 4.")
    lw.auto_filter.ref = f"A{HR}:{CL(len(names))}{R0 + max(nL, 1) - 1}"

    # ---------------- 7 Rules + Rule types
    rt = wb.create_sheet("Rule types")
    rt["A1"] = "Rule types you can use on 7 Rules"; rt["A1"].font = TITLE
    for j, h in enumerate(["Rule", "Can apply to", "Needs From/To", "Needs Number", "What it means", "Example"], start=1):
        c = rt.cell(3, j, h); c.font = HF; c.fill = PatternFill("solid", fgColor=NAVY)
    for i, row in enumerate(RULE_TYPES, start=4):
        for j, v in enumerate(row, start=1): c = rt.cell(i, j, v); c.font = BF; c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEF", (26, 30, 10, 10, 60, 44)): rt.column_dimensions[col].width = w
    nRT = len(RULE_TYPES); B.name("RuleList", f"'Rule types'!$A$4:$A${3 + nRT}")
    rcols = [("Applies to (kind)", 12), ("Applies to", 30), ("Subject (optional)", 28), ("Rule", 26), ("Week", 8), ("Day", 7), ("From period", 8), ("To period", 8),
             ("Number", 8), ("Must / Prefer", 9), ("Note", 40), ("Times", 20), ("Check", 36)]
    RIN = set(range(1, 12))
    rw = B.sheet(RU, "Step 7 · Rules: constraints on when things can happen", ["First pick what a rule applies to; the next drop-down then lists only those. Periods are numbered as on Period times (Times shows the clock).",
                 "Must = never broken; Prefer = kept where possible. See 'Rule types' for every rule and an example.", "Grey italic rows are examples: replace or delete them."], rcols, RIN)
    for i, x in enumerate(spec.get("rules", [])):
        for j, k in enumerate(["scope", "target", "subject", "rule", "week", "day", "from", "to", "number", "strength", "note"], start=1): rw.cell(R0 + i, j, x.get(k))
    rL = LAST(nR) + 80
    B.paint(rw, rL, RIN, 13)
    for i, x in enumerate(spec.get("rules", [])):
        if str(x.get("note") or "").startswith("example"):
            for j in range(1, 12): rw.cell(R0 + i, j).font = EXF
    SC = "{" + ",".join(f'"{s}"' for s in SCOPES) + "}"
    LISTS = "SchoolList,LevelList,ClassList,TeacherList,DeptList,SubjectList,BlockList,RoomList"
    rtA = f"'Rule types'!$A$4:$A${3 + nRT}"
    for r in range(R0, rL + 1):
        rw[f"L{r}"] = f'=IF(G{r}="","",TEXT(INDEX({PTB},G{r}+1),"h:mm AM/PM")&" - "&TEXT(INDEX({PTC},IF(H{r}="",G{r},H{r})+1),"h:mm AM/PM"))'
        nt, nn, scp = (f"INDEX('Rule types'!${c}$4:${c}${3 + nRT},MATCH(D{r},{rtA},0))" for c in ("C", "D", "B"))
        rw[f"M{r}"] = (f'=IF(A{r}="","",TRIM(IF(COUNTIF(CHOOSE(MATCH(A{r},{SC},0),{LISTS}),B{r})=0,"Pick the "&LOWER(A{r})&" from the list. ","")'
                       f'&IF(D{r}="","Pick a rule. ",IF(ISERROR(FIND(A{r},{scp})),D{r}&" does not apply to a "&LOWER(A{r})&". ","")'
                       f'&IF(AND(IFERROR({nt},"N")="Y",G{r}="",H{r}=""),"Give the period(s). ","")&IF(AND(IFERROR({nn},"N")="Y",I{r}=""),"Give the Number. ",""))'
                       f'&IF(AND(OR(D{r}="Same time across classes",D{r}="Not on the same day as"),C{r}=""),"Pick the Subject. ","")))')
        rw[f"M{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    rw.conditional_formatting.add(f"M{R0}:M{rL}", FormulaRule(formula=[f"LEN(M{R0})>0"], fill=PatternFill("solid", fgColor="F4B6B6")))
    B.dv(rw, f"A{R0}:A{rL}", "=RuleScopeList", "Pick what the rule applies to.")
    B.dv(rw, f"B{R0}:B{rL}", f"=CHOOSE(MATCH($A{R0},{SC},0),{LISTS})", "Pick from the list for the kind chosen in column A.")
    B.dv(rw, f"C{R0}:C{rL}", "=SubjectList", "Pick a subject."); B.dv(rw, f"D{R0}:D{rL}", "=RuleList", "Pick a rule (see Rule types).")
    B.dv(rw, f"E{R0}:E{rL}", "=WeekList", "Pick a week."); B.dv(rw, f"F{R0}:F{rL}", "=DayList", "All or a day.")
    B.dv(rw, f"G{R0}:H{rL}", "=PeriodList", "A period number (see Period times)."); B.dv(rw, f"J{R0}:J{rL}", "=StrengthList", "Must or Prefer.")

    # ---------------- Students (+ Combinations)
    stw = B.sheet("Students", "Students (optional)", ["With each student's class and groups (group codes from step 6, separated by commas), every student is checked for clashes.",
                  "Without it, students in a class are assumed to follow one level tag (e.g. all G3 or all Standard) across their banded subjects."],
                  [("Student ID", 12), ("Name", 30), ("Level", 8), ("Class", 9), ("Combination", 12), ("Groups (codes)", 60), ("Note", 30)], {1, 2, 3, 4, 5, 6, 7})
    for i, x in enumerate(spec.get("students", [])):
        for j, k in enumerate(["id", "name", "level", "class", "combination", "groups", "note"], start=1):
            v = x.get(k); stw.cell(R0 + i, j, ", ".join(v) if isinstance(v, list) else v)
    B.paint(stw, LAST(len(spec.get("students", []))), {1, 2, 3, 4, 5, 6, 7}, 7)
    B.dv(stw, f"C{R0}:C{R0 + 3000}", "=LevelList", "Pick a level."); B.dv(stw, f"D{R0}:D{R0 + 3000}", "=ClassList", "Pick a class.")
    order = ["README", "Settings", "Period times", "Day structure", V, C, S, T, BL, L, RU, "Rule types", "Students"]
    if spec.get("combinations"):
        cb = B.sheet("Combinations", "Subject combinations offered (optional)", ["Used to plan option blocks: subjects taken together must be in different blocks."],
                     [("Level", 8), ("Code", 8), ("Subjects", 70), ("Places", 8)], {1, 2, 3, 4})
        for i, x in enumerate(spec["combinations"]):
            for j, k in enumerate(["level", "code", "subjects", "places"], start=1):
                v = x.get(k); cb.cell(R0 + i, j, " + ".join(v) if isinstance(v, list) else v)
        B.paint(cb, LAST(len(spec["combinations"])), {1, 2, 3, 4}, 4)
        order.append("Combinations")
    order.append("Lists")
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)     # Excel works out every formula when the file is opened
    wb.save(path)
    return path
