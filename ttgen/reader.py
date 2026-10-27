"""Read a timetable setup workbook back into a spec, and work out what the formulas show (codes, periods, lesson pattern)
in Python so a file never opened in Excel still reads correctly."""
import datetime, re
import openpyxl
from .spec import defaults, lesson_split

HR, R0 = 5, 6

def _v(x):
    if isinstance(x, str):
        x = x.strip()
        return x or None
    return x

def _table(ws):
    hdr = [(_v(c.value) or "") for c in ws[HR]]
    out = []
    for row in ws.iter_rows(min_row=R0, values_only=True):
        d = {h: _v(v) for h, v in zip(hdr, row) if h}
        if any(v not in (None, "") for k, v in d.items() if not k.startswith(("Periods given", "Spare", "(key", "Lesson rows", "Home room of", "Lessons listed",
                                                                                 "Group", "Periods", "Lesson length", "Full", "Short", "Room kind", "Teacher periods",
                                                                                 "Check", "Times", "Students in", "Difference", "Longest"))):
            out.append(d)
    return hdr, out

def _first_col(row, prefix, n):
    return [row.get(f"{prefix} {i}") for i in range(1, n + 1) if row.get(f"{prefix} {i}")]

def read_template(path):
    wb = openpyxl.load_workbook(path, data_only=False)
    st = wb["Settings"]
    kv = {str(st.cell(r, 1).value).strip(): st.cell(r, 2).value for r in range(5, 40) if st.cell(r, 1).value}
    kind = str(kv.get("School type") or "secondary").lower()
    sc = defaults("primary" if kind.startswith("pri") else "secondary")
    t0 = kv.get("Period 0 starts at")
    if isinstance(t0, datetime.datetime): t0 = t0.time()
    sc.update(name=kv.get("School") or sc["name"], timetable=kv.get("Timetable") or sc["timetable"], weeks=int(kv.get("Weeks in cycle") or sc["weeks"]),
              period_minutes=int(kv.get("Period length (minutes)") or sc["period_minutes"]),
              start_time=t0.strftime("%H:%M") if isinstance(t0, datetime.time) else str(t0 or sc["start_time"]),
              periods_per_day=int(kv.get("Periods in a day") or sc["periods_per_day"]), lesson_length=int(kv.get("Standard lesson length (periods)") or sc["lesson_length"]))
    # lists
    ls = wb["Lists"]; heads = {str(ls.cell(3, j).value): j for j in range(1, ls.max_column + 1) if ls.cell(3, j).value}
    def col(name):
        j = heads.get(name); out = []
        if not j: return out
        for r in range(4, ls.max_row + 1):
            v = ls.cell(r, j).value
            if v not in (None, ""): out.append(v)
        return out
    if col("Levels"): sc["levels"] = [str(x) for x in col("Levels")]
    if col("Departments"): sc["departments"] = [str(x) for x in col("Departments")]
    if col("Room kinds"): sc["room_kinds"] = [str(x) for x in col("Room kinds")]
    tags, letters = col("Level tags"), [ls.cell(r, heads["Tag letter"]).value for r in range(4, 4 + len(col("Level tags")))] if "Tag letter" in heads else []
    if tags: sc["tags"] = [[str(a), str(b or "")] for a, b in zip(tags, letters + [""] * len(tags))]
    if col("Level") and col("Band"): sc["bands"] = dict(zip([str(x) for x in col("Level")], [str(x) for x in col("Band")]))
    # day structure
    ds = wb["Day structure"]; NP = sc["periods_per_day"]; day_structure = {}
    for r in range(6, ds.max_row + 1):
        lv, d = ds.cell(r, 1).value, ds.cell(r, 2).value
        if not lv or not d: continue
        codes = "".join((str(ds.cell(r, 3 + p).value or ".").strip()[:1].upper() or ".") for p in range(NP))
        day_structure.setdefault(str(lv), {})[str(d)] = codes
    _, ven = _table(wb["1 Venues"])
    venues = [dict(room=str(x["Room"]), kind=x.get("Kind"), capacity=x.get("Capacity"), building=x.get("Building"), floor=x.get("Floor"), note=x.get("Note")) for x in ven if x.get("Room")]
    _, cl = _table(wb["2 Classes"])
    classes = [{"class": str(x["Class"]), "level": x.get("Level"), "size": x.get("Students") or 0, "home_room": x.get("Home room"), "note": x.get("Note")} for x in cl if x.get("Class")]
    sh, su = _table(wb["3 Subjects"])
    lvcols = [h for h in sh if h in sc["levels"]]
    lenh = next((h for h in sh if h.startswith("Lesson length")), None); kindh = next((h for h in sh if h.startswith("Room kind")), None)
    subjects = [dict(subject=str(x["Subject"]), code=x.get("Code"), tag=x.get("Level tag") or "-", dept=x.get("Department"), room_kind=x.get(kindh),
                     lesson_length=x.get(lenh), periods={lv: x.get(lv) for lv in lvcols if x.get(lv)}, note=x.get("Note")) for x in su if x.get("Subject")]
    _, te = _table(wb["4 Teachers"])
    teachers = []
    for x in te:
        if not x.get("Teacher"): continue
        ct = [dict(subject=x.get(f"Can teach {p}"), levels=x.get(f"at {p}") or "All") for p in range(1, 9) if x.get(f"Can teach {p}")]
        teachers.append(dict(name=str(x["Teacher"]), short=x.get("Short"), dept=x.get("Department"), allowance=x.get("Allowance (periods)"),
                             provisional=x.get("Provisional"), can_teach=ct, note=x.get("Note")))
    _, bl = _table(wb["5 Blocks"])
    blocks = [dict(block=str(x["Block"]), level=x.get("Level"), classes=[str(c) for c in _first_col(x, "Class", 12)], note=x.get("Note")) for x in bl if x.get("Block")]
    lh, le = _table(wb["6 Lessons"])
    bh = next((h for h in lh if h.startswith("Block")), "Block")
    lessons = [dict(level=x.get("Level"), subject=x.get("Subject"), block=x.get(bh), classes=[str(c) for c in _first_col(x, "Class", 12)],
                    teachers=[str(t) for t in _first_col(x, "Teacher", 3)], students=x.get("Students"),
                    periods_override=x.get("Periods (only if different)"), code_override=x.get("Code (only to keep an old one)"), note=x.get("Note"))
               for x in le if x.get("Subject")]
    rules = []
    if "7 Rules" in wb.sheetnames:
        _, ru = _table(wb["7 Rules"])
        for x in ru:
            if not x.get("Applies to (kind)") or not x.get("Rule"): continue
            rules.append({"scope": x.get("Applies to (kind)"), "target": x.get("Applies to"), "subject": x.get("Subject (optional)"), "rule": x.get("Rule"),
                          "week": x.get("Week") or "Both", "day": x.get("Day") or "All", "from": x.get("From period"), "to": x.get("To period"),
                          "number": x.get("Number"), "strength": x.get("Must / Prefer") or "Must", "note": x.get("Note")})
    students = []
    if "Students" in wb.sheetnames:
        _, sx = _table(wb["Students"])
        for x in sx:
            if not (x.get("Student ID") or x.get("Name")): continue
            g = x.get("Groups (codes)")
            students.append(dict(id=x.get("Student ID"), name=x.get("Name"), level=x.get("Level"), **{"class": x.get("Class")}, combination=x.get("Combination"),
                                 groups=[s.strip() for s in str(g).split(",") if s.strip()] if g else []))
    return dict(school=sc, day_structure=day_structure, venues=venues, classes=classes, subjects=subjects, teachers=teachers,
                blocks=blocks, lessons=lessons, rules=rules, students=students)

def derive(spec):
    """add code, periods, lesson pattern, room kind to every lesson (same logic as the workbook formulas); returns problems found"""
    sc = spec["school"]; subj = {s["subject"]: s for s in spec["subjects"]}; letter = {a: b for a, b in sc["tags"]}
    seq = {}; problems = []
    for i, x in enumerate(spec["lessons"]):
        s = subj.get(x["subject"])
        if s is None: problems.append(f"Lesson row {i + 6}: subject '{x['subject']}' is not on Subjects."); continue
        p = x.get("periods_override") or (s.get("periods") or {}).get(x["level"])
        L = int(s.get("lesson_length") or sc["lesson_length"])
        full, short, L = lesson_split(p, L)
        if x.get("block"):
            k = (x["level"], x["subject"]); seq[k] = seq.get(k, 0) + 1
        else:
            seq.setdefault((x["level"], x["subject"]), seq.get((x["level"], x["subject"]), 0))
        if x.get("code_override"): code = str(x["code_override"])
        elif not x.get("block"): code = "Entire Class"
        else: code = f"{str(x['level'])[-1]}{s.get('code') or '??'}{letter.get(s.get('tag'), '')}{seq[(x['level'], x['subject'])]}"
        x.update(code=code, periods=int(p or 0), full=full, short=short, length=L, room_kind=s.get("room_kind") or ("classroom" if x.get("block") or x.get("code_override") else None),
                 tag=s.get("tag") or "-", dept=s.get("dept"))
        if not p: problems.append(f"Lesson row {i + 6}: {x['subject']} has no periods for {x['level']}.")
        if not x["classes"]: problems.append(f"Lesson row {i + 6}: no class.")
        if not x["teachers"]: problems.append(f"Lesson row {i + 6}: no teacher ({x['subject']} {', '.join(x['classes'])}).")
    return problems
