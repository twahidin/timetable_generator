"""Outputs for a solved timetable, built from one model so they always agree:

  * one combined Excel workbook (class, teacher, room and student grids, lesson list, teacher load)
  * one searchable HTML page (works offline; search, filters, free teachers and rooms, student timetable builder)
  * PDFs: one page per class, teacher, room and student, as printable books with bookmarks and as single files

The model comes from the solver result in memory (model_from_result) or from the setup workbook plus a timetable
workbook whose Lessons sheet uses LESSON_HEADERS (model_from_files), so a timetable made another way exports too."""
import collections, datetime, json, os, re, zipfile
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.comments import Comment
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.merge import MergedCellRange
from .reader import read_template, derive
from .report import LESSON_HEADERS, short_subject, clock
from .spec import periods_by_day, trim_day

ENTIRE = "Entire Class"
BREAK = {"F": "Assembly", "R": "Recess", "U": "Lunch"}
# (tint for the lesson, accent for its border and swatch); one per department, in the order departments appear
PALETTE = [("DCEBFB", "2F6FB5"), ("FDE6D2", "C2610F"), ("DCF1E1", "2E8B4E"), ("F9DCE3", "B23A5A"), ("FBEFC6", "9A7A0A"),
           ("E7DFF7", "6A4FB3"), ("D3EFEE", "1F8380"), ("EEE2D7", "87573A"), ("E2E7EE", "4A5A70"), ("F7DDEF", "A0408A"),
           ("E5F2D0", "58841C"), ("D6ECF7", "1B76A0"), ("F3E3D3", "9C5B1E"), ("E0E4F8", "4453A8")]
EXT = ("FFF1CC", "B86E00")
KINDS = ("classes", "teachers", "rooms", "students")


# ----------------------------------------------------------------------------------------------- names and labels
def abbrev(s):
    s = short_subject(str(s or ""))
    s = re.sub(r"Humanities \(Geography\)", "Hum Geog", s); s = re.sub(r"Humanities \(History\)", "Hum Hist", s)
    s = re.sub(r"^Geography\b", "Geog", s); s = re.sub(r"^History\b", "Hist", s)
    return s

def untagged(s):
    return re.sub(r"\s+(G[123]|Std|Fdn)$", "", abbrev(s))

def compact_classes(classes):
    cs = sorted(classes, key=lambda c: (len(c), c))
    if len(cs) <= 2: return ", ".join(cs)
    ms = [re.match(r"^(.*?)([A-Z])$", c) for c in cs]
    if all(ms) and len({x.group(1) for x in ms}) == 1:
        o = [ord(x.group(2)) for x in ms]
        if o == list(range(o[0], o[0] + len(o))): return f"{cs[0]}–{cs[-1]}"
    return ", ".join(cs)

def block_label(block, level):
    s = str(block or "")
    if level and s.lower().startswith(str(level).lower() + " "): s = s[len(str(level)) + 1:]
    return s

def safe_name(s):
    return re.sub(r'[\\/:*?"<>|]+', "-", str(s)).strip(" .") or "untitled"

def _week_day(name):
    parts = str(name).split()
    return (parts[0], parts[-1]) if len(parts) > 1 else (None, parts[0] if parts else "")


# ----------------------------------------------------------------------------------------------- the model
def _model(spec, days, lessons, notes=(), status=None):
    sc = spec["school"]; NP = int(sc["periods_per_day"]); ND = len(days); ds = spec["day_structure"]
    levels = [str(x) for x in sc["levels"]]
    keep = periods_by_day(sc, days)   # Settings 'Lessons end by (...)': later periods are no school on those days
    grid = {lv: [trim_day((ds.get(lv, {}).get(d) or "." * NP).ljust(NP, ".")[:NP], keep.get(d)) for d in days] for lv in levels}
    times = [[clock(sc, p), clock(sc, p + 1)] for p in range(NP)]
    classes = [dict(name=str(c["class"]), level=str(c.get("level") or ""), size=c.get("size") or "", home=c.get("home_room") or "") for c in spec["classes"]]
    seen = {c["name"] for c in classes}
    for L in lessons:
        for c in L["classes"]:
            if c not in seen: seen.add(c); classes.append(dict(name=c, level=L["level"], size="", home=""))
    teachers = [dict(name=str(t["name"]), short=str(t.get("short") or t["name"]), dept=t.get("dept") or "", allowance=t.get("allowance") or "") for t in spec["teachers"]]
    seen = {t["name"] for t in teachers}
    for L in lessons:
        for t in L["teachers"]:
            if t not in seen: seen.add(t); teachers.append(dict(name=t, short=t, dept="", allowance=""))
    rooms = [dict(name=str(v["room"]), kind=v.get("kind") or "classroom", capacity=v.get("capacity") or "") for v in spec["venues"]]
    seen = {r["name"] for r in rooms}
    for L in lessons:
        if L["room"] and L["room"] not in seen: seen.add(L["room"]); rooms.append(dict(name=L["room"], kind="to be confirmed" if L["room"] == "TBC" else "other", capacity=""))
    used_rooms = {L["room"] for L in lessons if L["room"]}
    rooms = [r for r in rooms if r["name"] in used_rooms] + [r for r in rooms if r["name"] not in used_rooms]
    depts = []
    for s in spec["subjects"]:
        k = s.get("dept") or untagged(s["subject"])
        if k not in depts: depts.append(k)
    subjects = {}
    for s in spec["subjects"]:
        k = s.get("dept") or untagged(s["subject"])
        subjects[str(s["subject"])] = dict(short=abbrev(s["subject"]), dept=s.get("dept") or "", color=depts.index(k) % len(PALETTE))
    for L in lessons:
        if L["subject"] not in subjects:
            k = untagged(L["subject"])
            if k not in depts: depts.append(k)
            subjects[L["subject"]] = dict(short=abbrev(L["subject"]), dept="", color=depts.index(k) % len(PALETTE))
    students = [dict(id=str(s.get("id") or s.get("name")), name=str(s.get("name") or s.get("id")), cls=str(s.get("class") or ""),
                     level=str(s.get("level") or ""), groups=[str(g) for g in s.get("groups") or []]) for s in spec.get("students", [])]
    cap = {r["name"]: 1 for r in rooms}
    for r in spec.get("rules", []):
        if r.get("rule") == "At most N at once" and r.get("scope") == "Room" and r.get("target") in cap and r.get("number"):
            cap[r["target"]] = int(r["number"])
    planned = {}
    for x in spec["lessons"]:
        if not x.get("periods") or not x.get("classes"): continue
        if x.get("code") == ENTIRE:
            for c in x["classes"]: planned[("C", c, x["subject"])] = planned.get(("C", c, x["subject"]), 0) + int(x["periods"])
        else: planned[("G", x["code"])] = int(x["periods"])
    m = dict(school=dict(name=sc.get("name") or "School", timetable=sc.get("timetable") or "Timetable", weeks=int(sc.get("weeks") or 1),
                         period_minutes=int(sc["period_minutes"]), start_time=str(sc["start_time"]), kind=sc.get("kind") or ""),
             days=list(days), NP=NP, ND=ND, times=times, grid=grid, keep=keep, levels=levels, classes=classes, teachers=teachers, rooms=rooms,
             subjects=subjects, lessons=lessons, students=students, room_cap=cap, planned=planned, notes=list(notes or []), status=status,
             generated=datetime.date.today().strftime("%d %b %Y"))
    m["class_ix"] = {c["name"]: c for c in classes}; m["teacher_ix"] = {t["name"]: t for t in teachers}; m["room_ix"] = {r["name"]: r for r in rooms}
    m["checks"] = check_model(m)
    return m

def _lesson_order(res):
    return sorted(res["lessons"], key=lambda x: (x["start"], x["unit"]["level"], x["unit"]["name"]))

def model_from_result(res):
    """the solver result after report.assign_rooms (write_timetable does that)"""
    P = res["problem"]; NP = P.NP; out = []
    for i, x in enumerate(_lesson_order(res), start=1):
        for r in x["part"]:
            out.append(dict(id=i, d=x["start"] // NP, p=x["start"] % NP, n=int(x["ln"]), level=str(r["level"]),
                            block=x["unit"]["name"] if x["unit"]["kind"] == "block" else "",
                            classes=[str(c) for c in (x["unit"]["classes"] if x["unit"]["kind"] == "class" else r["classes"])],
                            subject=str(r["subject"]), group=str(r["code"]), teachers=[str(t) for t in r["teachers"]],
                            room=str(x.get("rooms", {}).get(id(r), "") or ""), ext=bool(x["ext"])))
    return _model(P.spec, P.days, out, notes=res.get("notes", []), status=res.get("status"))

def _split_list(v):
    return [s.strip() for s in str(v or "").split(",") if s.strip()]

def model_from_files(setup_path, timetable_path):
    spec = read_template(setup_path); derive(spec)
    sc = spec["school"]; NP = int(sc["periods_per_day"]); ds = spec["day_structure"]
    lv0 = next((lv for lv in sc["levels"] if lv in ds), None)
    if lv0 is None: raise ValueError("The setup workbook's Day structure is empty.")
    days = list(ds[lv0].keys()); dix = {d.lower(): i for i, d in enumerate(days)}
    wb = openpyxl.load_workbook(timetable_path, data_only=True)
    if "Lessons" not in wb.sheetnames: raise ValueError("The timetable workbook has no Lessons sheet.")
    rows = wb["Lessons"].iter_rows(values_only=True)
    hdr = [str(h).strip() if h is not None else "" for h in next(rows)]
    ix = {h.lower(): i for i, h in enumerate(hdr)}
    def g(row, *names):
        for n in names:
            i = ix.get(n.lower())
            if i is not None and i < len(row) and row[i] not in (None, ""): return row[i]
        return None
    out, ids, bad = [], {}, []
    for k, row in enumerate(rows, start=2):
        if not any(v not in (None, "") for v in row): continue
        day = str(g(row, "Day") or "").strip(); p = g(row, "Start period")
        if day.lower() not in dix or p is None: bad.append(k); continue
        p = int(p); n = g(row, "Periods")
        n = int(n) if n is not None else int(g(row, "End period") or p) - p + 1
        subject = str(g(row, "Subject") or ""); group = str(g(row, "Group") or ENTIRE); block = str(g(row, "Block") or "")
        lid = g(row, "Lesson ID")
        if lid is None: lid = ids.setdefault((dix[day.lower()], p, n, block or f"{group}|{g(row, 'Classes')}"), len(ids) + 1)
        out.append(dict(id=int(lid), d=dix[day.lower()], p=p, n=n, level=str(g(row, "Level") or ""), block=block,
                        classes=_split_list(g(row, "Classes")), subject=subject, group=group, teachers=_split_list(g(row, "Teachers")),
                        room=str(g(row, "Room") or ""), ext=str(g(row, "After school") or "").strip().lower() in ("yes", "y", "true", "1", "x")))
    notes, status = [], None
    if "Read me" in wb.sheetnames:
        sect = None
        for (v,) in wb["Read me"].iter_rows(min_col=1, max_col=1, values_only=True):
            if v is None: continue
            v = str(v)
            if v in ("Checks", "After-school lessons", "Rules checked after solving", "Notes", "How to read", "Contents"): sect = v; continue
            if "Solver status:" in v: status = v.split("Solver status:")[-1].strip(" .")
            if sect in ("Notes", "Rules checked after solving"): notes.append(v)
    if bad: notes.append(f"{len(bad)} rows on the Lessons sheet were skipped (unknown day or no start period): rows {', '.join(map(str, bad[:10]))}.")
    return _model(spec, days, out, notes=notes, status=status)


# ----------------------------------------------------------------------------------------------- independent checks
def check_model(m):
    NP, L = m["NP"], m["lessons"]; days = m["days"]; probs = []
    where = lambda d, q: f"{days[d]} period {q} ({m['times'][q][0]})"
    tb = collections.defaultdict(set)
    for x in L:
        for t in x["teachers"]:
            for q in range(x["p"], x["p"] + x["n"]): tb[(t, x["d"], q)].add(x["id"])
    def spans(keys):
        """(who, d, q) -> [(who, d, 'Odd Mon 10.30-11.30')] with consecutive periods joined"""
        by = collections.defaultdict(list)
        for w, d, q in keys: by[(w, d)].append(q)
        out = []
        for (w, d), qs in sorted(by.items()):
            qs = sorted(set(qs)); a = prev = qs[0]
            for q in qs[1:] + [None]:
                if q is not None and q == prev + 1: prev = q; continue
                out.append((w, d, f"{days[d]} {m['times'][a][0]}-{m['times'][prev][1]} (periods {a}-{prev})" if prev > a else where(d, a)))
                if q is not None: a = prev = q
        return out
    probs += [f"Teacher double-booked: {t} on {w}." for t, d, w in spans({k for k, v in tb.items() if len(v) > 1})][:30]
    rb = collections.Counter()
    for x in L:
        if x["room"] and x["room"] != "TBC":
            for q in range(x["p"], x["p"] + x["n"]): rb[(x["room"], x["d"], q)] += 1
    probs += [f"Room over capacity: {r} on {w} (holds {m['room_cap'].get(r, 1)} at once)." for r, d, w in spans({k for k, n in rb.items() if n > m["room_cap"].get(k[0], 1)})][:30]
    occ = collections.defaultdict(lambda: collections.defaultdict(set)); whole = collections.defaultdict(set)
    for x in L:
        for c in x["classes"]:
            for q in range(x["p"], x["p"] + x["n"]):
                occ[c][(x["d"], q)].add(x["id"])
                if x["group"] == ENTIRE: whole[(c, x["d"], q)].add(x["id"])
    side = 0; ccl = set()
    for c, slots in occ.items():
        for (d, q), ids in slots.items():
            if len(ids) > 1:
                if whole.get((c, d, q)): ccl.add((c, d, q))
                else: side += 1
    probs += [f"Class clash: {c} on {w}." for c, d, w in spans(ccl)][:30]
    for x in L:
        lv = x["level"] if x["level"] in m["grid"] else m["class_ix"].get((x["classes"] or [""])[0], {}).get("level")
        k = m.get("keep", {}).get(m["days"][x["d"]])
        if k is not None and x["p"] + x["n"] > k:   # Settings 'Lessons end by': runs past the time for its day
            probs.append(f"{x['subject']} {x['group']} ends after {m['times'][min(k, m['NP'] - 1)][0]} on {where(x['d'], x['p'])} (Settings: Lessons end by)."); continue
        if lv in m["grid"] and any(m["grid"][lv][x["d"]][q] not in "LA" for q in range(x["p"], min(m["NP"], x["p"] + x["n"]))):
            probs.append(f"{x['subject']} {x['group']} sits in a break or outside the day on {where(x['d'], x['p'])}.")
    placed = collections.Counter(); seen = set()
    for x in L:
        if x["group"] == ENTIRE:
            for c in x["classes"]:
                k = (x["id"], c)
                if k not in seen: seen.add(k); placed[("C", c, x["subject"])] += x["n"]
        else:
            k = (x["id"], x["group"])
            if k not in seen: seen.add(k); placed[("G", x["group"])] += x["n"]
    for k, want in m["planned"].items():
        if placed.get(k, 0) != want:
            what = f"{k[2]} for {k[1]}" if k[0] == "C" else f"group {k[1]}"
            probs.append(f"{what}: {placed.get(k, 0)} periods timetabled, {want} planned.")
    by_group = collections.defaultdict(set); whole_pos = collections.defaultdict(set)
    for x in L:
        pos = {(x["d"], q) for q in range(x["p"], x["p"] + x["n"])}
        if x["group"] == ENTIRE:
            for c in x["classes"]: whole_pos[c] |= pos
        else: by_group[x["group"]] |= pos
    sclash, unknown = [], collections.Counter()
    for s in m["students"]:
        seen_q = collections.Counter(whole_pos.get(s["cls"], ()))
        for gc in s["groups"]:
            if gc not in by_group: unknown[gc] += 1; continue
            seen_q.update(by_group[gc])
        if any(v > 1 for v in seen_q.values()): sclash.append(s)
    if sclash: probs.append(f"{len(sclash)} students have two lessons at the same time, e.g. {', '.join(s['id'] + ' ' + s['name'] for s in sclash[:5])}.")
    if unknown: probs.append(f"Students list group codes that have no lessons: {', '.join(sorted(unknown)[:12])}.")
    lids = collections.OrderedDict()
    for x in L:
        if x["ext"]: lids.setdefault(x["id"], []).append(x)
    after = []
    for lid, xs in lids.items():
        x = xs[0]; who = ", ".join(sorted({y["group"] if y["group"] != ENTIRE else ", ".join(y["classes"]) for y in xs}))
        after.append(f"{days[x['d']]} {m['times'][x['p']][0]}-{m['times'][x['p'] + x['n'] - 1][1]}: {x['block'] or x['subject']} ({who})")
    return dict(problems=probs, after_school=after, side_by_side=side, lessons=len({x["id"] for x in L}), rows=len(L),
                students=len(m["students"]), ok=not probs)


# ----------------------------------------------------------------------------------------------- views (one per class, teacher, room, student)
def _codes_for(m, levels):
    levels = [lv for lv in dict.fromkeys(levels) if lv in m["grid"]]
    if not levels: return None
    out = []
    for d in range(m["ND"]):
        row = []
        for p in range(m["NP"]):
            cs = {m["grid"][lv][d][p] for lv in levels}
            row.append(cs.pop() if len(cs) == 1 else "L" if cs & set("LA") else ("R" if "R" in cs else sorted(cs)[0]))
        out.append("".join(row))
    return out

def _by_id(rows):
    g = collections.OrderedDict()
    for x in rows: g.setdefault(x["id"], []).append(x)
    return g

def _common(vals):
    c = collections.Counter(v for v in vals if v)
    return ", ".join(v for v, _ in c.most_common(2)) if c else ""

def _tshort(m, names):
    return ", ".join(m["teacher_ix"].get(t, {}).get("short") or t for t in names)

def _color(m, subject):
    return m["subjects"].get(subject, {}).get("color", 0)

def class_view(m, c):
    info = m["class_ix"][c]; rows = [x for x in m["lessons"] if c in x["classes"]]
    bgroups = collections.defaultdict(set)
    for x in rows:
        if x["block"]: bgroups[x["block"]].add(x["group"])
    cells, used = [], set()
    for lid, xs in _by_id(rows).items():
        x0 = xs[0]; used |= {(x0["d"], q) for q in range(x0["p"], x0["p"] + x0["n"])}
        if len(xs) == 1 and x0["group"] == ENTIRE:
            title, lines = abbrev(x0["subject"]), [_tshort(m, x0["teachers"]), x0["room"]]
        elif x0["block"]:
            groups = {x["group"] for x in xs}
            if groups < bgroups[x0["block"]]:
                subs = list(dict.fromkeys(abbrev(x["subject"]) for x in xs))
                title = (" · ".join(subs) + " only") if len(subs) <= 2 else f"{block_label(x0['block'], info['level'])}: {len(groups)} of {len(bgroups[x0['block']])} groups"
                lines = [", ".join(sorted(groups))]
            else:
                subs = list(dict.fromkeys(untagged(x["subject"]) for x in xs))
                summary = " · ".join(subs) if len(subs) > 1 else " · ".join(dict.fromkeys(abbrev(x["subject"]).replace(subs[0], "").strip() or subs[0] for x in xs))
                title, lines = block_label(x0["block"], info["level"]) or summary, [summary, f"{len(xs)} groups"]
        else:
            title, lines = abbrev(x0["subject"]), [x0["group"], _tshort(m, x0["teachers"]), x0["room"]]
        cells.append(dict(d=x0["d"], p=x0["p"], n=x0["n"], title=title, lines=[l for l in lines if l], color=_color(m, x0["subject"]),
                          ext=any(x["ext"] for x in xs), id=lid))
    legend = []; per = collections.Counter(); ent = collections.OrderedDict()
    for lid, xs in _by_id(rows).items():
        for x in xs:
            k = (x["block"], x["group"], x["subject"]) if x["group"] != ENTIRE else ("", ENTIRE, x["subject"])
            if (lid, k) in per: continue
            per[(lid, k)] = 1; e = ent.setdefault(k, dict(x=x, periods=0, rooms=[], teachers=x["teachers"])); e["periods"] += x["n"]; e["rooms"].append(x["room"])
    blocks = collections.OrderedDict()
    for k, e in ent.items():
        x = e["x"]
        if k[1] == ENTIRE:
            legend.append(dict(color=_color(m, x["subject"]), label=abbrev(x["subject"]),
                               text=f"{x['subject']} · {', '.join(e['teachers']) or 'teacher to be named'} · {_common(e['rooms']) or '-'} · {e['periods']} periods"))
        else: blocks.setdefault(x["block"] or x["group"], []).append(e)
    for b, es in blocks.items():
        x = es[0]["x"]
        parts = [f"{abbrev(e['x']['subject'])} {e['x']['group']} ({_tshort(m, e['teachers'])}, {_common(e['rooms']) or '-'})" for e in es]
        legend.append(dict(color=_color(m, x["subject"]), label=block_label(b, info["level"]) or abbrev(x["subject"]),
                           text=f"{max(e['periods'] for e in es)} periods · " + "; ".join(parts)))
    nper = len(used)
    stats = [(info["level"], "level"), (str(info["size"] or "-"), "students"), (info["home"] or "-", "home room"), (str(nper), "periods a cycle")]
    sub = f"{info['level']} · {info['size']} students · home room {info['home'] or '-'}"
    return dict(kind="classes", key=c, kicker="Class", title=f"Class {c}", subtitle=sub, stats=stats, codes=_codes_for(m, [info["level"]]),
                cells=cells, legend=legend, legend_title="Subjects, teachers and rooms", group=info["level"])

def teacher_view(m, t):
    info = m["teacher_ix"][t]; rows = [x for x in m["lessons"] if t in x["teachers"]]
    cells, per_day = [], collections.Counter(); ent = collections.OrderedDict(); total = 0; ext = 0
    for lid, xs in _by_id(rows).items():
        x0 = xs[0]; who = ", ".join(dict.fromkeys(x["group"] if x["group"] != ENTIRE else ", ".join(x["classes"]) for x in xs))
        cls = sorted({c for x in xs for c in x["classes"]})
        lines = [who if x0["group"] == ENTIRE else f"{who} ({compact_classes(cls)})", ", ".join(dict.fromkeys(x["room"] for x in xs if x["room"]))]
        cells.append(dict(d=x0["d"], p=x0["p"], n=x0["n"], title=" + ".join(dict.fromkeys(abbrev(x["subject"]) for x in xs)), lines=[l for l in lines if l],
                          color=_color(m, x0["subject"]), ext=any(x["ext"] for x in xs), id=lid))
        total += x0["n"]; per_day[x0["d"]] += x0["n"]; ext += x0["n"] if x0["ext"] else 0
        for x in xs:
            k = (x["subject"], x["group"] if x["group"] != ENTIRE else ", ".join(x["classes"]))
            e = ent.setdefault(k, dict(x=x, periods=0, rooms=[])); e["periods"] += x["n"]; e["rooms"].append(x["room"])
    legend = []
    for (s, who), e in ent.items():
        x = e["x"]; extra = f" ({compact_classes(x['classes'])})" if x["group"] != ENTIRE else ""
        legend.append(dict(color=_color(m, s), label=abbrev(s), text=f"{who}{extra} · {e['periods']} periods · {_common(e['rooms']) or '-'}"))
    busiest = max(per_day.values()) if per_day else 0
    stats = [(str(total), "periods a cycle"), (str(info.get("allowance") or "-"), "allowance"), (str(busiest), "most in a day"), (str(len(per_day)), "days teaching")]
    sub = " · ".join(x for x in (f"Department {info['dept']}" if info["dept"] else "", f"short name {info['short']}" if info["short"] != t else "",
                                     f"{ext} after-school periods" if ext else "") if x)
    lv = [x["level"] for x in rows]
    return dict(kind="teachers", key=t, kicker="Teacher", title=t, subtitle=sub or "Teacher", stats=stats, codes=_codes_for(m, lv), cells=cells,
                legend=legend, legend_title="Teaching", group=info["dept"] or "Other")

def room_view(m, r):
    info = m["room_ix"][r]; rows = [x for x in m["lessons"] if x["room"] == r]
    cells, ent, used = [], collections.OrderedDict(), set()
    for lid, xs in _by_id(rows).items():
        x0 = xs[0]; used |= {(x0["d"], q) for q in range(x0["p"], x0["p"] + x0["n"])}
        who = ", ".join(dict.fromkeys(x["group"] if x["group"] != ENTIRE else ", ".join(x["classes"]) for x in xs))
        cells.append(dict(d=x0["d"], p=x0["p"], n=x0["n"], title=" + ".join(dict.fromkeys(abbrev(x["subject"]) for x in xs)),
                          lines=[who, _tshort(m, list(dict.fromkeys(t for x in xs for t in x["teachers"])))], color=_color(m, x0["subject"]),
                          ext=any(x["ext"] for x in xs), id=lid))
        for x in xs:
            k = (x["subject"], x["group"] if x["group"] != ENTIRE else ", ".join(x["classes"]))
            e = ent.setdefault(k, dict(x=x, periods=0)); e["periods"] += x["n"]
    legend = [dict(color=_color(m, s), label=abbrev(s), text=f"{who} · {_tshort(m, e['x']['teachers'])} · {e['periods']} periods") for (s, who), e in ent.items()]
    lv = [x["level"] for x in rows]; codes = _codes_for(m, lv)
    open_ = sum(ch in "L" for row in (codes or []) for ch in row) or 1
    stats = [(info["kind"], "kind"), (str(m["room_cap"].get(r, 1)), "groups at once"), (str(len(used)), "periods used"), (f"{round(100 * len(used) / open_)}%", "of lesson time")]
    return dict(kind="rooms", key=r, kicker="Room", title=r, subtitle=f"{info['kind'].capitalize()}" + (f" · holds {info['capacity']}" if info["capacity"] else ""),
                stats=stats, codes=codes, cells=cells, legend=legend, legend_title="Used by", group=info["kind"])

def student_rows(m, cls, groups):
    gs = set(groups)
    return [x for x in m["lessons"] if (x["group"] == ENTIRE and cls in x["classes"]) or (x["group"] != ENTIRE and x["group"] in gs)]

def student_view(m, s):
    rows = student_rows(m, s["cls"], s["groups"]); cells, ent, used = [], collections.OrderedDict(), set()
    for lid, xs in _by_id(rows).items():
        x0 = xs[0]; used |= {(x0["d"], q) for q in range(x0["p"], x0["p"] + x0["n"])}
        cells.append(dict(d=x0["d"], p=x0["p"], n=x0["n"], title=" + ".join(dict.fromkeys(abbrev(x["subject"]) for x in xs)),
                          lines=[_tshort(m, x0["teachers"]), x0["room"]], color=_color(m, x0["subject"]), ext=any(x["ext"] for x in xs), id=lid))
        for x in xs:
            e = ent.setdefault(x["subject"], dict(x=x, periods=0, rooms=[])); e["periods"] += x["n"]; e["rooms"].append(x["room"])
    legend = [dict(color=_color(m, s_), label=abbrev(s_), text=f"{s_} · {e['x']['group'] if e['x']['group'] != ENTIRE else 'with the class'} · "
                                                             f"{', '.join(e['x']['teachers']) or '-'} · {_common(e['rooms']) or '-'} · {e['periods']} periods")
              for s_, e in ent.items()]
    lv = m["class_ix"].get(s["cls"], {}).get("level") or s["level"]
    stats = [(s["cls"], "class"), (lv, "level"), (str(len(s["groups"])), "groups"), (str(len(used)), "periods a cycle")]
    return dict(kind="students", key=s["id"], kicker="Student", title=s["name"], subtitle=f"{s['id']} · Class {s['cls']} · {lv}",
                stats=stats, codes=_codes_for(m, [lv]), cells=cells, legend=legend, legend_title="Subjects, teachers and rooms", group=s["cls"])

def all_views(m, kind):
    if kind == "classes":
        for c in m["classes"]:
            if any(c["name"] in x["classes"] for x in m["lessons"]): yield class_view(m, c["name"])
    elif kind == "teachers":
        busy = {t for x in m["lessons"] for t in x["teachers"]}
        for t in sorted(m["teachers"], key=lambda t: (str(t["dept"]).lower(), t["name"].lower())):
            if t["name"] in busy: yield teacher_view(m, t["name"])
    elif kind == "rooms":
        busy = {x["room"] for x in m["lessons"]}
        for r in m["rooms"]:
            if r["name"] in busy: yield room_view(m, r["name"])
    elif kind == "students":
        order = {c["name"]: i for i, c in enumerate(m["classes"])}
        for s in sorted(m["students"], key=lambda s: (order.get(s["cls"], 999), s["id"])): yield student_view(m, s)


# ----------------------------------------------------------------------------------------------- PDF
_FONTS = None
def _fonts():
    global _FONTS
    if _FONTS: return _FONTS
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    here = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fonts")
    try:
        pdfmetrics.registerFont(TTFont("Carlito", os.path.join(here, "Carlito-Regular.ttf")))
        pdfmetrics.registerFont(TTFont("Carlito-Bold", os.path.join(here, "Carlito-Bold.ttf")))
        _FONTS = ("Carlito", "Carlito-Bold")
    except Exception:
        _FONTS = ("Helvetica", "Helvetica-Bold")
    return _FONTS

def _hex(c):
    from reportlab.lib.colors import HexColor
    return HexColor("#" + c)

INK, SUB, MUTED, LINE, FAINT = "1B2330", "4A5566", "7A8494", "E3E7EC", "F6F7F9"

def _fit(text, font, size, maxw, minsize=5.0):
    from reportlab.pdfbase.pdfmetrics import stringWidth
    s = size
    while s > minsize and stringWidth(text, font, s) > maxw: s -= 0.25
    if stringWidth(text, font, s) <= maxw: return text, s
    while text and stringWidth(text + "…", font, s) > maxw: text = text[:-1]
    return (text + "…") if text else "", s

def _wrap(text, font, size, width, first_width=None):
    from reportlab.pdfbase.pdfmetrics import stringWidth
    words, lines, cur = str(text).split(" "), [], ""
    w = first_width if first_width is not None else width
    for wd in words:
        t = (cur + " " + wd).strip()
        if stringWidth(t, font, size) <= w or not cur: cur = t
        else: lines.append(cur); cur = wd; w = width
    if cur: lines.append(cur)
    return lines

def _lanes(cells):
    """cells that overlap in time share the row height"""
    out = {}
    for d in {c["d"] for c in cells}:
        cs = sorted((c for c in cells if c["d"] == d), key=lambda c: (c["p"], -c["n"]))
        cluster, end = [], -1
        def flush(cl):
            lanes = []
            for c in cl:
                for i, e in enumerate(lanes):
                    if e <= c["p"]: lanes[i] = c["p"] + c["n"]; out[id(c)] = [i, 0]; break
                else: lanes.append(c["p"] + c["n"]); out[id(c)] = [len(lanes) - 1, 0]
            for c in cl: out[id(c)][1] = len(lanes)
        for c in cs:
            if cluster and c["p"] >= end: flush(cluster); cluster, end = [], -1
            cluster.append(c); end = max(end, c["p"] + c["n"])
        if cluster: flush(cluster)
    return out

def _legend_layout(v, R, B, width, max_h=150):
    from reportlab.pdfbase.pdfmetrics import stringWidth
    ents = v["legend"]; lay = None
    for size in (6.6, 6.2, 5.8, 5.4):
        lh = size * 1.28; ncol = 3 if len(ents) > 4 else 2
        cw = (width - (ncol - 1) * 14) / ncol
        items = []
        for e in ents:
            lw = stringWidth(e["label"], B, size) + 4
            lines = _wrap(e["text"], R, size, cw - 9, first_width=cw - 9 - lw)
            items.append((e, lw, lines, len(lines) * lh + 1.6))
        def pack(limit):
            cols, acc = [[]], 0
            for it in items:
                if cols[-1] and acc + it[3] > limit: cols.append([]); acc = 0
                cols[-1].append(it); acc += it[3]
            return cols
        limit = max([it[3] for it in items] + [sum(it[3] for it in items) / ncol])
        while len(pack(limit)) > ncol: limit += lh / 2
        cols = pack(limit); h = max(sum(it[3] for it in c) for c in cols) if items else 0
        lay = dict(size=size, lh=lh, cw=cw, cols=cols, h=h + 14, dropped=0)
        if h + 14 <= max_h: return lay
    # still too tall at the smallest size: keep what fits and say how many more there are
    room = max_h - 14 - lay["lh"]; kept, n = [], 0
    for c in lay["cols"]:
        acc, cc = 0, []
        for it in c:
            if acc + it[3] > room: n += 1; continue
            cc.append(it); acc += it[3]
        kept.append(cc)
    lay.update(cols=kept, h=max_h, dropped=n)
    return lay

def draw_page(c, m, v, page=None):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfbase.pdfmetrics import stringWidth
    R, B = _fonts(); W, H = landscape(A4); M = 26; NP, ND = m["NP"], m["ND"]; sc = m["school"]
    # header
    c.setFillColor(_hex(MUTED)); c.setFont(B, 7)
    c.drawString(M, H - 26, f"{sc['name'].upper()}  ·  {sc['timetable'].upper()}")
    c.setFont(R, 7); c.drawRightString(W - M, H - 26, f"Generated {m['generated']}")
    kick = v["kicker"].upper(); kw = stringWidth(kick, B, 6.6) + 10
    c.setFillColor(_hex("26364A")); c.roundRect(M, H - 52, kw, 13, 3, stroke=0, fill=1)
    c.setFillColor(_hex("FFFFFF")); c.setFont(B, 6.6); c.drawString(M + 5, H - 47.5, kick)
    title, ts = _fit(v["title"], B, 19, 420, 12)
    c.setFillColor(_hex(INK)); c.setFont(B, ts); c.drawString(M + kw + 8, H - 51.5, title)
    c.setFillColor(_hex(SUB)); c.setFont(R, 8.6); c.drawString(M, H - 66, _fit(v["subtitle"], R, 8.6, 420)[0])
    x = W - M
    for val, lab in reversed(v["stats"]):
        vw = max(stringWidth(str(val), B, 12), stringWidth(lab, R, 6.4)) + 4
        c.setFillColor(_hex(INK)); c.setFont(B, 12); c.drawRightString(x, H - 52, _fit(str(val), B, 12, 110)[0])
        c.setFillColor(_hex(MUTED)); c.setFont(R, 6.4); c.drawRightString(x, H - 61, lab)
        x -= min(vw, 112) + 16
    c.setStrokeColor(_hex(LINE)); c.setLineWidth(0.8); c.line(M, H - 74, W - M, H - 74)
    # legend size decides the grid height (rows never shorter than min_rh)
    foot = 20; gtop = H - 84; hdr = 21; wcol = 13 if m["school"]["weeks"] > 1 else 0; dcol = 30
    min_rh = 27 if ND > 6 else 40
    lay = _legend_layout(v, R, B, W - 2 * M, max_h=gtop - hdr - ND * min_rh - 8 - foot - 6)
    leg_top = foot + 6 + lay["h"]
    gx = M + wcol + dcol; cw = (W - M - gx) / NP
    rh = min(72, (gtop - hdr - leg_top - 8) / ND)
    gbot = gtop - hdr - rh * ND
    leg_top = gbot - 12
    # period header
    for p in range(NP):
        cx = gx + p * cw + cw / 2
        c.setFillColor(_hex(MUTED)); c.setFont(R, 5.6); c.drawCentredString(cx, gtop - 8, str(p))
        c.setFillColor(_hex(INK)); c.setFont(B, 6.6 if cw > 22 else 5.8); c.drawCentredString(cx, gtop - 16.5, m["times"][p][0])
    # rows
    per_week = ND // max(1, m["school"]["weeks"])
    for d in range(ND):
        y = gtop - hdr - (d + 1) * rh
        wk, dn = _week_day(m["days"][d])
        c.setFillColor(_hex("FFFFFF" if d % 2 == 0 else "FBFCFD")); c.rect(gx, y, cw * NP, rh, stroke=0, fill=1)
        c.setFillColor(_hex(INK)); c.setFont(B, 8); c.drawCentredString(M + wcol + dcol / 2, y + rh / 2 - 2.8, dn)
        if wcol and d % per_week == 0:
            yb = y - rh * (per_week - 1)
            c.setFillColor(_hex("EEF1F4")); c.roundRect(M, yb + 1, wcol - 2, rh * per_week - 2, 2, stroke=0, fill=1)
            c.saveState(); c.translate(M + wcol / 2 + 1.6, yb + rh * per_week / 2); c.rotate(90)
            c.setFillColor(_hex(SUB)); c.setFont(B, 6); c.drawCentredString(0, 0, f"{(wk or '').upper()} WEEK"); c.restoreState()
        codes = v["codes"][d] if v["codes"] else None
        if codes:
            p = 0
            while p < NP:
                ch = codes[p]; q = p
                while q + 1 < NP and codes[q + 1] == ch: q += 1
                if ch in BREAK or ch == ".":
                    c.setFillColor(_hex("EDF0F3" if ch in BREAK else "F7F8FA")); c.rect(gx + p * cw, y, (q - p + 1) * cw, rh, stroke=0, fill=1)
                    if ch in BREAK:
                        lab, ls = _fit(BREAK[ch], R, 6.2, (q - p + 1) * cw - 3, 4.6)
                        c.setFillColor(_hex(MUTED)); c.setFont(R, ls); c.drawCentredString(gx + (p + q + 1) * cw / 2, y + rh / 2 - 2, lab)
                elif ch == "A":
                    c.setFillColor(_hex("FFFBF0")); c.rect(gx + p * cw, y, (q - p + 1) * cw, rh, stroke=0, fill=1)
                p = q + 1
    c.setStrokeColor(_hex("EEF0F3")); c.setLineWidth(0.4)
    for p in range(NP + 1): c.line(gx + p * cw, gbot, gx + p * cw, gtop - hdr)
    for d in range(ND + 1):
        y = gtop - hdr - d * rh
        strong = d % per_week == 0
        c.setStrokeColor(_hex("AEB6C1" if strong and d not in (0, ND) else LINE)); c.setLineWidth(0.9 if strong and d not in (0, ND) else 0.5)
        c.line(M + wcol, y, gx + NP * cw, y)
    # lessons
    lanes = _lanes(v["cells"])
    for cell in v["cells"]:
        lane, nl = lanes.get(id(cell), (0, 1))
        bh = rh / nl; x0 = gx + cell["p"] * cw + 1; w = cell["n"] * cw - 2
        y0 = gtop - hdr - (cell["d"] + 1) * rh + rh - (lane + 1) * bh + 1; h = bh - 2
        tint, acc = EXT if cell["ext"] else PALETTE[cell["color"] % len(PALETTE)]
        c.setFillColor(_hex(tint)); c.setStrokeColor(_hex(acc)); c.setLineWidth(0.7)
        if cell["ext"]: c.setDash(2, 1.4)
        c.roundRect(x0, y0, w, h, 2.2, stroke=1, fill=1); c.setDash()
        c.setFillColor(_hex(acc)); c.rect(x0 + 0.35, y0 + 0.35, 1.8, h - 0.7, stroke=0, fill=1)
        ts_, ls_ = (7.4, 6.3) if w > 34 else (6.4, 5.6)
        lines = [(cell["title"], B, ts_, INK)] + [(l, R, ls_, SUB if i == 0 else MUTED) for i, l in enumerate(cell["lines"])]
        if cell["ext"]: lines.append(("after school", B, 5.4, "8A5300"))
        room = h - 3
        while lines and sum(s * 1.16 for _, _, s, _ in lines) > room: lines.pop()
        if not lines: lines = [(cell["title"], B, max(4.6, min(ts_, h - 2)), INK)]
        tot = sum(s * 1.16 for _, _, s, _ in lines); ty = y0 + h / 2 + tot / 2
        for txt, f, s, col in lines:
            t2, s2 = _fit(str(txt), f, s, w - 5.5, 4.4)
            ty -= s * 1.16; c.setFillColor(_hex(col)); c.setFont(f, s2); c.drawString(x0 + 3.6, ty + s * 0.22, t2)
    # legend
    ly = leg_top
    c.setFillColor(_hex(INK)); c.setFont(B, 7.4); c.drawString(M, ly - 8, v["legend_title"])
    size, lh = lay["size"], lay["lh"]
    for ci, col in enumerate(lay["cols"]):
        lx = M + ci * (lay["cw"] + 14); yy = ly - 14
        for e, lw, lines, hh in col:
            tint, acc = PALETTE[e["color"] % len(PALETTE)]
            c.setFillColor(_hex(tint)); c.setStrokeColor(_hex(acc)); c.setLineWidth(0.6); c.roundRect(lx, yy - size + 0.6, 5.6, 5.6, 1.2, stroke=1, fill=1)
            c.setFillColor(_hex(INK)); c.setFont(B, size); c.drawString(lx + 9, yy - size + 1.2, e["label"])
            c.setFillColor(_hex(SUB)); c.setFont(R, size)
            for i, ln in enumerate(lines):
                c.drawString(lx + 9 + (lw if i == 0 else 0), yy - size + 1.2 - i * lh, ln)
            yy -= hh
    if lay.get("dropped"):
        c.setFillColor(_hex(MUTED)); c.setFont(R, size); c.drawString(M, foot + 4, f"+ {lay['dropped']} more: see the Excel workbook or the HTML page.")
    # footer
    c.setStrokeColor(_hex(LINE)); c.setLineWidth(0.5); c.line(M, foot, W - M, foot)
    c.setFillColor(_hex(MUTED)); c.setFont(R, 6.2)
    note = "Times are period start times." + (" Shaded columns are breaks; pale yellow cells are after-school slots, used only when needed." if v["codes"] else "")
    c.drawString(M, foot - 9, note)
    if page: c.drawRightString(W - M, foot - 9, f"Page {page[0]} of {page[1]}")

def _canvas(path, m, title):
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4, landscape
    c = canvas.Canvas(path, pagesize=landscape(A4), pageCompression=1)
    c.setTitle(title); c.setAuthor(m["school"]["name"]); c.setSubject(m["school"]["timetable"]); c.setCreator("school-timetable-generator")
    return c

BOOK = {"classes": "Classes", "teachers": "Teachers", "rooms": "Rooms", "students": "Students"}

def write_pdfs(m, outdir, kinds=KINDS, individual=True, log=print):
    """writes outdir/<Kind> - all.pdf (one page each, bookmarked) and, if individual, outdir/<Kind>/<name>.pdf"""
    os.makedirs(outdir, exist_ok=True); files = []
    for kind in kinds:
        vs = list(all_views(m, kind))
        if not vs: continue
        path = os.path.join(outdir, f"{BOOK[kind]} - all.pdf")
        c = _canvas(path, m, f"{m['school']['name']} {m['school']['timetable']}: {BOOK[kind].lower()}")
        last_group = None
        for i, v in enumerate(vs, start=1):
            key = f"{kind}-{i}"; c.bookmarkPage(key)
            if v["group"] != last_group:
                c.addOutlineEntry(str(v["group"] or "-"), key, level=0); last_group = v["group"]
            c.addOutlineEntry(f"{v['key']} {v['title']}" if kind == "students" else v["title"], key, level=1)
            draw_page(c, m, v, page=(i, len(vs))); c.showPage()
        c.showOutline(); c.save(); files.append(path)
        if individual:
            for v in vs:
                sub = os.path.join(outdir, BOOK[kind], safe_name(v["group"])) if kind == "students" else os.path.join(outdir, BOOK[kind])
                os.makedirs(sub, exist_ok=True)
                name = f"{v['key']} {v['title']}" if kind == "students" else v["key"]
                p = os.path.join(sub, safe_name(name) + ".pdf"); c = _canvas(p, m, v["title"]); draw_page(c, m, v); c.showPage(); c.save()
        log(f"  PDF {BOOK[kind].lower()}: {len(vs)} pages")
    return files


# ----------------------------------------------------------------------------------------------- Excel
XF = lambda **k: Font(name="Calibri", size=k.pop("size", 9), **k)
NAVY = "26364A"
THIN = Side(style="thin", color="C9D0D8")

def _merge(ws, r1, c1, r2, c2):
    """merge without openpyxl's containment scan (quadratic on big sheets); our ranges never overlap"""
    ws.merged_cells.ranges.add(MergedCellRange(ws, f"{CL(c1)}{r1}:{CL(c2)}{r2}"))

def write_excel(m, path, log=print):
    ch = m["checks"]; NP, ND = m["NP"], m["ND"]; sc = m["school"]
    wb = openpyxl.Workbook(); rm = wb.active; rm.title = "Read me"
    lines = [(f"{sc['name']} · {sc['timetable']} timetable", "t"),
             (f"{len(m['classes'])} classes, {len([t for t in m['teachers'] if any(t['name'] in x['teachers'] for x in m['lessons'])])} teachers, "
              f"{ch['lessons']} lessons placed" + (f", {len(m['students'])} students" if m["students"] else "") + "." + (f" Solver status: {m['status']}." if m.get("status") else ""), ""),
             ("Checks", "h")]
    lines += [("All checks passed: no teacher, room, class or student is double-booked; every subject has its planned periods; nothing sits in a break.", "ok")] \
        if ch["ok"] else [(p, "bad") for p in ch["problems"][:60]]
    if ch["side_by_side"]: lines.append((f"{ch['side_by_side']} class-periods have two banded lessons side by side for different students (for example extra periods for one group).", ""))
    lines += [("After-school lessons", "h")] + ([(a, "") for a in ch["after_school"]] or [("None.", "")])
    if m["notes"]: lines += [("Notes", "h")] + [(n, "") for n in m["notes"]]
    lines += [("How to read", "h"),
              ("Classes, Teachers, Rooms and Students sheets: one timetable each, days down the side and periods across (start times in the header). Each prints on its own landscape page.", ""),
              ("Banded lessons on a class timetable show the block and its subjects; the lines under the grid list every group, teacher and room. Hover a banded cell for details.", ""),
              ("Lessons sheet: one row per group per lesson (filter it like any table). Teacher load: periods per teacher. Contents: a link to every timetable.", "")]
    for i, (t, k) in enumerate(lines, start=1):
        c = rm.cell(i, 1, t); c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(name="Calibri", size=16, bold=True, color=NAVY) if k == "t" else Font(name="Calibri", size=12, bold=True, color=NAVY) if k == "h" else \
            Font(name="Calibri", size=10, color="2E7D32" if k == "ok" else "B00020" if k == "bad" else "1B2330")
    rm.column_dimensions["A"].width = 120
    rm.page_setup.fitToWidth = 1; rm.page_setup.fitToHeight = 0; rm.sheet_properties.pageSetUpPr.fitToPage = True
    cs = wb.create_sheet("Contents"); contents_row = 1
    fills = {i: PatternFill("solid", fgColor=t) for i, (t, _) in enumerate(PALETTE)}
    extfill = PatternFill("solid", fgColor=EXT[0]); brk = PatternFill("solid", fgColor="EDF0F3"); noschool = PatternFill("solid", fgColor="F7F8FA")
    hdrfill = PatternFill("solid", fgColor=NAVY); wrapc = Alignment(wrap_text=True, vertical="center", horizontal="center")
    borders = {a: Border(left=Side(style="thick", color=a), top=Side(style="thin", color="FFFFFF"), bottom=Side(style="thin", color="FFFFFF")) for _, a in PALETTE + [EXT]}
    links = collections.OrderedDict()

    def grid_sheet(name, kind):
        vs = list(all_views(m, kind))
        if not vs: return
        ws = wb.create_sheet(name); ws.column_dimensions["A"].width = 10
        for p in range(NP): ws.column_dimensions[CL(p + 2)].width = 6.4 if NP > 16 else 9.5
        r0 = 1; links[name] = []
        for v in vs:
            links[name].append((v, r0))
            ws.cell(r0, 1, f"{v['title']}").font = XF(bold=True, size=13, color=NAVY)
            extra = [f"{a} {b}" for a, b in v["stats"] if str(a) not in v["subtitle"]]
            ws.cell(r0 + 1, 1, v["subtitle"] + ("   ·   " + "   ".join(extra) if extra else "")).font = XF(size=9, color="4A5566")
            hr = r0 + 2; ws.cell(hr, 1, "Day").font = XF(bold=True, color="FFFFFF"); ws.cell(hr, 1).fill = hdrfill
            for p in range(NP):
                c = ws.cell(hr, p + 2, f"{p}\n{m['times'][p][0]}"); c.font = XF(bold=True, color="FFFFFF", size=8); c.fill = hdrfill; c.alignment = wrapc
            ws.row_dimensions[hr].height = 26
            at = collections.defaultdict(list)
            for cell in v["cells"]:
                for q in range(cell["p"], cell["p"] + cell["n"]): at[(cell["d"], q)].append(cell)
            for d in range(ND):
                rr = hr + 1 + d; ws.row_dimensions[rr].height = 40
                c = ws.cell(rr, 1, m["days"][d]); c.font = XF(bold=True); c.alignment = Alignment(vertical="center")
                p = 0
                while p < NP:
                    cs = at.get((d, p))
                    if cs:
                        key = tuple(id(x) for x in cs); q = p
                        while q + 1 < NP and tuple(id(x) for x in at.get((d, q + 1), [])) == key: q += 1
                        if len(cs) == 1: txt = "\n".join([cs[0]["title"]] + cs[0]["lines"])
                        else: txt = "\n".join(f"{x['title']}" + (f" ({x['lines'][0]})" if x["lines"] else "") for x in cs)
                        c = ws.cell(rr, p + 2, txt); c.font = XF(size=8); c.alignment = wrapc
                        c.fill = extfill if any(x["ext"] for x in cs) else fills[cs[0]["color"] % len(PALETTE)]
                        c.border = borders[EXT[1] if any(x["ext"] for x in cs) else PALETTE[cs[0]["color"] % len(PALETTE)][1]]
                        if kind == "classes" and cs[0]["lines"] and "groups" in cs[0]["lines"][-1]:
                            det = [f"{x['group']}: {x['subject']} - {', '.join(x['teachers'])} - {x['room']}" for x in m["lessons"] if x["id"] == cs[0]["id"] and v["key"] in x["classes"]]
                            c.comment = Comment("\n".join(det), "Timetable"); c.comment.width = 340; c.comment.height = 30 + 14 * len(det)
                        if q > p: _merge(ws, rr, p + 2, rr, q + 2)
                        p = q + 1; continue
                    code = v["codes"][d][p] if v["codes"] else "L"
                    if code in BREAK or code == ".":
                        q = p
                        while q + 1 < NP and not at.get((d, q + 1)) and v["codes"][d][q + 1] == code: q += 1
                        c = ws.cell(rr, p + 2, BREAK.get(code, "")); c.fill = brk if code in BREAK else noschool; c.font = XF(size=7, color="7A8494"); c.alignment = wrapc
                        if q > p: _merge(ws, rr, p + 2, rr, q + 2)
                        p = q + 1; continue
                    p += 1
            lr = hr + ND + 1
            ws.cell(lr, 1, v["legend_title"]).font = XF(bold=True, size=9, color=NAVY)
            for e in v["legend"]:
                lr += 1
                c = ws.cell(lr, 1, e["label"]); c.font = XF(bold=True, size=8); c.fill = fills[e["color"] % len(PALETTE)]; c.alignment = Alignment(vertical="top", wrap_text=True)
                c2 = ws.cell(lr, 2, e["text"]); c2.font = XF(size=8); c2.alignment = Alignment(wrap_text=True, vertical="top")
                _merge(ws, lr, 2, lr, NP + 1)
                ws.row_dimensions[lr].height = 12 * max(1, -(-len(e["text"]) // (NP * (6 if NP > 16 else 9))))
            r0 = lr + 2
            ws.row_breaks.append(Break(id=lr + 1))
        ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = ws.page_margins.right = 0.3; ws.page_margins.top = ws.page_margins.bottom = 0.4
        ws.sheet_view.zoomScale = 90
        log(f"  Excel {name.lower()}: {len(vs)} timetables")

    grid_sheet("Classes", "classes"); grid_sheet("Teachers", "teachers"); grid_sheet("Rooms", "rooms"); grid_sheet("Students", "students")
    # contents with links
    r = contents_row
    cs.cell(r, 1, "Contents: click a name to open its timetable").font = XF(bold=True, size=13, color=NAVY); r += 2
    for name, items in links.items():
        cs.cell(r, 1, f"{name} ({len(items)})").font = XF(bold=True, size=11, color=NAVY); r += 1
        per_row = 6
        for i in range(0, len(items), per_row):
            for j, (v, row) in enumerate(items[i:i + per_row]):
                label = f"{v['key']} {v['title']}" if name == "Students" else v["key"]
                c = cs.cell(r, 1 + j, label); c.hyperlink = f"#'{name}'!A{row}"; c.font = XF(size=9, color="1B5FAA", underline="single")
            r += 1
        r += 1
    for j in range(1, 7): cs.column_dimensions[CL(j)].width = 28
    cs.page_setup.fitToWidth = 1; cs.page_setup.fitToHeight = 0; cs.sheet_properties.pageSetUpPr.fitToPage = True
    # lessons (hand-over format) and teacher load
    ws = wb.create_sheet("Lessons"); ws.append(LESSON_HEADERS)
    for x in sorted(m["lessons"], key=lambda x: (x["id"], x["group"])):
        ws.append([x["id"], m["days"][x["d"]], x["p"], x["p"] + x["n"] - 1, x["n"], m["times"][x["p"]][0], m["times"][min(NP - 1, x["p"] + x["n"] - 1)][1],
                   x["level"], x["block"], ", ".join(x["classes"]), x["subject"], x["group"], ", ".join(x["teachers"]), x["room"], "yes" if x["ext"] else ""])
    for c in ws[1]: c.font = Font(name="Calibri", bold=True, color="FFFFFF", size=10); c.fill = hdrfill
    for col, w in zip("ABCDEFGHIJKLMNO", (8, 11, 7, 7, 7, 7, 7, 7, 24, 22, 34, 12, 34, 14, 8)): ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    ws = wb.create_sheet("Teacher load")
    ws.append(["Teacher", "Short", "Department", "Allowance", "Periods timetabled", "Most periods in a day", "Days teaching", "After-school periods"])
    for t in sorted(m["teachers"], key=lambda t: (str(t["dept"]), t["name"])):
        per_day = collections.Counter(); ex = 0
        for lid, xs in _by_id([x for x in m["lessons"] if t["name"] in x["teachers"]]).items():
            per_day[xs[0]["d"]] += xs[0]["n"]; ex += xs[0]["n"] if xs[0]["ext"] else 0
        ws.append([t["name"], t["short"], t["dept"], t["allowance"], sum(per_day.values()), max(per_day.values()) if per_day else 0, len(per_day), ex])
    for c in ws[1]: c.font = Font(name="Calibri", bold=True, color="FFFFFF", size=10); c.fill = hdrfill; c.alignment = Alignment(wrap_text=True)
    for col, w in zip("ABCDEFGH", (28, 14, 12, 10, 12, 12, 10, 12)): ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    wb.save(path)
    return path


# ----------------------------------------------------------------------------------------------- HTML
def html_data(m, view="full", student_id=None, app=None):
    """data for the finder page. view: full (committee/teacher: everything), student (class timetables and the
    builder, no student list), personal (one student's own lessons only). Filtering happens here, on the server."""
    pal = [[t, a] for t, a in PALETTE]
    d = dict(school=m["school"], generated=m["generated"], days=m["days"], NP=m["NP"], times=m["times"], grid=m["grid"], levels=m["levels"],
             classes=[[c["name"], c["level"], c["size"], c["home"]] for c in m["classes"]],
             teachers=[[t["name"], t["short"], t["dept"], t["allowance"]] for t in m["teachers"]],
             rooms=[[r["name"], r["kind"], m["room_cap"].get(r["name"], 1)] for r in m["rooms"]],
             subjects={k: [v["short"], v["color"]] for k, v in m["subjects"].items()}, palette=pal, ext=list(EXT),
             lessons=[[x["id"], x["d"], x["p"], x["n"], x["level"], x["block"], x["classes"], x["subject"], x["group"], x["teachers"], x["room"], int(x["ext"])]
                      for x in m["lessons"]],
             students=[[s["id"], s["name"], s["cls"], s["groups"]] for s in m["students"]],
             checks=dict(ok=m["checks"]["ok"], problems=m["checks"]["problems"], after=m["checks"]["after_school"], notes=m["notes"]), app=app)
    if view in ("student", "personal"):
        d["teachers"] = [[t[0], t[1], "", ""] for t in d["teachers"]]
        d["checks"] = dict(ok=m["checks"]["ok"], problems=[], after=[], notes=[])
        d["rooms"] = []
    if view == "student":
        d["students"] = []
    if view == "personal":
        s = next(x for x in m["students"] if x["id"] == student_id); gs = set(s["groups"])
        d["lessons"] = [x for x in d["lessons"] if (x[8] == ENTIRE and s["cls"] in x[6]) or (x[8] != ENTIRE and x[8] in gs)]
        d["classes"] = [c for c in d["classes"] if c[0] == s["cls"]]
        names = {t for x in d["lessons"] for t in x[9]}
        d["teachers"] = [t for t in d["teachers"] if t[0] in names]
        d["students"] = [[s["id"], s["name"], s["cls"], s["groups"]]]
    return d

def render_html(m, view="full", student_id=None, app=None, title=None):
    here = os.path.dirname(os.path.abspath(__file__))
    tpl = open(os.path.join(here, "timetable_page.html"), encoding="utf-8").read()
    data = json.dumps(html_data(m, view, student_id, app), ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    title = title or f"{m['school']['name']} · {m['school']['timetable']}"
    return tpl.replace("/*__DATA__*/null", data).replace("__TITLE__", title.replace("&", "&amp;").replace("<", "&lt;"))

def write_html(m, path, view="full"):
    with open(path, "w", encoding="utf-8") as f: f.write(render_html(m, view))
    return path

def custom_view(m, cls, groups):
    """a student's timetable from their class and chosen groups (the finder's 'Build a timetable')"""
    info = m["class_ix"][cls]
    v = student_view(m, dict(id="custom", name=f"Class {cls}", cls=cls, level=info["level"], groups=list(groups)))
    v["subtitle"] = f"Class {cls} · {info['level']} · groups {', '.join(groups) or '-'}"
    return v

def group_options(m, cls):
    """{block or group: [(group code, subject, teachers)]} for one class, for the builder"""
    out = collections.OrderedDict()
    for x in m["lessons"]:
        if x["group"] != ENTIRE and cls in x["classes"]:
            out.setdefault(x["block"] or x["group"], collections.OrderedDict()).setdefault(x["group"], (x["subject"], x["teachers"]))
    return out

def pdf_bytes(m, v):
    import io
    buf = io.BytesIO(); c = _canvas(buf, m, v["title"]); draw_page(c, m, v); c.showPage(); c.save()
    return buf.getvalue()

def code_slips_pdf(m, codes, url, out):
    """printable slips (10 to an A4 page, a new page per class) with each student's personal code"""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    R, B = _fonts(); W, H = A4; c = canvas.Canvas(out, pagesize=A4); c.setTitle(f"{m['school']['name']} timetable codes")
    order = {cl["name"]: i for i, cl in enumerate(m["classes"])}
    studs = sorted(m["students"], key=lambda s: (order.get(s["cls"], 999), s["id"]))
    by = collections.OrderedDict()
    for s in studs: by.setdefault(s["cls"], []).append(s)
    M, cols, rows = 28, 2, 5; sw, sh = (W - 2 * M) / cols, (H - 2 * M - 16) / rows
    for cls, ss in by.items():
        for k in range(0, len(ss), cols * rows):
            c.setFont(B, 9); c.setFillColor(_hex(MUTED)); c.drawString(M, H - M + 6, f"{m['school']['name']} · {m['school']['timetable']} · Class {cls} · cut along the dashed lines")
            for i, s in enumerate(ss[k:k + cols * rows]):
                x = M + (i % cols) * sw; y = H - M - 16 - (i // cols + 1) * sh
                c.setDash(3, 3); c.setStrokeColor(_hex("B8BEC6")); c.setLineWidth(0.6); c.rect(x + 4, y + 4, sw - 8, sh - 8); c.setDash()
                c.setFillColor(_hex(MUTED)); c.setFont(R, 8.5); c.drawString(x + 16, y + sh - 24, f"{m['school']['name']} · your timetable code")
                c.setFillColor(_hex(INK)); c.setFont(B, 13); c.drawString(x + 16, y + sh - 44, _fit(s["name"], B, 13, sw - 40)[0])
                c.setFillColor(_hex(SUB)); c.setFont(R, 9.5); c.drawString(x + 16, y + sh - 59, f"Class {s['cls']} · {s['id']}")
                c.setFillColor(_hex(INK)); c.setFont("Courier-Bold", 24); c.drawString(x + 16, y + sh - 98, codes.get(s["id"], "-"))
                c.setFillColor(_hex(SUB)); t, ts = _fit(f"Go to {url}", B, 8.5, sw - 36, 6); c.setFont(B, ts); c.drawString(x + 16, y + 42, t)
                c.setFont(R, 8.2); c.drawString(x + 16, y + 30, "and enter this code to see your own timetable.")
                c.drawString(x + 16, y + 18, "Keep it to yourself: anyone with the code can see it.")
            c.showPage()
    c.save()
    return out


# ----------------------------------------------------------------------------------------------- everything at once
def export_all(m, outdir, pdf=True, individual=True, zip_individual=True, html=True, excel=True, log=print):
    os.makedirs(outdir, exist_ok=True); out = {}
    base = safe_name(f"{m['school']['name']} {m['school']['timetable']}")
    if excel: out["excel"] = write_excel(m, os.path.join(outdir, f"{base} - all timetables.xlsx"), log=log)
    if html: out["html"] = write_html(m, os.path.join(outdir, f"{base} - timetable finder.html"))
    if pdf:
        pdir = os.path.join(outdir, "PDF")
        out["pdf_books"] = write_pdfs(m, pdir, individual=individual, log=log)
        if individual and zip_individual:
            zp = os.path.join(outdir, f"{base} - individual PDFs.zip")
            with zipfile.ZipFile(zp, "w", zipfile.ZIP_DEFLATED) as z:
                for kind in KINDS:
                    root = os.path.join(pdir, BOOK[kind])
                    for dp, _, fs in os.walk(root):
                        for f in sorted(fs): z.write(os.path.join(dp, f), os.path.relpath(os.path.join(dp, f), pdir))
            out["pdf_zip"] = zp
    return out
